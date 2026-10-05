import os
os.environ["TESTING"] = "True"
import openpyxl
from sheets_manager import SheetsManager
from receipt_processor import FinancialReportPDFGenerator, FinancialReportExcelGenerator

sm = SheetsManager()
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "static", "receipts")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def test_payments_report_data_and_generation():
    print("Testing Payments Report data aggregation, PDF, and Excel generation...")
    data = sm.get_payments_report_data()
    assert "records" in data
    assert "rows" in data
    assert "kpis" in data
    assert "total_collected" in data["kpis"]
    assert isinstance(data["records"], list)
    print(f"  Found {len(data['records'])} payment records. Total: ${data['kpis']['total_collected']:.2f}")

    # Generate PDF
    pdf_path = os.path.join(OUTPUT_DIR, "test_reporte_pagos.pdf")
    FinancialReportPDFGenerator.generate_payments_report_pdf(
        dest_path=pdf_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(pdf_path), f"PDF file was not created: {pdf_path}"
    with open(pdf_path, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-", "Generated file is not a valid PDF"
    print(f"  Payments PDF successfully created: {pdf_path}")

    # Generate Excel
    excel_path = os.path.join(OUTPUT_DIR, "test_reporte_pagos.xlsx")
    FinancialReportExcelGenerator.generate_payments_report_excel(
        dest_path=excel_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(excel_path), f"Excel file was not created: {excel_path}"
    wb = openpyxl.load_workbook(excel_path)
    assert "Reporte Pagos" in wb.sheetnames or len(wb.sheetnames) > 0
    ws = wb.active
    assert ws["A1"].value == "CONDOMINIO EL MIRADOR"
    print(f"  Payments Excel successfully created: {excel_path}")
    print("[OK] Payments Report test PASSED!")

def test_debtors_report_data_and_generation():
    print("\nTesting Debtors Report data aggregation, PDF, and Excel generation...")
    data = sm.get_debtors_report_data(status_filter="all")
    assert "records" in data
    assert "rows" in data
    assert "kpis" in data
    assert "total_debt" in data["kpis"]
    assert "debtors_count" in data["kpis"]
    assert isinstance(data["records"], list)
    print(f"  Found {len(data['records'])} apartment records. Total debt: ${data['kpis']['total_debt']:.2f}")

    # Generate PDF
    pdf_path = os.path.join(OUTPUT_DIR, "test_reporte_deudores.pdf")
    FinancialReportPDFGenerator.generate_debtors_report_pdf(
        dest_path=pdf_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(pdf_path), f"PDF file was not created: {pdf_path}"
    with open(pdf_path, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-", "Generated file is not a valid PDF"
    print(f"  Debtors PDF successfully created: {pdf_path}")

    # Generate Excel
    excel_path = os.path.join(OUTPUT_DIR, "test_reporte_deudores.xlsx")
    FinancialReportExcelGenerator.generate_debtors_report_excel(
        dest_path=excel_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(excel_path), f"Excel file was not created: {excel_path}"
    wb = openpyxl.load_workbook(excel_path)
    assert len(wb.sheetnames) > 0
    ws = wb.active
    assert ws["A1"].value == "CONDOMINIO EL MIRADOR"
    print(f"  Debtors Excel successfully created: {excel_path}")
    print("[OK] Debtors Report test PASSED!")

def test_unreconciled_report_data_and_generation():
    print("\nTesting Unreconciled Report data aggregation, PDF, and Excel generation...")
    data = sm.get_unreconciled_report_data()
    assert "records" in data
    assert "rows" in data
    assert "kpis" in data
    assert "total_unreconciled_amount" in data["kpis"]
    assert isinstance(data["records"], list)
    print(f"  Found {len(data['records'])} unreconciled records. Total: ${data['kpis']['total_unreconciled_amount']:.2f}")

    # Generate PDF
    pdf_path = os.path.join(OUTPUT_DIR, "test_reporte_sin_conciliar.pdf")
    FinancialReportPDFGenerator.generate_unreconciled_report_pdf(
        dest_path=pdf_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(pdf_path), f"PDF file was not created: {pdf_path}"
    with open(pdf_path, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-", "Generated file is not a valid PDF"
    print(f"  Unreconciled PDF successfully created: {pdf_path}")

    # Generate Excel
    excel_path = os.path.join(OUTPUT_DIR, "test_reporte_sin_conciliar.xlsx")
    FinancialReportExcelGenerator.generate_unreconciled_report_excel(
        dest_path=excel_path,
        report_data=data,
        condo_name="Condominio El Mirador",
        currency="$"
    )
    assert os.path.exists(excel_path), f"Excel file was not created: {excel_path}"
    wb = openpyxl.load_workbook(excel_path)
    assert len(wb.sheetnames) > 0
    ws = wb.active
    assert ws["A1"].value == "CONDOMINIO EL MIRADOR"
    print(f"  Unreconciled Excel successfully created: {excel_path}")
    print("[OK] Unreconciled Report test PASSED!")

def test_filtered_report_data():
    print("\nTesting filtered queries...")
    data_filtered = sm.get_payments_report_data(unit="101")
    for r in data_filtered["records"]:
        if r["unit"] and r["unit"] != "N/A":
            assert str(r["unit"]) == "101"
    print("  Unit filter verified.")

    debtors_only = sm.get_debtors_report_data(status_filter="debtors_only")
    for r in debtors_only["records"]:
        assert r["total_debt"] > 0.01
    print("  Debtors only filter verified.")
def test_paid_additional_charges_report_amounts():
    print("\nTesting Paid Additional Charges / Deuda Inicial in Payments Report...")
    # Add a temporary additional charge
    sm.data["additional_charges"].append({
        "id": "TEST-AC-REPORT-99",
        "unit": "909",
        "type": "Deuda Inicial",
        "description": "Deuda prueba",
        "amount": 0.0,
        "paid_amount": 75.50,
        "status": "Pagado",
        "payment_date": "2026-09-20",
        "reference": "REF99999",
        "paid_by": "Residente Prueba"
    })
    
    data = sm.get_payments_report_data(unit="909")
    target_row = next((r for r in data["rows"] if r["id"] == "TEST-AC-REPORT-99"), None)
    assert target_row is not None, "Target paid additional charge not found in report rows"
    assert target_row["base_amount"] == 75.50, f"Expected base_amount 75.50, got {target_row['base_amount']}"
    assert target_row["late_fee"] == 0.0, f"Expected late_fee 0.0, got {target_row['late_fee']}"
    assert target_row["amount"] == 75.50, f"Expected amount 75.50, got {target_row['amount']}"
    assert data["kpis"]["total_collected"] >= 75.50
    assert data["kpis"]["total_charges"] >= 75.50
    
    # Clean up
    sm.data["additional_charges"] = [c for c in sm.data["additional_charges"] if c["id"] != "TEST-AC-REPORT-99"]
    print("[OK] Paid Additional Charges in Payments Report test PASSED!")

def test_overdue_aliquots_debt_and_late_fees():
    print("\nTesting Overdue Aliquot Late Fee and Debt Calculation in Debtors Report...")
    test_unit = "TEST-DEBTOR-99"
    # Create test unit
    sm.data["units"] = [u for u in sm.data["units"] if u["id"] != test_unit]
    sm.data["units"].append({
        "id": test_unit,
        "owner": "Propietario Moroso Test",
        "phone1": "0991234567",
        "email1": "moroso@test.com",
        "aliquot_base": 70.0
    })
    
    # 1. Add an overdue aliquot (Enero 2026 -> definitely past the 15th)
    aliquot_overdue = {
        "id": f"{test_unit}-2026-Enero",
        "unit": test_unit,
        "month": "Enero",
        "year": 2026,
        "amount": 70.0,
        "late_fee": 0.0,
        "paid_amount": 0.0,
        "status": "Pendiente",
        "payment_date": "",
        "reference": ""
    }
    
    # 2. Add an additional charge of $30.00
    charge_extra = {
        "id": f"AC-{test_unit}",
        "unit": test_unit,
        "type": "Multa Convivencia",
        "description": "Ruidos molestos",
        "amount": 30.0,
        "paid_amount": 0.0,
        "status": "Pendiente",
        "payment_date": "",
        "reference": ""
    }
    
    sm.data["aliquots"].append(aliquot_overdue)
    sm.data["additional_charges"].append(charge_extra)
    
    # Query debtors report for this unit
    rep = sm.get_debtors_report_data(unit=test_unit)
    target_row = next((r for r in rep["rows"] if r["unit"] == test_unit), None)
    
    assert target_row is not None, "Target debtor unit not found in report"
    assert target_row["aliquots_debt"] == 70.0, f"Expected aliquots_debt 70.0, got {target_row['aliquots_debt']}"
    assert target_row["late_fees"] == 10.0, f"Expected late_fees 10.0, got {target_row['late_fees']}"
    assert target_row["charges_debt"] == 30.0, f"Expected charges_debt 30.0, got {target_row['charges_debt']}"
    # Total debt must be 70 (aliquot) + 10 (late fee) + 30 (charge) = 110.00
    assert target_row["total_debt"] == 110.0, f"Expected total_debt 110.0, got {target_row['total_debt']}"
    assert "Multa: $10.00" in target_row["periods_str"]
    
    # Check get_debtors summary as well
    debtors_sum = sm.get_debtors()
    d_entry = next((d for d in debtors_sum if d["unit"] == test_unit), None)
    assert d_entry is not None
    assert d_entry["total_debt"] == 110.0
    assert d_entry["late_fees"] == 10.0
    
    # Clean up test records
    sm.data["units"] = [u for u in sm.data["units"] if u["id"] != test_unit]
    sm.data["aliquots"] = [a for a in sm.data["aliquots"] if a["id"] != aliquot_overdue["id"]]
    sm.data["additional_charges"] = [c for c in sm.data["additional_charges"] if c["id"] != charge_extra["id"]]
    
    print("[OK] Overdue Aliquot Late Fee and Debt Calculation test PASSED!")

if __name__ == "__main__":
    test_payments_report_data_and_generation()
    test_debtors_report_data_and_generation()
    test_unreconciled_report_data_and_generation()
    test_filtered_report_data()
    test_paid_additional_charges_report_amounts()
    test_overdue_aliquots_debt_and_late_fees()
    print("\nALL FINANCIAL REPORT TESTS COMPLETED SUCCESSFULLY!")
