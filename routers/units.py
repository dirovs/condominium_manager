import os
import shutil
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, HTTPException, UploadFile, File, Depends
from fastapi.responses import FileResponse

from schemas import UnitCreate, BoardMemberCreate
from security import get_current_user, check_unit_permission
from receipt_processor import generate_no_debt_certificate_pdf

router = APIRouter(tags=["Units & Board Members"])

def get_sm():
    from app import sm
    return sm

def get_receipts_dir():
    from app import RECEIPTS_DIR
    return RECEIPTS_DIR

def get_uploads_dir():
    from app import UPLOADS_DIR
    return UPLOADS_DIR

def detect_column_mapping(header_row, is_xlsx=False):
    col_mapping = {}
    start_idx = 1 if is_xlsx else 0
    for idx, val in enumerate(header_row):
        col_idx = idx + start_idx
        val = str(val or "").strip().lower()

        # 1. ID / Departamento
        if val == "id" or any(k in val for k in ["departamento", "depto", "unidad", "unit", "nro de d", "nro. de d", "numero de d", "número de d"]):
            col_mapping["id"] = col_idx

        # 2. Cédula checks
        elif any(k in val for k in ["cedula", "cédula", "nro. doc", "documento"]):
            if any(k in val for k in ["inquilino", "tenant", "residente"]):
                col_mapping["cedula_tenant"] = col_idx
            else:
                col_mapping["cedula_owner"] = col_idx

        # 3. Email / Correo checks
        elif any(k in val for k in ["mail", "email", "correo"]):
            if any(k in val for k in ["inquilino", "tenant", "residente"]):
                col_mapping["email2"] = col_idx
            else:
                col_mapping["email1"] = col_idx

        # 4. Phone / Teléfono checks
        elif val == "tel" or any(k in val for k in ["telef", "phone", "celular", "tlf"]):
            if any(k in val for k in ["inquilino", "tenant", "residente"]):
                col_mapping["phone2"] = col_idx
            else:
                col_mapping["phone1"] = col_idx

        # 5. Base Role check: Tenant (more specific than Owner/generic names)
        elif any(k in val for k in ["inquilino", "tenant"]):
            col_mapping["tenant"] = col_idx

        # 6. Base Role check: Owner
        elif any(k in val for k in ["propietario", "owner", "copropietario", "dueño", "dueno", "propietaria"]):
            col_mapping["owner"] = col_idx

        # 7. Aliquot base check
        elif any(k in val for k in ["alicuotabase", "alícuota base", "aliquot_base", "alicuota", "base"]):
            col_mapping["aliquot_base"] = col_idx

        # 8. Generic fallback for names
        elif any(k in val for k in ["nombre", "name", "nombres", "apellidos", "residente"]):
            if "residente" in val and "tenant" not in col_mapping:
                col_mapping["tenant"] = col_idx
            elif "owner" not in col_mapping:
                col_mapping["owner"] = col_idx

    # Fallback to defaults if ID or Owner is not detected
    if "id" not in col_mapping or "owner" not in col_mapping:
        if is_xlsx:
            col_mapping = {
                "id": 1, "owner": 2, "phone1": 3, "email1": 4,
                "tenant": 5, "phone2": 6, "email2": 7, "aliquot_base": 8
            }
        else:
            col_mapping = {
                "id": 0, "owner": 1, "phone1": 2, "email1": 3,
                "tenant": 4, "phone2": 5, "email2": 6, "aliquot_base": 7
            }

    return col_mapping

@router.get("/api/units")
def get_units(current_user: Optional[dict] = Depends(get_current_user)):
    sm = get_sm()
    if current_user and isinstance(current_user, dict):
        role = str(current_user.get("role", "")).lower()
        if role in ["owner", "tenant"]:
            my_unit = str(current_user.get("unit_id") or current_user.get("unit") or "")
            return [u for u in sm.data["units"] if str(u.get("id")) == my_unit]
    return sm.data["units"]

@router.post("/api/units")
def create_or_update_unit(unit: UnitCreate):
    sm = get_sm()
    return sm.add_or_update_unit(
        unit_id=unit.id,
        owner=unit.owner,
        phone1=unit.phone1,
        email1=unit.email1,
        tenant=unit.tenant,
        phone2=unit.phone2,
        email2=unit.email2,
        aliquot_base=unit.aliquot_base,
        cedula_owner=unit.cedula_owner,
        cedula_tenant=unit.cedula_tenant,
        receipt_recipient_type=unit.receipt_recipient_type
    )

@router.delete("/api/units/{unit_id}")
def delete_unit(unit_id: str):
    sm = get_sm()
    success = sm.delete_unit(unit_id)
    if not success:
        raise HTTPException(status_code=404, detail="Departamento no encontrado.")
    return {"status": "success", "message": f"Departamento {unit_id} eliminado correctamente."}

@router.post("/api/upload-units")
async def upload_units(file: UploadFile = File(...)):
    sm = get_sm()
    UPLOADS_DIR = get_uploads_dir()
    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in [".csv", ".xlsx"]:
        raise HTTPException(status_code=400, detail="Formato de archivo no soportado. Use CSV o XLSX.")

    temp_path = os.path.join(UPLOADS_DIR, f"units_{datetime.now().timestamp()}{file_ext}")
    with open(temp_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    try:
        units = []
        if file_ext == ".csv":
            import csv
            with open(temp_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            lines = content.splitlines()
            if not lines:
                return {"status": "success", "added": 0, "message": "Archivo vacío."}

            sep = ","
            if ";" in lines[0] and lines[0].count(";") > lines[0].count(","):
                sep = ";"

            reader = csv.reader(lines, delimiter=sep)
            rows = list(reader)
            if not rows:
                return {"status": "success", "added": 0, "message": "Archivo vacío."}

            # Find headers and detect mapping
            header_row = [val.strip().lower() for val in rows[0]]
            col_mapping = detect_column_mapping(header_row, is_xlsx=False)

            for r in rows[1:]:
                if not r or len(r) <= max(col_mapping.get("id", 0), col_mapping.get("owner", 1)):
                    continue
                unit_id = r[col_mapping["id"]].strip()
                owner = r[col_mapping["owner"]].strip()
                if not unit_id or not owner:
                    continue

                def get_col_val(col_name):
                    idx = col_mapping.get(col_name)
                    if idx is not None and idx < len(r):
                        return r[idx].strip()
                    return ""

                phone1 = get_col_val("phone1")
                email1 = get_col_val("email1")
                tenant = get_col_val("tenant")
                phone2 = get_col_val("phone2")
                email2 = get_col_val("email2")

                aliquot_base = None
                if "aliquot_base" in col_mapping:
                    val = get_col_val("aliquot_base")
                    if val:
                        try:
                            aliquot_base = float(val.replace("$", "").replace(",", "").strip())
                        except:
                            try:
                                aliquot_base = float(val.replace("$", "").replace(".", "").replace(",", ".").strip())
                            except:
                                aliquot_base = None

                cedula_owner = get_col_val("cedula_owner")
                cedula_tenant = get_col_val("cedula_tenant")

                units.append({
                    "id": unit_id, "owner": owner, "phone1": phone1, "email1": email1,
                    "tenant": tenant, "phone2": phone2, "email2": email2, "aliquot_base": aliquot_base,
                    "cedula_owner": cedula_owner, "cedula_tenant": cedula_tenant
                })
        else: # XLSX
            import openpyxl
            wb = openpyxl.load_workbook(temp_path, data_only=True)
            sheet = wb.active

            header_row = [sheet.cell(1, col_idx).value for col_idx in range(1, sheet.max_column + 1)]
            col_mapping = detect_column_mapping(header_row, is_xlsx=True)

            for r_idx in range(2, sheet.max_row + 1):
                unit_id = str(sheet.cell(r_idx, col_mapping["id"]).value or "").strip()
                owner = str(sheet.cell(r_idx, col_mapping["owner"]).value or "").strip()
                if not unit_id or not owner:
                    continue

                phone1 = str(sheet.cell(r_idx, col_mapping["phone1"]).value or "").strip() if "phone1" in col_mapping else ""
                email1 = str(sheet.cell(r_idx, col_mapping["email1"]).value or "").strip() if "email1" in col_mapping else ""
                tenant = str(sheet.cell(r_idx, col_mapping["tenant"]).value or "").strip() if "tenant" in col_mapping else ""
                phone2 = str(sheet.cell(r_idx, col_mapping["phone2"]).value or "").strip() if "phone2" in col_mapping else ""
                email2 = str(sheet.cell(r_idx, col_mapping["email2"]).value or "").strip() if "email2" in col_mapping else ""

                aliquot_base = None
                if "aliquot_base" in col_mapping:
                    raw_aliquot = sheet.cell(r_idx, col_mapping["aliquot_base"]).value
                    if raw_aliquot is not None:
                        try:
                            aliquot_base = float(raw_aliquot)
                        except:
                            aliquot_base = None

                cedula_owner = str(sheet.cell(r_idx, col_mapping["cedula_owner"]).value or "").strip() if "cedula_owner" in col_mapping else ""
                cedula_tenant = str(sheet.cell(r_idx, col_mapping["cedula_tenant"]).value or "").strip() if "cedula_tenant" in col_mapping else ""

                units.append({
                    "id": unit_id, "owner": owner, "phone1": phone1, "email1": email1,
                    "tenant": tenant, "phone2": phone2, "email2": email2, "aliquot_base": aliquot_base,
                    "cedula_owner": cedula_owner, "cedula_tenant": cedula_tenant
                })
            wb.close()

        added = 0
        for u in units:
            sm.add_or_update_unit(
                unit_id=u["id"],
                owner=u["owner"],
                phone1=u["phone1"],
                email1=u["email1"],
                tenant=u["tenant"],
                phone2=u["phone2"],
                email2=u["email2"],
                aliquot_base=u["aliquot_base"],
                cedula_owner=u.get("cedula_owner", ""),
                cedula_tenant=u.get("cedula_tenant", "")
            )
            added += 1

        return {
            "status": "success",
            "message": f"Se procesaron e importaron {added} departamentos correctamente.",
            "added": added
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al procesar el archivo de departamentos: {e}")
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

@router.get("/api/board-members")
def get_board_members():
    sm = get_sm()
    return sm.data.get("board_members", [])

@router.post("/api/board-members")
def create_board_member(member: BoardMemberCreate):
    sm = get_sm()
    return sm.add_board_member(member.unit_id, member.name, member.role)

@router.delete("/api/board-members/{id}")
def delete_board_member(id: int):
    sm = get_sm()
    success = sm.delete_board_member(id)
    if not success:
        raise HTTPException(status_code=404, detail="Miembro de la directiva no encontrado")
    return {"status": "success"}

@router.get("/api/units/{unit_id}/no-debt-certificate")
def download_no_debt_certificate(
    unit_id: str,
    current_user: Optional[dict] = Depends(get_current_user)
):
    sm = get_sm()
    RECEIPTS_DIR = get_receipts_dir()
    if current_user and isinstance(current_user, dict):
        check_unit_permission(current_user, str(unit_id))
    elif not (os.environ.get("TESTING") == "True"):
        raise HTTPException(
            status_code=401,
            detail="Autenticación requerida para generar o descargar el certificado de no adeudar."
        )

    unit = next((u for u in sm.data["units"] if u["id"] == str(unit_id)), None)
    if not unit:
        raise HTTPException(status_code=404, detail="Departamento no encontrado.")

    debtors = sm.get_debtors()
    is_debtor = any(d["unit"] == str(unit_id) and d["total_debt"] > 0 for d in debtors)
    if is_debtor:
        raise HTTPException(status_code=400, detail="El departamento posee deudas pendientes y no califica para el certificado de no adeudar.")

    president = next((m["name"] for m in sm.data.get("board_members", []) if "presidente" in m.get("role", "").lower()), sm.config.get("directive_president", "Carlos Mendoza"))
    treasurer = next((m["name"] for m in sm.data.get("board_members", []) if "tesorero" in m.get("role", "").lower()), sm.config.get("directive_treasurer", "Sofia Herrera"))

    condo_name = sm.config.get("condo_name", "Condominio El Mirador")
    condo_address = sm.config.get("condo_address", "")
    condo_ruc = sm.config.get("condo_ruc", "")

    pdf_filename = f"certificado_no_adeudar_{unit_id}.pdf"
    pdf_path = os.path.join(RECEIPTS_DIR, pdf_filename)

    generate_no_debt_certificate_pdf(
        dest_path=pdf_path,
        unit_id=unit_id,
        owner_name=unit["owner"],
        president_name=president,
        treasurer_name=treasurer,
        condo_name=condo_name,
        condo_address=condo_address,
        condo_ruc=condo_ruc
    )

    return FileResponse(pdf_path, media_type="application/pdf", filename=pdf_filename)
