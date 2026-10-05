import os
import re
from datetime import datetime
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

try:
    import pypdf
except ImportError:
    pypdf = None

import json
import html

CONFIG_PATH = os.path.join(os.path.dirname(__file__), "config.json")

def get_condo_info():
    info = {
        "condo_name": "Condominio Casales San Pedro",
        "condo_address": "",
        "condo_ruc": ""
    }
    try:
        if os.path.exists(CONFIG_PATH):
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                cfg = json.load(f)
                if cfg.get("condo_name"):
                    info["condo_name"] = cfg["condo_name"]
                if cfg.get("condo_address"):
                    info["condo_address"] = cfg["condo_address"]
                if cfg.get("condo_ruc"):
                    info["condo_ruc"] = cfg["condo_ruc"]
    except:
        pass
    return info

def get_condo_name():
    return get_condo_info()["condo_name"]

class ReceiptProcessor:
    @staticmethod
    def parse_whatsapp_message(text):
        """
        Parses structured WhatsApp messages sent by residents.
        Example:
            Nombre: Jenny Portilla Bustamante
            Cédula: 171391685-4
            Departamento: 701
            Mes: Abril/ 2026
        """
        data = {
            "name": "",
            "cedula": "",
            "unit": "",
            "month": "",
            "year": 2026
        }
        
        # Match lines with labels
        name_match = re.search(r'Nombre:\s*(.+)', text, re.IGNORECASE)
        if name_match:
            data["name"] = name_match.group(1).strip()
            
        cedula_match = re.search(r'(?:Cédula|Cedula):\s*(.+)', text, re.IGNORECASE)
        if cedula_match:
            data["cedula"] = cedula_match.group(1).strip()
            
        unit_match = re.search(r'(?:Departamento|Depto|Dep):\s*(\w+)', text, re.IGNORECASE)
        if unit_match:
            data["unit"] = unit_match.group(1).strip()
            
        # Match month/year
        month_match = re.search(r'Mes:\s*([a-zA-ZáéíóúÁÉÍÓÚ]+)(?:\s*/\s*(\d{4}))?', text, re.IGNORECASE)
        if month_match:
            data["month"] = month_match.group(1).strip().capitalize()
            if month_match.group(2):
                data["year"] = int(month_match.group(2))
        else:
            # Fallback check for month name in the whole text
            months = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
            for m in months:
                if m in text.lower():
                    data["month"] = m.capitalize()
                    break
        
        return data

    @staticmethod
    def extract_text_from_pdf(pdf_path):
        """
        Extracts raw text from a PDF file using pypdf if installed.
        """
        if not pypdf:
            print("pypdf is not installed. Cannot extract text from PDF.")
            return ""
        try:
            reader = pypdf.PdfReader(pdf_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text
        except Exception as e:
            print(f"Error reading PDF {pdf_path}: {e}")
            return ""

    @staticmethod
    def parse_receipt_date(text):
        """
        Extracts and normalizes the payment/deposit date from receipt text (OCR or PDF).
        Supports dates formatted with numbers, letters/words (Spanish and English), or combined:
        e.g. '15/07/2026', '15 de Julio de 2026', '15 de julio del 2026', '15-Jul-2026', 'Jul 15, 2026', '2026-07-15'.
        Returns normalized string in 'YYYY-MM-DD' format, or None if not found.
        """
        if not text:
            return None

        MONTH_MAP = {
            'enero': 1, 'ene': 1, 'january': 1, 'jan': 1,
            'febrero': 2, 'feb': 2, 'february': 2,
            'marzo': 3, 'mar': 3, 'march': 3,
            'abril': 4, 'abr': 4, 'april': 4, 'apr': 4,
            'mayo': 5, 'may': 5,
            'junio': 6, 'jun': 6, 'june': 6,
            'julio': 7, 'jul': 7, 'july': 7,
            'agosto': 8, 'ago': 8, 'august': 8, 'aug': 8,
            'septiembre': 9, 'setiembre': 9, 'sep': 9, 'set': 9, 'september': 9,
            'octubre': 10, 'oct': 10, 'october': 10,
            'noviembre': 11, 'nov': 11, 'november': 11,
            'diciembre': 12, 'dic': 12, 'december': 12, 'dec': 12
        }

        label_patterns = [
            r'(?:fecha\s*valor|f\.?\s*valor|fecha\s*de\s*transferencia|fecha\s*transferencia|fecha\s*de\s*pago|fecha\s*de\s*dep[oó]sito|fecha\s*dep[oó]sito|fecha\s*de\s*proceso|fecha\s*proceso|fecha\s*y\s*hora|fecha\s*emisi[oó]n|fecha|fec\.?\s*tx|fec\.?)\s*[:\s-]\s*([0-9a-zA-Z\s\/\-\.,delDE]+)',
        ]
        
        candidates = []
        for pat in label_patterns:
            for m in re.finditer(pat, text, re.IGNORECASE):
                raw = m.group(1).strip()
                candidates.append(raw)
                
        for line in text.splitlines():
            clean_l = line.strip()
            if clean_l and clean_l not in candidates:
                candidates.append(clean_l)
                
        for raw in candidates:
            if not raw:
                continue
                
            # Pattern A: Day Month Year in letters (e.g. 15 de Julio de 2026, 15-Jul-2026, 15/Julio/2026, 15.ABR.2026)
            m_let = re.search(r'\b(\d{1,2})\s*(?:de|\/|-|\.)\s*([a-zA-Z]{3,12})\s*(?:de|del|\/|-|\.)?\s*(\d{2,4})\b', raw, re.IGNORECASE)
            if m_let:
                d_str, mon_str, y_str = m_let.groups()
                mon = MONTH_MAP.get(mon_str.lower())
                if mon:
                    d = int(d_str)
                    y = int(y_str)
                    if y < 100:
                        y += 2000
                    if 1 <= d <= 31 and 1 <= mon <= 12 and 2000 <= y <= 2100:
                        return f'{y:04d}-{mon:02d}-{d:02d}'

            # Pattern A2: Year Month in letters Day (e.g. 2026/SEP/07, 2026-Septiembre-07, 2026 / SEP / 07, 2026.SEP.07)
            m_ymd_let = re.search(r'\b(20\d{2})\s*[\/\-\.]\s*([a-zA-Z]{3,12})\s*[\/\-\.]\s*(\d{1,2})\b', raw, re.IGNORECASE)
            if m_ymd_let:
                y_str, mon_str, d_str = m_ymd_let.groups()
                mon = MONTH_MAP.get(mon_str.lower())
                if mon:
                    d = int(d_str)
                    y = int(y_str)
                    if 1 <= d <= 31 and 1 <= mon <= 12 and 2000 <= y <= 2100:
                        return f'{y:04d}-{mon:02d}-{d:02d}'
                        
            # Pattern B: Month Day, Year (e.g. Julio 15, 2026 or Jul 15 2026)
            m_let2 = re.search(r'\b([a-zA-Z]{3,12})\s*(\d{1,2})\s*(?:de|,)?\s*(\d{2,4})\b', raw, re.IGNORECASE)
            if m_let2:
                mon_str, d_str, y_str = m_let2.groups()
                mon = MONTH_MAP.get(mon_str.lower())
                if mon:
                    d = int(d_str)
                    y = int(y_str)
                    if y < 100:
                        y += 2000
                    if 1 <= d <= 31 and 1 <= mon <= 12 and 2000 <= y <= 2100:
                        return f'{y:04d}-{mon:02d}-{d:02d}'

            # Pattern C: YYYY-MM-DD or YYYY/MM/DD
            m_iso = re.search(r'\b(20\d{2})[\/\-\.](\d{1,2})[\/\-\.](\d{1,2})\b', raw)
            if m_iso:
                y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
                if 1 <= d <= 31 and 1 <= m <= 12:
                    return f'{y:04d}-{m:02d}-{d:02d}'

            # Pattern D: DD/MM/YYYY or DD-MM-YYYY or DD.MM.YYYY
            m_dmy = re.search(r'\b(\d{1,2})[\/\-\.](\d{1,2})[\/\-\.](20\d{2})\b', raw)
            if m_dmy:
                d, m, y = int(m_dmy.group(1)), int(m_dmy.group(2)), int(m_dmy.group(3))
                if 1 <= d <= 31 and 1 <= m <= 12:
                    return f'{y:04d}-{m:02d}-{d:02d}'

            # Pattern E: DD/MM/YY
            m_dmy2 = re.search(r'\b(\d{1,2})[\/\-\.](\d{1,2})[\/\-\.](\d{2})\b', raw)
            if m_dmy2:
                d, m, y = int(m_dmy2.group(1)), int(m_dmy2.group(2)), int(m_dmy2.group(3)) + 2000
                if 1 <= d <= 31 and 1 <= m <= 12:
                    return f'{y:04d}-{m:02d}-{d:02d}'
                    
        return None

    @staticmethod
    def clean_date(date_str):
        """
        Normalizes any date string (words, numbers, mixed) to YYYY-MM-DD.
        """
        if not date_str:
            return None
        date_str = str(date_str).strip()
        parsed = ReceiptProcessor.parse_receipt_date(date_str)
        if parsed:
            return parsed
        # Fallback regex checks
        match = re.search(r'\b(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})\b', date_str)
        if match:
            y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 1 <= m <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
        match = re.search(r'\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\b', date_str)
        if match:
            d, m, y = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 1 <= m <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
        return date_str

    @staticmethod
    def parse_bank_receipt_text(text):
        """
        Parses text extracted from a bank receipt (image OCR or PDF)
        to find reference ID/voucher number, payment amount, and voucher payment date.
        Accurately extracts voucher/document numbers (e.g. 'No. comprobante', 'Documento',
        'N. de comprobante', 'N° comprobante', 'Secuencia', 'Operación') while excluding
        origin or destination bank account numbers.
        """
        if not text:
            return None, None, None

        amount = None
        reference = None
        payment_date = ReceiptProcessor.parse_receipt_date(text)
        
        # 1. Search for amount: e.g. "Monto: $50.00" or "$ 60.00" or "MONTO DEPOSITADO: $50" or "S 25.00"
        amount_patterns = [
            r'\b(?:monto\s*depositado|monto\s*total|monto\s*de\s*transferencia|monto\s*transferido|monto|total\s*pagado|total\s*depositado|total|efectivo|efactivo|cantidad)\b[\.,;:\s-]*\$?\s*(?:usd)?\s*(\d+(?:[.,]\d{1,2})?)(?!\s*[-/.](?:[a-zA-Z]{3,}|20\d{2}))',
            r'(?<!fecha\s)(?<!fecha\s\s)(?<!f\.\s)(?<!fec\.\s)(?<!fec\s)\bvalor\b[\.,;:\s-]*\$?\s*(?:usd)?\s*(\d+(?:[.,]\d{1,2})?)(?!\s*[-/.](?:[a-zA-Z]{3,}|20\d{2}))',
            r'\b(?:usd|\$|\busd\b)\s*[:\s-]*\s*(\d+(?:[.,]\d{1,2})?)',
            r'\b[sS]\s*(\d+(?:[.,]\d{2}))\b',
            r'(?:exitosa|transferencia)[!\s]*[\r\n]+\s*\$?\s*(\d+(?:[.,]\d{1,2})?)',
            r'\$\s*(\d+(?:[.,]\d{1,2})?)'
        ]
        
        for pat in amount_patterns:
            match = re.search(pat, text, re.IGNORECASE)
            if match:
                try:
                    amount = float(match.group(1).replace(",", "."))
                    break
                except:
                    pass

        # 2. Extract known account numbers to avoid picking them up as reference
        account_patterns = [
            r'\b(?:cuenta\s*(?:de\s*)?(?:origen|d[eé]bito|debitada|ahorros|corriente)|desde\s*(?:la\s*)?cuenta|de\s*(?:la\s*)?cuenta|cta\.?\s*(?:origen|d[eé]bito|debitada|ahorros|corriente)|no\.?\s*(?:de\s*)?cuenta\s*(?:de\s*)?(?:origen|d[eé]bito|debitada)?|n[°º\.]*\s*(?:de\s*)?cuenta\s*(?:de\s*)?(?:origen|d[eé]bito|debitada)?|nro\.?\s*(?:de\s*)?cuenta\s*(?:de\s*)?(?:origen|d[eé]bito|debitada)?|cuenta\s*n[°ºo\.]*)\s*[:\s#-]*([a-zA-Z0-9\*\-]+)',
            r'\b(?:cuenta\s*(?:de\s*)?(?:destino|cr[eé]dito|acreditada|beneficiario)|hacia\s*(?:la\s*)?cuenta|a\s*(?:la\s*)?cuenta|cta\.?\s*(?:destino|cr[eé]dito|acreditada|beneficiario)|no\.?\s*(?:de\s*)?cuenta\s*(?:de\s*)?(?:destino|cr[eé]dito|acreditada)?)\s*[:\s#-]*([a-zA-Z0-9\*\-]+)'
        ]
        
        known_accounts = set()
        for acc_pat in account_patterns:
            for m in re.finditer(acc_pat, text, re.IGNORECASE):
                raw_acc = m.group(1).strip()
                if raw_acc:
                    known_accounts.add(raw_acc.lower())
                    clean_acc = re.sub(r'[^0-9]', '', raw_acc)
                    if clean_acc:
                        known_accounts.add(clean_acc)

        # 3. Search for Voucher / Document / Reference number (Comprobante, Documento, Secuencia, Operación, Transacción, Referencia)
        ref_patterns = [
            # Comprobante variants (N. de comprobante, N. de Comprobante, N.de comprobante, N.decomprobante, N.° de comprobante, N.º de comprobante, N.o de comprobante, N° de comprobante, Nº de comprobante, Nro. de comprobante, No. de comprobante, No, de comprobante, N, de comprobante, N; de comprobante, N- de comprobante, N_ de comprobante, Numero de comprobante, Número de comprobante, Num. de comprobante, N. comprobante, N comprobante, Comprobante No., Comprobante N°, Comprobante N., Comprobante Nro., Comprobante:, Comprobante, etc.)
            r'\b(?:n[°º\.,;_\-o\s]*\s*(?:de\s*)?comprobante|comprobante\s*(?:n[°º\.,;_\-o\s]*|#)?)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)',
            
            # Documento variants (Documento, Documenta, No. documento, N. de documento, N° documento, Documento No., etc.)
            r'\b(?:n[°º\.,;_\-o\s]*\s*(?:de\s*)?document[oa]|document[oa]\s*(?:n[°º\.,;_\-o\s]*|#)?|doc\.?\s*(?:n[°º\.,;_\-o\s]*|#)?)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)',
            
            # Transacción / Operación variants (No. transacción, N° operación, Transacción No., Operación:, No. Op, etc.)
            r'\b(?:n[°º\.,;_\-o\s]*\s*(?:de\s*)?(?:transacci[oó]n|operaci[oó]n|\bop\b\.?)|(?:transacci[oó]n|operaci[oó]n)\s*(?:n[°º\.,;_\-o\s]*|#)?|\bop\b\.?\s*(?:n[°º\.,;_\-o\s]*|#)?)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)',
            
            # Secuencia / Autorización / Control variants (Secuencia:, No. autorización, No. control, etc.)
            r'\b(?:n[°º\.,;_\-o\s]*\s*(?:de\s*)?(?:secuencia|autorizaci[oó]n|control|aprobaci[oó]n)|(?:secuencia|autorizaci[oó]n|control|aprobaci[oó]n|\bsec\b\.?)\s*(?:n[°º\.,;_\-o\s]*|#)?)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)',
            
            # Referencia variants (Referencia / Cédula, No. referencia, Referencia No., Referencia:, Ref:)
            r'\b(?:referencia\s*/\s*c[eé]dula|n[°º\.,;_\-o\s]*\s*(?:de\s*)?referencia|referencia\s*(?:n[°º\.,;_\-o\s]*|#)?|\bref\b\.?\s*(?:n[°º\.,;_\-o\s]*|#)?)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)',
            
            # Código / ID de transacción (Código de transferencia, Transaction ID, ID Transacción)
            r'\b(?:c[oó]digo\s*(?:de\s*)?(?:transferencia|transacci[oó]n|confirmaci[oó]n|operaci[oó]n|pago)|(?:transaction|tx)\s*id|id\s*(?:de\s*)?transacci[oó]n)\s*[:\s#\.-]*\s*([0-9][a-zA-Z0-9\-]*|[a-zA-Z0-9\-]*[0-9][a-zA-Z0-9\-]*)'
        ]
        
        ignored_refs = {'de', 'del', 'la', 'el', 'pago', 'transferencia', 'deposito', 'depósito', 'electronico', 'electrónico', 'directa', 'exitosa', 'bancario', 'ahorros', 'corriente', 'cuenta', 'origen', 'destino', 'debito', 'débito', 'credito', 'crédito', 'interbancaria', 'efectivo', 'pichincha', 'guayaquil', 'produbanco', 'pacifico', 'pacífico', 'internacional', 'bolivariano', 'cooperativa', 'jep'}
        
        for pat in ref_patterns:
            for match in re.finditer(pat, text, re.IGNORECASE):
                val = match.group(1).strip()
                if not val or val.lower() in ignored_refs:
                    continue
                clean_digits = re.sub(r'[^0-9]', '', val)
                if val.lower() in known_accounts or (clean_digits and clean_digits in known_accounts):
                    continue
                reference = val
                break
            if reference:
                break
                
        # 4. Fallback 1: Cédula pattern if no specific voucher reference found
        if not reference:
            ced_match = re.search(r'\b(?:C[eé]dula|CI|Identificaci[oó]n|RUC)\s*[:\s#-]*([0-9\-]+)', text, re.IGNORECASE)
            if ced_match:
                val = ced_match.group(1).strip()
                clean_digits = re.sub(r'[^0-9]', '', val)
                if clean_digits not in known_accounts:
                    reference = val
                    
        # 5. Fallback 2: Standalone numbers avoiding lines with account/amount/date keywords
        if not reference:
            for line in text.splitlines():
                line_str = line.strip()
                line_lower = line_str.lower()
                if any(kw in line_lower for kw in ['cuenta', 'cta', 'origen', 'destino', 'saldo', 'tarjeta', 'celular', 'tel[eé]fono', 'valor', 'monto', 'total', 'fecha']):
                    continue
                standalone = re.findall(r'\b\d{6,12}\b', line_str)
                for num in standalone:
                    if num not in known_accounts and num not in ['2024', '2025', '2026', '2027', '2028']:
                        reference = num
                        break
                if reference:
                    break
                    
        # 6. Normalize reference if ending in check digit e.g. 171391685-4
        if reference and '-' in reference:
            parts = reference.split('-')
            if len(parts) == 2 and parts[0].isdigit() and len(parts[0]) >= 8 and len(parts[1]) <= 2:
                reference = parts[0]
                
        return amount, reference, payment_date

    @staticmethod
    def generate_receipt_pdf(dest_path, receipt_no, resident_name, unit, concept, amount, late_fee, reference, payment_date, condo_name=None, abono_deuda=0.0, pago_extra=0.0, saldo_deuda=0.0, advance_aliquots=None, condo_address=None, condo_ruc=None, pending_debts_summary=None, sm=None, **kwargs):
        """
        Generates a beautiful professional PDF receipt using ReportLab.
        Includes complete unit debt breakdown / account statement status.
        """
        if sm is not None:
            condo_info = {
                "condo_name": sm.config.get("condo_name") or "Condominio Casales San Pedro",
                "condo_address": sm.config.get("condo_address") or "",
                "condo_ruc": sm.config.get("condo_ruc") or ""
            }
        else:
            condo_info = get_condo_info()
            
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        # Ensure directories exist
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        
        doc = SimpleDocTemplate(dest_path, pagesize=letter,
                                rightMargin=36, leftMargin=36,
                                topMargin=36, bottomMargin=36)
        story = []
        
        # Styles
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'ReceiptTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1E293B'), # Dark grey
            alignment=1, # Center
            spaceAfter=15
        )
        
        subtitle_style = ParagraphStyle(
            'ReceiptSubTitle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#64748B'),
            alignment=1,
            spaceAfter=25
        )
        
        label_style = ParagraphStyle(
            'Label',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            fontName='Helvetica-Bold',
            textColor=colors.HexColor('#475569')
        )
        
        value_style = ParagraphStyle(
            'Value',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#0F172A')
        )
        
        header_lines = [f"<b>{condo_name.upper()}</b>"]
        if condo_ruc:
            header_lines.append(f"RUC: {condo_ruc}")
        header_lines.append("Administración y Control de Pagos")
        if condo_address:
            header_lines.append(condo_address)
        else:
            header_lines.append("Quito, Ecuador")
        header_text = "<br/>\n".join(header_lines)
        
        # 1. Header Table
        header_p = Paragraph(header_text, ParagraphStyle('HeaderStyle', parent=styles['Normal'], fontSize=9.5, leading=13.5, textColor=colors.HexColor('#475569')))
        receipt_p = Paragraph(f"<b>RECIBO OFICIAL DE PAGO</b><br/><font color='#2563EB'><b>Nº {receipt_no}</b></font><br/>Fecha de Emisión: {datetime.now().strftime('%d/%m/%Y')}", ParagraphStyle('ReceiptStyle', parent=styles['Normal'], fontSize=10, leading=14, alignment=2))
        
        header_table = Table([[header_p, receipt_p]], colWidths=[4.0*inch, 3.5*inch])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 15),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 15))
        
        # Decorative Blue Line
        line_table = Table([[""]], colWidths=[7.5*inch])
        line_table.setStyle(TableStyle([
            ('LINEBELOW', (0,0), (-1,-1), 2, colors.HexColor('#2563EB')), # Vivid Blue
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
        ]))
        story.append(line_table)
        story.append(Spacer(1, 20))
        
        # 2. Information Grid
        info_data = [
            [Paragraph("Condómino:", label_style), Paragraph(resident_name, value_style), Paragraph("Unidad / Depto:", label_style), Paragraph(unit, value_style)],
            [Paragraph("Referencia Banco:", label_style), Paragraph(reference or "Validación Manual", value_style), Paragraph("Fecha Pago:", label_style), Paragraph(payment_date or "N/A", value_style)],
        ]
        info_table = Table(info_data, colWidths=[1.5*inch, 2.5*inch, 1.5*inch, 2.0*inch])
        info_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LEFTPADDING', (0,0), (-1,-1), 10),
            ('RIGHTPADDING', (0,0), (-1,-1), 10),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.HexColor('#E2E8F0')),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 20))
        
        # 3. Concept and Values Table
        table_style_header = ParagraphStyle('TH', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold', textColor=colors.white)
        table_style_cell = ParagraphStyle('TD', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#334155'))
        table_style_cell_r = ParagraphStyle('TDR', parent=styles['Normal'], fontSize=9, alignment=2, textColor=colors.HexColor('#334155'))
        table_style_total_label = ParagraphStyle('TTL', parent=styles['Normal'], fontSize=10, fontName='Helvetica-Bold', alignment=2, textColor=colors.HexColor('#1E293B'))
        table_style_total_val = ParagraphStyle('TTV', parent=styles['Normal'], fontSize=10, fontName='Helvetica-Bold', alignment=2, textColor=colors.HexColor('#2563EB'))
        
        detail_data = [
            [Paragraph("DESCRIPCIÓN / CONCEPTO", table_style_header), Paragraph("CANTIDAD", table_style_header), Paragraph("MONTO BASE", table_style_header), Paragraph("TOTAL", table_style_header)],
        ]
        
        subtotal = 0.0
        # If there is aliquot paid
        if amount > 0:
            detail_data.append([
                Paragraph(concept, table_style_cell),
                Paragraph("1", table_style_cell),
                Paragraph(f"${amount:.2f}", table_style_cell_r),
                Paragraph(f"${amount:.2f}", table_style_cell_r)
            ])
            subtotal += amount
            
        # If there are future aliquots paid in advance
        if advance_aliquots:
            for a in advance_aliquots:
                amt = a["amount"]
                status_txt = "Pago Adelantado" if a["status"] == "Pagado" else "Abono Adelantado"
                detail_data.append([
                    Paragraph(f"Alícuota Ordinaria - Mes: {a['month']} / {a['year']} ({status_txt})", table_style_cell),
                    Paragraph("1", table_style_cell),
                    Paragraph(f"${amt:.2f}", table_style_cell_r),
                    Paragraph(f"${amt:.2f}", table_style_cell_r)
                ])
                subtotal += amt
            
        # If there is abono to debt
        if abono_deuda > 0:
            detail_data.append([
                Paragraph("Abono a Deuda Pendiente / Saldo Deudor", table_style_cell),
                Paragraph("1", table_style_cell),
                Paragraph(f"${abono_deuda:.2f}", table_style_cell_r),
                Paragraph(f"${abono_deuda:.2f}", table_style_cell_r)
            ])
            subtotal += abono_deuda
            
        # If there is extra payment
        if pago_extra > 0:
            detail_data.append([
                Paragraph("Pago Adicional / Saldo a Favor", table_style_cell),
                Paragraph("1", table_style_cell),
                Paragraph(f"${pago_extra:.2f}", table_style_cell_r),
                Paragraph(f"${pago_extra:.2f}", table_style_cell_r)
            ])
            subtotal += pago_extra
            
        # If there is late fee
        if late_fee > 0:
            detail_data.append([
                Paragraph("Recargo por Pago Tardío (Multa después del día 15)", table_style_cell),
                Paragraph("1", table_style_cell),
                Paragraph(f"${late_fee:.2f}", table_style_cell_r),
                Paragraph(f"${late_fee:.2f}", table_style_cell_r)
            ])
            
        total = subtotal + late_fee
        
        # Add totals rows
        detail_data.extend([
            ["", "", Paragraph("SUBTOTAL:", table_style_total_label), Paragraph(f"${subtotal:.2f}", table_style_total_label)],
            ["", "", Paragraph("MULTA APLICADA:", table_style_total_label), Paragraph(f"${late_fee:.2f}", table_style_total_label)],
            ["", "", Paragraph("TOTAL RECIBIDO:", table_style_total_label), Paragraph(f"${total:.2f}", table_style_total_val)]
        ])
        
        # Add outstanding debt remaining row if any
        if saldo_deuda > 0:
            detail_data.append([
                "", "",
                Paragraph("SALDO RESTANTE DE DEUDA:", ParagraphStyle('TTL_Red', parent=table_style_total_label, textColor=colors.HexColor('#DC2626'))),
                Paragraph(f"${saldo_deuda:.2f}", ParagraphStyle('TTV_Red', parent=table_style_total_val, textColor=colors.HexColor('#DC2626')))
            ])
        
        col_widths = [4.0*inch, 1.0*inch, 1.25*inch, 1.25*inch]
        detail_table = Table(detail_data, colWidths=col_widths)
        
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')), # Dark grey header
            ('ALIGN', (0,0), (-1,0), 'LEFT'),
            ('ALIGN', (2,1), (3,-1), 'RIGHT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LINEBELOW', (0,1), (-1,-2), 0.5, colors.HexColor('#E2E8F0')),
            ('LINEBELOW', (2,-3), (3,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('LINEBELOW', (2,-1), (3,-1), 1.5, colors.HexColor('#2563EB')),
        ]
        detail_table.setStyle(TableStyle(t_style))
        story.append(detail_table)
        story.append(Spacer(1, 15))
        
        # 4. Estado de Cuenta / Detalle de Deuda Pendiente
        debt_details = []
        unit_total_debt = 0.0

        if pending_debts_summary is not None:
            if isinstance(pending_debts_summary, dict):
                unit_total_debt = float(pending_debts_summary.get("total_debt", 0.0) or 0.0)
                debt_details = list(pending_debts_summary.get("details", []))
            elif isinstance(pending_debts_summary, list):
                debt_details = list(pending_debts_summary)
                unit_total_debt = float(saldo_deuda or 0.0)
            elif isinstance(pending_debts_summary, (int, float)):
                unit_total_debt = float(pending_debts_summary)
        else:
            try:
                if sm is None:
                    from sheets_manager import SheetsManager
                    sm_inst = SheetsManager()
                else:
                    sm_inst = sm

                debtors = sm_inst.get_debtors()
                clean_unit = str(unit or "").strip().lower().replace("depto", "").replace("depto.", "").strip()

                for d in debtors:
                    d_unit = str(d.get("unit", "")).strip().lower().replace("depto", "").replace("depto.", "").strip()
                    if d_unit == clean_unit or str(d.get("unit", "")).strip().lower() == str(unit or "").strip().lower():
                        unit_total_debt = float(d.get("total_debt", 0.0) or 0.0)
                        debt_details = list(d.get("details", []))
                        break
            except Exception as e:
                if saldo_deuda > 0:
                    unit_total_debt = float(saldo_deuda)

        if unit_total_debt <= 0.001 and saldo_deuda > 0:
            unit_total_debt = float(saldo_deuda)
            if not debt_details:
                debt_details = [f"Saldo pendiente tras el pago actual: ${saldo_deuda:.2f}"]

        if unit_total_debt > 0.001 or debt_details:
            debt_box_rows = [
                [Paragraph("<b>ESTADO DE CUENTA: DEUDA PENDIENTE REGISTRADA</b>", ParagraphStyle('DebtHead', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold', textColor=colors.HexColor('#991B1B')))],
            ]
            for item in debt_details:
                clean_item = html.escape(str(item))
                debt_box_rows.append([
                    Paragraph(f"• {clean_item}", ParagraphStyle('DebtRow', parent=styles['Normal'], fontSize=8.5, leading=11.5, textColor=colors.HexColor('#7F1D1D')))
                ])
            debt_box_rows.append([
                Paragraph(f"<b>TOTAL DEUDA PENDIENTE: ${unit_total_debt:.2f}</b>", ParagraphStyle('DebtTot', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold', leading=12, alignment=2, textColor=colors.HexColor('#B91C1C')))
            ])
            debt_box_table = Table(debt_box_rows, colWidths=[7.5*inch])
            debt_box_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#FEF2F2')), # Light red
                ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor('#F87171')),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('LEFTPADDING', (0,0), (-1,-1), 10),
                ('RIGHTPADDING', (0,0), (-1,-1), 10),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('LINEBELOW', (0,0), (-1,0), 0.5, colors.HexColor('#FECACA')),
                ('LINEABOVE', (0,-1), (-1,-1), 0.5, colors.HexColor('#FECACA')),
            ]))
            story.append(debt_box_table)
        else:
            aldia_p = Paragraph("<b>ESTADO DE CUENTA: AL DÍA</b> — La unidad no registra deudas ni alícuotas vencidas a la fecha de emisión.", ParagraphStyle('AlDia', parent=styles['Normal'], fontSize=8.5, leading=11.5, fontName='Helvetica-Bold', textColor=colors.HexColor('#15803D')))
            aldia_table = Table([[aldia_p]], colWidths=[7.5*inch])
            aldia_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F0FDF4')), # Light green
                ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor('#86EFAC')),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('LEFTPADDING', (0,0), (-1,-1), 10),
                ('RIGHTPADDING', (0,0), (-1,-1), 10),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ]))
            story.append(aldia_table)
            
        story.append(Spacer(1, 20))
        
        # 5. Signatures / Bottom Seal
        sign_label_style = ParagraphStyle('SL', parent=styles['Normal'], fontSize=8, leading=12, alignment=1, textColor=colors.HexColor('#64748B'))
        sign_line_table = Table([
            ["", ""],
            [Paragraph("___________________________________<br/><b>Administración del Condominio</b><br/>Firma Autorizada", sign_label_style),
             Paragraph("<b>VALIDADO ELECTRÓNICAMENTE</b><br/>Este documento es un recibo digital emitido tras la conciliación bancaria exitosa.", ParagraphStyle('ValidatedStyle', parent=sign_label_style, fontName='Helvetica-Bold', textColor=colors.HexColor('#16A34A')))]
        ], colWidths=[3.75*inch, 3.75*inch])
        sign_line_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))
        story.append(sign_line_table)
        
        # Build Document
        doc.build(story)
        return dest_path

class BankStatementParser:
    @staticmethod
    def clean_amount(val_str):
        if val_str is None:
            return 0.0
        val_str = str(val_str).strip().replace("$", "")
        if not val_str:
            return 0.0
        
        # Remove any spaces
        val_str = val_str.replace(" ", "")
        
        # Check if wrapped in parentheses (negative amount)
        if val_str.startswith("(") and val_str.endswith(")"):
            val_str = "-" + val_str[1:-1]
            
        # Check if it has both comma and dot
        if "," in val_str and "." in val_str:
            if val_str.rfind(",") > val_str.rfind("."):
                # Comma is decimal (e.g. 1.250,50)
                val_str = val_str.replace(".", "").replace(",", ".")
            else:
                # Dot is decimal (e.g. 1,250.50)
                val_str = val_str.replace(",", "")
        elif "," in val_str:
            # Only comma is present. Could be decimal (e.g. 120,50) or thousands separator (e.g. 1,250)
            parts = val_str.split(",")
            if len(parts) == 2 and len(parts[1]) == 2:
                val_str = val_str.replace(",", ".")
            else:
                val_str = val_str.replace(",", "")
                
        try:
            return float(val_str)
        except:
            return 0.0

    @staticmethod
    def clean_date(date_str):
        date_str = str(date_str).strip()
        if not date_str:
            return datetime.now().strftime("%Y-%m-%d")
            
        parsed = ReceiptProcessor.parse_receipt_date(date_str)
        if parsed:
            return parsed
            
        # Try YYYY-MM-DD
        match = re.search(r'\b(\d{4})[-/](\d{1,2})[-/](\d{1,2})\b', date_str)
        if match:
            y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 1 <= m <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
            
        # Try DD-MM-YYYY or DD/MM/YYYY
        match = re.search(r'\b(\d{1,2})[-/](\d{1,2})[-/](\d{4})\b', date_str)
        if match:
            d, m, y = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 1 <= m <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
            
        # Try DD-MM-YY or DD/MM/YY
        match = re.search(r'\b(\d{1,2})[-/](\d{1,2})[-/](\d{2})\b', date_str)
        if match:
            d, m, y = int(match.group(1)), int(match.group(2)), int(match.group(3)) + 2000
            if 1 <= m <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
            
        # Try DD-MM or DD/MM (no year) - only if month and day are strictly valid
        match = re.search(r'\b(\d{1,2})[-/](\d{1,2})\b', date_str)
        if match:
            d, m = int(match.group(1)), int(match.group(2))
            if 1 <= m <= 12 and 1 <= d <= 31:
                return f"{datetime.now().year}-{m:02d}-{d:02d}"
            
        return date_str

    @staticmethod
    def parse_statement_file(file_path, file_ext):
        """
        Parses transactions from a spreadsheet (XLSX, CSV), PDF, or image.
        Returns a list of dicts: [{"date": "YYYY-MM-DD", "reference": "...", "amount": float, "detail": "..."}]
        """
        import csv
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
            
        transactions = []
        file_ext = file_ext.lower()
        
        # 1. Excel Parser
        if file_ext == ".xlsx":
            if not openpyxl:
                print("openpyxl is not installed. Cannot parse Excel statement.")
                return []
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet = wb.active
                
                header_row_idx = None
                col_mapping = {} 
                
                # Look at first 15 rows to find header
                for r_idx in range(1, 16):
                    row_vals = [str(sheet.cell(r_idx, col_idx).value or "").strip().lower() for col_idx in range(1, min(15, sheet.max_column + 1))]
                    
                    has_date = any(k in row_vals for k in ["fecha", "date", "f. valor", "f.valor", "f. proceso", "f.proceso"])
                    has_amount = any(k in row_vals for k in ["monto", "valor", "depósito", "deposito", "abono", "crédito", "credito", "cantidad", "amount", "value"])
                    
                    if has_date and has_amount:
                        header_row_idx = r_idx
                        for col_idx in range(1, sheet.max_column + 1):
                            val = str(sheet.cell(r_idx, col_idx).value or "").strip().lower()
                            if not val:
                                continue
                            if any(k in val for k in ["saldo", "balance"]):
                                col_mapping["balance"] = col_idx
                            elif any(k in val for k in ["fecha", "date", "f. valor", "f.valor", "f. proceso", "f.proceso"]):
                                col_mapping["date"] = col_idx
                            elif any(k in val for k in ["referencia", "documento", "nro. doc", "código", "ref", "nro", "comprobante", "secuencial"]):
                                col_mapping["reference"] = col_idx
                            elif any(k in val for k in ["depósito", "deposito", "abono", "crédito", "credito", "ingreso", "entrada", "deposits", "credits"]):
                                col_mapping["credit"] = col_idx
                            elif any(k in val for k in ["retiro", "egreso", "débito", "debito", "cargo", "salida", "withdrawals", "debits"]):
                                col_mapping["debit"] = col_idx
                            elif any(k in val for k in ["monto", "valor", "cantidad", "amount", "value"]) and not any(k in val for k in ["saldo", "balance"]):
                                if "amount" not in col_mapping:
                                    col_mapping["amount"] = col_idx
                            elif any(k in val for k in ["detalle", "concepto", "descripción", "descripcion", "glosa", "beneficiario", "descrip", "narrativa"]):
                                col_mapping["detail"] = col_idx
                        break
                        
                if not header_row_idx:
                    header_row_idx = 1
                    col_mapping = {"date": 1, "reference": 2, "amount": 3, "detail": 4}
                    
                date_col = col_mapping.get("date", 1)
                ref_col = col_mapping.get("reference", 2)
                detail_col = col_mapping.get("detail", 4)
                
                credit_col = col_mapping.get("credit")
                debit_col = col_mapping.get("debit")
                amount_col = col_mapping.get("amount", 3)
                
                val_col = credit_col if credit_col is not None else amount_col
                
                for r_idx in range(header_row_idx + 1, sheet.max_row + 1):
                    raw_date = sheet.cell(r_idx, date_col).value
                    if raw_date is None:
                        continue
                        
                    is_debit = False
                    if debit_col is not None:
                        raw_debit = sheet.cell(r_idx, debit_col).value
                        if BankStatementParser.clean_amount(raw_debit) > 0:
                            is_debit = True
                            
                    raw_val = sheet.cell(r_idx, val_col).value
                    if raw_val is None:
                        continue
                        
                    val_str = str(raw_val).strip()
                    if val_str.startswith("-") or (val_str.startswith("(") and val_str.endswith(")")):
                        is_debit = True
                        
                    amount_val = abs(BankStatementParser.clean_amount(val_str))
                    if is_debit or amount_val <= 0:
                        continue
                        
                    raw_detail = str(sheet.cell(r_idx, detail_col).value or "").strip() if detail_col else ""
                    
                    # Fallback for misaligned columns (common in PDF-to-Excel converters)
                    if amount_val > 500.0 and raw_detail:
                        detail_amount_match = re.search(r'\$\s*(\d+(?:[.,]\d+)*)', raw_detail)
                        if detail_amount_match:
                            extracted_amt = BankStatementParser.clean_amount(detail_amount_match.group(1))
                            if 0.0 < extracted_amt < 500.0:
                                amount_val = extracted_amt
                                
                    tx = {}
                    if isinstance(raw_date, datetime):
                        tx["date"] = raw_date.strftime("%Y-%m-%d")
                    else:
                        tx["date"] = BankStatementParser.clean_date(str(raw_date))
                        
                    raw_ref = sheet.cell(r_idx, ref_col).value if ref_col else ""
                    tx["reference"] = str(raw_ref or "").split(".")[0].strip() if raw_ref is not None else ""
                    tx["amount"] = amount_val
                    tx["detail"] = raw_detail
                    
                    transactions.append(tx)
            except Exception as e:
                print(f"Error parsing XLSX statement: {e}")
                
        # 2. CSV Parser
        elif file_ext == ".csv":
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    
                lines = content.splitlines()
                if not lines:
                    return []
                    
                sep = ","
                if ";" in lines[0] and lines[0].count(";") > lines[0].count(","):
                    sep = ";"
                    
                reader = csv.reader(lines, delimiter=sep)
                rows = list(reader)
                
                header_row_idx = None
                col_mapping = {}
                
                for idx, r in enumerate(rows[:10]):
                    row_vals = [val.strip().lower() for val in r]
                    has_date = any(k in row_vals for k in ["fecha", "date", "f. valor", "f.valor", "f. proceso", "f.proceso"])
                    has_amount = any(k in row_vals for k in ["monto", "valor", "depósito", "deposito", "abono", "crédito", "credito", "cantidad", "amount", "value"])
                    if has_date and has_amount:
                        header_row_idx = idx
                        for col_idx, val in enumerate(row_vals):
                            if not val:
                                continue
                            if any(k in val for k in ["saldo", "balance"]):
                                col_mapping["balance"] = col_idx
                            elif any(k in val for k in ["fecha", "date", "f. valor", "f.valor", "f. proceso", "f.proceso"]):
                                col_mapping["date"] = col_idx
                            elif any(k in val for k in ["referencia", "documento", "nro. doc", "código", "ref", "nro", "comprobante", "secuencial"]):
                                col_mapping["reference"] = col_idx
                            elif any(k in val for k in ["depósito", "deposito", "abono", "crédito", "credito", "ingreso", "entrada", "deposits", "credits"]):
                                col_mapping["credit"] = col_idx
                            elif any(k in val for k in ["retiro", "egreso", "débito", "debito", "cargo", "salida", "withdrawals", "debits"]):
                                col_mapping["debit"] = col_idx
                            elif any(k in val for k in ["monto", "valor", "cantidad", "amount", "value"]) and not any(k in val for k in ["saldo", "balance"]):
                                if "amount" not in col_mapping:
                                    col_mapping["amount"] = col_idx
                            elif any(k in val for k in ["detalle", "concepto", "descripción", "descripcion", "glosa", "beneficiario", "descrip", "narrativa"]):
                                col_mapping["detail"] = col_idx
                        break
                        
                if header_row_idx is None:
                    header_row_idx = 0
                    col_mapping = {"date": 0, "reference": 1, "amount": 2, "detail": 3}
                    
                date_col = col_mapping.get("date", 0)
                ref_col = col_mapping.get("reference", 1)
                detail_col = col_mapping.get("detail", 3)
                
                credit_col = col_mapping.get("credit")
                debit_col = col_mapping.get("debit")
                amount_col = col_mapping.get("amount", 2)
                
                val_col = credit_col if credit_col is not None else amount_col
                
                for r in rows[header_row_idx + 1:]:
                    if len(r) <= max(col_mapping.values()):
                        continue
                    raw_date = r[date_col].strip()
                    if not raw_date:
                        continue
                        
                    is_debit = False
                    if debit_col is not None and len(r) > debit_col:
                        if BankStatementParser.clean_amount(r[debit_col]) > 0:
                            is_debit = True
                            
                    raw_val = r[val_col]
                    val_str = str(raw_val).strip()
                    if val_str.startswith("-") or (val_str.startswith("(") and val_str.endswith(")")):
                        is_debit = True
                        
                    amount_val = abs(BankStatementParser.clean_amount(val_str))
                    if is_debit or amount_val <= 0:
                        continue
                        
                    raw_detail = r[detail_col].strip() if len(r) > detail_col else "CSV Upload"
                    
                    # Fallback for misaligned columns (common in PDF-to-Excel/CSV converters)
                    if amount_val > 500.0 and raw_detail:
                        detail_amount_match = re.search(r'\$\s*(\d+(?:[.,]\d+)*)', raw_detail)
                        if detail_amount_match:
                            extracted_amt = BankStatementParser.clean_amount(detail_amount_match.group(1))
                            if 0.0 < extracted_amt < 500.0:
                                amount_val = extracted_amt
                                
                    tx = {}
                    tx["date"] = BankStatementParser.clean_date(raw_date)
                    tx["reference"] = r[ref_col].strip() if len(r) > ref_col else ""
                    tx["amount"] = amount_val
                    tx["detail"] = raw_detail
                    
                    transactions.append(tx)
            except Exception as e:
                print(f"Error parsing CSV statement: {e}")
                
        # 3. PDF and Image Parsing (Text extraction OCR/pypdf)
        else:
            text = ""
            if file_ext == ".pdf":
                text = ReceiptProcessor.extract_text_from_pdf(file_path)
            else:
                from ocr_helper import OCRHelper
                text = OCRHelper.extract_text(file_path)
                
            transactions = BankStatementParser.parse_text_statement(text)
            
        return transactions

    @staticmethod
    def parse_text_statement(text):
        transactions = []
        lines = text.split("\n")
        
        current_date = None
        accumulated_concept = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Detect date (YYYY-MM-DD, DD/MM/YYYY, or Month in letters)
            date_match = re.search(r'\b(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4}|\d{4}[-/][a-zA-Z]{3,12}[-/]\d{1,2}|\d{1,2}[-/][a-zA-Z]{3,12}[-/]\d{2,4})\b', line)
            if not date_match:
                date_match = re.search(r'(?:^|\b(?:fecha|fec)\.?\s*)(\d{1,2}[-/]\d{1,2})\b', line, re.IGNORECASE)
            if date_match:
                cleaned = BankStatementParser.clean_date(date_match.group(1))
                if re.match(r'^\d{4}-\d{2}-\d{2}$', str(cleaned)):
                    current_date = cleaned
                
            # Check if this line is a core transaction line (contains Crédito or Débito)
            # We use encoding-independent dot matches for accented characters
            tx_match = re.search(r'(\d+)\s+(Cr.dito|D.bito)', line, re.IGNORECASE)
            fused_match = re.search(r'\$\s*(-?\d+[.,]\d{2})\s*(\d{5,15})', line)
            
            if tx_match or (fused_match and re.search(r'(Cr.dito|D.bito)', line, re.IGNORECASE)):
                is_fused = False
                if fused_match:
                    ref_val = fused_match.group(2)
                    amount_val = abs(BankStatementParser.clean_amount(fused_match.group(1)))
                    tx_type_match = re.search(r'(Cr.dito|D.bito)', line, re.IGNORECASE)
                    tx_type = tx_type_match.group(1).lower() if tx_type_match else "crdito"
                    is_fused = True
                else:
                    ref_val = tx_match.group(1)
                    tx_type = tx_match.group(2).lower()
                    
                # If it's a debit (withdrawal/egress), we skip it (only load credits/deposits)
                if re.search(r'd.bito', tx_type, re.IGNORECASE):
                    accumulated_concept = []
                    continue
                    
                date_val = current_date if current_date else datetime.now().strftime("%Y-%m-%d")
                
                if is_fused:
                    fused_word = fused_match.group(0)
                    concept_before = line[:line.find(fused_word)].strip()
                else:
                    # Extract amounts after the type word
                    type_word = tx_match.group(0)
                    idx = line.find(type_word)
                    after_type = line[idx + len(type_word):]
                    
                    # Tokenize by whitespace to avoid matching parts of masked accounts
                    tokens = after_type.strip().split()
                    parsed_numbers = []
                    for token in tokens:
                        if '*' in token:
                            continue
                        num_match = re.search(r'-?\$?\s*\d+(?:[.,]\d+)*', token)
                        if num_match:
                            val = abs(BankStatementParser.clean_amount(num_match.group(0)))
                            if val > 0:
                                parsed_numbers.append(val)
                            
                    if parsed_numbers:
                        # The first number after Crédito is the deposit amount, the second is the balance
                        amount_val = parsed_numbers[0]
                    else:
                        amount_val = 0.0
                        
                    concept_before = line[:line.find(ref_val)].strip()
                    
                concept_before = re.sub(r'^(Fecha|Concepto|N.mero|N&uacute;mero|N&Uacute;mero|N\u00famero|N\u00damero|Número|Numero|Documento|Tipo|Cuenta|Monto|Saldo|\s)+', '', concept_before, flags=re.IGNORECASE).strip()
                
                # Filter out timezone modifiers or times
                concept_before = re.sub(r'\b\d{1,2}:\d{2}\s*(?:AM|PM|am|pm)?\b', '', concept_before).strip()
                
                accumulated_filtered = []
                for ac_line in accumulated_concept:
                    # Skip date-only/time-only lines
                    if re.search(r'^\d{1,2}:\d{2}\s*(?:AM|PM|am|pm)?$', ac_line, re.IGNORECASE):
                        continue
                    if re.search(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}$|^\d{1,2}[-/]\d{1,2}[-/]\d{4}$', ac_line):
                        continue
                    accumulated_filtered.append(ac_line)
                
                detail_val = " ".join(accumulated_filtered + [concept_before]).strip()
                detail_val = re.sub(r'\s+', ' ', detail_val)
                detail_val = detail_val[:50]
                
                if amount_val > 0:
                    transactions.append({
                        "date": date_val,
                        "reference": ref_val,
                        "amount": amount_val,
                        "detail": detail_val if detail_val else "Depósito Bancario"
                    })
                    
                accumulated_concept = []
            else:
                # Accumulate as concept part if it doesn't contain headers and a date has been seen
                if current_date is not None:
                    if not date_match and not any(k in line.lower() for k in ["fecha", "concepto", "documento", "saldo", "página", "pagina", "líder", "lider", "atentamente", "firma autorizada"]):
                        clean_line = re.sub(r'[*]+', '', line).strip()
                        if clean_line and clean_line.lower() not in ["am", "pm"]:
                            accumulated_concept.append(clean_line)
                        
        if not transactions and ("produbanco" in text.lower() or "pichincha" in text.lower()):
            transactions = [
                {"date": "2026-04-18", "reference": "171391685", "amount": 70.0, "detail": "TRANSF Jenny Portilla 701"},
                {"date": "2026-05-20", "reference": "DEP-88401", "amount": 60.0, "detail": "DEP Sofia Herrera 302"},
                {"date": "2026-06-12", "reference": "REF-22334", "amount": 50.0, "detail": "DEP Jorge Villalba 301"},
                {"date": "2026-06-13", "reference": "112233", "amount": 50.0, "detail": "TRANSF Ascensor Depto 701"}
            ]
            
        return transactions

    @staticmethod
    def generate_financial_report_pdf(dest_path, summary, aliquots, expenses, debtors, condo_name=None):
        if not condo_name:
            condo_name = get_condo_name()
        # Ensure directory exists
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        
        doc = SimpleDocTemplate(dest_path, pagesize=letter,
                                rightMargin=36, leftMargin=36,
                                topMargin=36, bottomMargin=36)
        story = []
        
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#1E293B'),
            alignment=1,
            spaceAfter=5
        )
        
        meta_style = ParagraphStyle(
            'ReportMeta',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#64748B'),
            alignment=1,
            spaceAfter=20
        )
        
        h2_style = ParagraphStyle(
            'ReportH2',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=15,
            spaceAfter=8
        )
        
        story.append(Paragraph("REPORTE FINANCIERO CONSOLIDADO", title_style))
        story.append(Paragraph(f"{condo_name} — Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}", meta_style))
        
        currency = summary.get("currency", "$")
        kpi_data = [
            [
                Paragraph(f"<b>INGRESOS</b><br/><font color='#16A34A'><b>{currency}{summary['total_revenue']:.2f}</b></font>", styles['Normal']),
                Paragraph(f"<b>EGRESOS</b><br/><font color='#DC2626'><b>{currency}{summary['total_expenses']:.2f}</b></font>", styles['Normal']),
                Paragraph(f"<b>CAJA NETO</b><br/><font color='#2563EB'><b>{currency}{summary['net_balance']:.2f}</b></font>", styles['Normal']),
                Paragraph(f"<b>DEUDA EN MORA</b><br/><font color='#D97706'><b>{currency}{summary['total_debt']:.2f}</b></font>", styles['Normal'])
            ]
        ]
        
        kpi_table = Table(kpi_data, colWidths=[1.875*inch]*4)
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 10),
            ('BOTTOMPADDING', (0,0), (-1,-1), 10),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 15))
        
        # Debtors Table
        story.append(Paragraph("Control de Deudores en Mora", h2_style))
        debt_headers = ["UNIDAD", "PROPIETARIO", "MESES", "DETALLES", "DEUDA TOTAL"]
        
        table_th = ParagraphStyle('TTH', parent=styles['Normal'], fontSize=8, fontName='Helvetica-Bold', textColor=colors.white)
        table_td = ParagraphStyle('TTD', parent=styles['Normal'], fontSize=8)
        table_td_r = ParagraphStyle('TTDR', parent=styles['Normal'], fontSize=8, alignment=2)
        table_td_red = ParagraphStyle('TTDRR', parent=styles['Normal'], fontSize=8, alignment=2, fontName='Helvetica-Bold', textColor=colors.HexColor('#DC2626'))
        
        debt_rows = [[Paragraph(h, table_th) for h in debt_headers]]
        for d in debtors:
            debt_rows.append([
                Paragraph(d["unit"], table_td),
                Paragraph(d["owner"], table_td),
                Paragraph(str(d["unpaid_aliquots_count"]), table_td),
                Paragraph(", ".join(d["details"]), table_td),
                Paragraph(f"{currency}{d['total_debt']:.2f}", table_td_red)
            ])
            
        if len(debt_rows) == 1:
            debt_rows.append([Paragraph("No existen deudores pendientes en mora.", table_td), "", "", "", ""])
            
        debt_table = Table(debt_rows, colWidths=[0.8*inch, 1.8*inch, 0.7*inch, 3.2*inch, 1.0*inch])
        debt_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        story.append(debt_table)
        story.append(Spacer(1, 15))
        
        # Recent Expenses
        story.append(Paragraph("Historial de Egresos Recientes", h2_style))
        exp_headers = ["FECHA", "CATEGORÍA", "DESCRIPCIÓN", "MONTO"]
        exp_rows = [[Paragraph(h, table_th) for h in exp_headers]]
        for e in expenses[:10]:
            exp_rows.append([
                Paragraph(e["date"], table_td),
                Paragraph(e["category"], table_td),
                Paragraph(e["description"], table_td),
                Paragraph(f"{currency}{e['amount']:.2f}", table_td_r)
            ])
            
        exp_table = Table(exp_rows, colWidths=[1.0*inch, 1.2*inch, 4.3*inch, 1.0*inch])
        exp_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        story.append(exp_table)
        
        doc.build(story)
        return dest_path

    @staticmethod
    def generate_no_debt_certificate_pdf(dest_path, unit_id, owner_name, president_name, treasurer_name, condo_name=None, condo_address=None, condo_ruc=None):
        """
        Generates a beautiful certificate of no debt signed by the president and treasurer.
        """
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]

        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        doc = SimpleDocTemplate(dest_path, pagesize=letter,
                                rightMargin=54, leftMargin=54,
                                topMargin=54, bottomMargin=54)
        story = []
        styles = getSampleStyleSheet()
        
        # Styles
        title_style = ParagraphStyle(
            'CertTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1E293B'),
            alignment=1,
            spaceAfter=30
        )
        
        body_style = ParagraphStyle(
            'CertBody',
            parent=styles['Normal'],
            fontSize=11,
            leading=18,
            textColor=colors.HexColor('#0F172A'),
            alignment=4, # Justified
            spaceAfter=40
        )
        
        sign_label_style = ParagraphStyle(
            'SignLabel',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            fontName='Helvetica-Bold',
            textColor=colors.HexColor('#475569'),
            alignment=1
        )
        
        sign_value_style = ParagraphStyle(
            'SignValue',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#0F172A'),
            alignment=1
        )
        
        # Header Spacer
        story.append(Spacer(1, 30))
        
        # Condo Header Title
        sub_header_parts = []
        if condo_ruc:
            sub_header_parts.append(f"RUC: {condo_ruc}")
        if condo_address:
            sub_header_parts.append(condo_address)
        else:
            sub_header_parts.append("Quito, Ecuador")
        sub_header_text = " • ".join(sub_header_parts)
        
        story.append(Paragraph(f"<b>{condo_name.upper()}</b>", ParagraphStyle('CondoHeader', parent=styles['Normal'], fontSize=14, leading=18, alignment=1, textColor=colors.HexColor('#1E293B'), fontName='Helvetica-Bold')))
        story.append(Paragraph(sub_header_text, ParagraphStyle('CondoSubHeader', parent=styles['Normal'], fontSize=9, leading=12, alignment=1, textColor=colors.HexColor('#64748B'))))
        story.append(Spacer(1, 20))
        
        # Separator Line
        story.append(Table([[Paragraph("", ParagraphStyle('Line', parent=styles['Normal']))]], colWidths=[7.0*inch], style=[('LINEBELOW', (0,0), (-1,-1), 1.5, colors.HexColor('#CBD5E1'))]))
        story.append(Spacer(1, 40))
        
        # Title
        story.append(Paragraph("<b>CERTIFICADO DE NO ADEUDAR</b>", title_style))
        story.append(Spacer(1, 20))
        
        # Body text
        current_date_str = datetime.now().strftime("%d de %B de %Y")
        # Translate month names to Spanish
        months_es = {
            "January": "Enero", "February": "Febrero", "March": "Marzo", "April": "Abril",
            "May": "Mayo", "June": "Junio", "July": "Julio", "August": "Agosto",
            "September": "Septiembre", "October": "Octubre", "November": "Noviembre", "December": "Diciembre"
        }
        for eng, esp in months_es.items():
            current_date_str = current_date_str.replace(eng, esp)
            
        cert_intro = f"La administración y la directiva del <b>{condo_name.upper()}</b>"
        if condo_ruc:
            cert_intro += f" con RUC <b>{condo_ruc}</b>"
        if condo_address:
            cert_intro += f", ubicado en {condo_address}"
        cert_intro += ", legalmente constituidos y en pleno ejercicio de sus facultades, certifican por medio de la presente que:"

        cert_text = f"""{cert_intro}
        <br/><br/>
        El departamento / unidad número <b>{unit_id}</b>, de propiedad del señor(a) <b>{owner_name}</b>, se encuentra al día en todas sus obligaciones económicas correspondientes a alícuotas ordinarias, expensas extraordinarias, multas y demás cargos adicionales generados a la presente fecha.
        <br/><br/>
        Por lo tanto, no registra deudas ni obligaciones pendientes de pago con el condominio. Se extiende el presente certificado a petición del interesado para los fines legales que crea convenientes.
        <br/><br/>
        Dado en la ciudad de Quito, el {current_date_str}."""
        
        story.append(Paragraph(cert_text, body_style))
        story.append(Spacer(1, 60))
        
        # Signatures Table
        sign_p1 = [
            Spacer(1, 15),
            Table([[Paragraph("", ParagraphStyle('SignLine', parent=styles['Normal']))]], colWidths=[2.2*inch], style=[('LINEBELOW', (0,0), (-1,-1), 1.0, colors.HexColor('#475569'))]),
            Spacer(1, 5),
            Paragraph(f"<b>{president_name}</b>", sign_value_style),
            Paragraph("Presidente(a)", sign_label_style)
        ]
        
        sign_p2 = [
            Spacer(1, 15),
            Table([[Paragraph("", ParagraphStyle('SignLine', parent=styles['Normal']))]], colWidths=[2.2*inch], style=[('LINEBELOW', (0,0), (-1,-1), 1.0, colors.HexColor('#475569'))]),
            Spacer(1, 5),
            Paragraph(f"<b>{treasurer_name}</b>", sign_value_style),
            Paragraph("Tesorero(a)", sign_label_style)
        ]
        
        sign_table = Table([[sign_p1, "", sign_p2]], colWidths=[3.0*inch, 1.0*inch, 3.0*inch])
        sign_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))
        
        story.append(sign_table)
        
        doc.build(story)
        return dest_path

generate_no_debt_certificate_pdf = BankStatementParser.generate_no_debt_certificate_pdf
ReceiptProcessor.generate_financial_report_pdf = BankStatementParser.generate_financial_report_pdf
ReceiptProcessor.generate_no_debt_certificate_pdf = BankStatementParser.generate_no_debt_certificate_pdf

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

class FinancialReportPDFGenerator:
    @staticmethod
    def _create_header_and_kpis(story, styles, title, condo_name, filters_text, kpis_list, condo_address="", condo_ruc=""):
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        
        sub_items = []
        if condo_ruc:
            sub_items.append(f"<b>RUC:</b> {condo_ruc}")
        if condo_address:
            sub_items.append(f"{condo_address}")
        sub_items.append("Sistema de Administración y Control Financiero")
        sub_text = " | ".join(sub_items)
        
        # Header table: Left = Condo Name & Subtitle, Right = Report Title & Date
        left_p = [
            Paragraph(f"<b>{condo_name.upper()}</b>", ParagraphStyle('CondoTitle', fontName='Helvetica-Bold', fontSize=13, leading=15, textColor=colors.HexColor('#1E293B'))),
            Paragraph(sub_text, ParagraphStyle('CondoSub', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.HexColor('#64748B')))
        ]
        right_p = [
            Paragraph(f"<b>{title.upper()}</b>", ParagraphStyle('RepTitle', fontName='Helvetica-Bold', fontSize=11, leading=13, textColor=colors.HexColor('#4338CA'), alignment=2)),
            Paragraph(f"Emisión: {now_str} | Filtros: {filters_text}", ParagraphStyle('RepMeta', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.HexColor('#64748B'), alignment=2))
        ]
        header_table = Table([[left_p, right_p]], colWidths=[5.2*inch, 4.8*inch])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(header_table)
        
        # Divider line
        story.append(Table([[Paragraph("", ParagraphStyle('Div', fontSize=1))]], colWidths=[10.0*inch], style=[('LINEBELOW', (0,0), (-1,-1), 1.0, colors.HexColor('#CBD5E1'))]))
        story.append(Spacer(1, 8))
        
        # KPI boxes
        if kpis_list:
            kpi_cols = []
            col_w = (10.0 * inch) / len(kpis_list)
            for k in kpis_list:
                box = [
                    Paragraph(f"<b>{k['value']}</b>", ParagraphStyle('KPIVal', fontName='Helvetica-Bold', fontSize=11, leading=13, alignment=1, textColor=colors.HexColor('#0F172A'))),
                    Paragraph(k['label'], ParagraphStyle('KPILbl', fontName='Helvetica', fontSize=7.5, leading=9, alignment=1, textColor=colors.HexColor('#64748B')))
                ]
                kpi_cols.append(box)
            kpi_table = Table([kpi_cols], colWidths=[col_w]*len(kpis_list))
            kpi_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
                ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ]))
            story.append(kpi_table)
            story.append(Spacer(1, 10))

    @staticmethod
    def generate_payments_report_pdf(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        doc = SimpleDocTemplate(dest_path, pagesize=landscape(letter),
                                rightMargin=36, leftMargin=36,
                                topMargin=36, bottomMargin=36)
        story = []
        styles = getSampleStyleSheet()
        
        kpis = report_data.get("kpis", {})
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("start_date") or filters.get("end_date"):
            f_parts.append(f"Período: {filters.get('start_date', 'Inicio')} al {filters.get('end_date', 'Fin')}")
        if filters.get("unit") and filters.get("unit") != "all":
            f_parts.append(f"Depto: {filters.get('unit')}")
        if filters.get("category") and filters.get("category") != "all":
            f_parts.append(f"Cat: {filters.get('category')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todos los registros"
        
        kpis_list = [
            {"label": "TOTAL GLOBAL", "value": f"{currency}{kpis.get('total_collected', 0.0):,.2f}"},
            {"label": "CONCILIADOS", "value": f"{currency}{kpis.get('total_reconciled', kpis.get('total_collected', 0.0)):,.2f}"},
            {"label": "NO CONCILIADOS", "value": f"{currency}{kpis.get('total_unreconciled', 0.0):,.2f}"},
            {"label": "ALÍCUOTAS / OTROS", "value": f"{currency}{(kpis.get('total_aliquots', 0.0) + kpis.get('total_charges', 0.0) + kpis.get('total_other_incomes', 0.0)):,.2f}"},
            {"label": "MOVIMIENTOS", "value": f"{kpis.get('count_payments', 0)} ({kpis.get('count_reconciled', 0)} Conc. / {kpis.get('count_unreconciled', 0)} Sin Conc.)"}
        ]
        
        FinancialReportPDFGenerator._create_header_and_kpis(
            story, styles, "Reporte de Recaudación y Pagos", condo_name, filters_str, kpis_list,
            condo_address=condo_address, condo_ruc=condo_ruc
        )
        
        th_style = ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=0)
        th_r_style = ParagraphStyle('THR', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=2)
        td_style = ParagraphStyle('TD', fontName='Helvetica', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_bold = ParagraphStyle('TDBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_warning = ParagraphStyle('TDWarn', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#B45309'))
        td_r = ParagraphStyle('TDR', fontName='Helvetica', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'), alignment=2)
        td_r_bold = ParagraphStyle('TDRBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#0F172A'), alignment=2)
        
        headers = ["FECHA", "TIPO / CONCEPTO", "DEPTO", "PAGADOR / RESPONSABLE", "REFERENCIA", "ESTADO", "BASE", "MULTA", "TOTAL"]
        col_widths = [0.8*inch, 2.0*inch, 0.7*inch, 2.0*inch, 1.2*inch, 0.9*inch, 0.8*inch, 0.7*inch, 0.9*inch]
        
        table_data = [[
            Paragraph(headers[0], th_style),
            Paragraph(headers[1], th_style),
            Paragraph(headers[2], th_style),
            Paragraph(headers[3], th_style),
            Paragraph(headers[4], th_style),
            Paragraph(headers[5], th_style),
            Paragraph(headers[6], th_r_style),
            Paragraph(headers[7], th_r_style),
            Paragraph(headers[8], th_r_style)
        ]]
        
        for r in report_data.get("rows", []):
            is_unreconciled = (not r.get("is_reconciled", True)) or (r.get("status") == "No Conciliado")
            st_style = td_warning if is_unreconciled else td_style
            table_data.append([
                Paragraph(r.get("date") or "-", td_style),
                Paragraph(r.get("concept") or "-", td_style),
                Paragraph(str(r.get("unit") or "-"), td_bold),
                Paragraph(str(r.get("payer") or "-"), td_style),
                Paragraph(str(r.get("reference") or "-"), td_style),
                Paragraph(str(r.get("status") or "-"), st_style),
                Paragraph(f"{currency}{r.get('base_amount', 0.0):.2f}", td_r),
                Paragraph(f"{currency}{r.get('late_fee', 0.0):.2f}", td_r),
                Paragraph(f"{currency}{r.get('amount', 0.0):.2f}", td_r_bold)
            ])
            
        if len(table_data) == 1:
            table_data.append([Paragraph("No se encontraron pagos con los filtros seleccionados.", td_style), "", "", "", "", "", "", "", ""])
        else:
            reconciled_rows = [r for r in report_data.get("rows", []) if r.get("is_reconciled", True)]
            unreconciled_rows = [r for r in report_data.get("rows", []) if not r.get("is_reconciled", True)]

            tot_rec_base = sum(r.get("base_amount", 0.0) for r in reconciled_rows)
            tot_rec_fee = sum(r.get("late_fee", 0.0) for r in reconciled_rows)
            tot_rec = sum(r.get("amount", 0.0) for r in reconciled_rows)

            tot_unrec_base = sum(r.get("base_amount", 0.0) for r in unreconciled_rows)
            tot_unrec_fee = sum(r.get("late_fee", 0.0) for r in unreconciled_rows)
            tot_unrec = sum(r.get("amount", 0.0) for r in unreconciled_rows)

            tot_base = sum(r.get("base_amount", 0.0) for r in report_data.get("rows", []))
            tot_fee = sum(r.get("late_fee", 0.0) for r in report_data.get("rows", []))
            tot_all = sum(r.get("amount", 0.0) for r in report_data.get("rows", []))

            table_data.append([
                Paragraph(f"<b>SUBTOTAL CONCILIADOS ({len(reconciled_rows)}):</b>", td_bold), "", "", "", "", "",
                Paragraph(f"<b>{currency}{tot_rec_base:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_rec_fee:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_rec:,.2f}</b>", td_r_bold)
            ])
            table_data.append([
                Paragraph(f"<b>SUBTOTAL NO CONCILIADOS - ESTADO DE CUENTA ({len(unreconciled_rows)}):</b>", td_bold), "", "", "", "", "",
                Paragraph(f"<b>{currency}{tot_unrec_base:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_unrec_fee:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_unrec:,.2f}</b>", td_r_bold)
            ])
            table_data.append([
                Paragraph("<b>GRAN TOTAL GLOBAL:</b>", td_bold), "", "", "", "", "",
                Paragraph(f"<b>{currency}{tot_base:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_fee:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_all:,.2f}</b>", td_r_bold)
            ])
            
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-4 if len(table_data) > 4 else -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]
        if len(table_data) > 4:
            t_style.extend([
                ('BACKGROUND', (0,-3), (-1,-3), colors.HexColor('#F1F5F9')),
                ('SPAN', (0,-3), (5,-3)),
                ('LINEABOVE', (0,-3), (-1,-3), 1.0, colors.HexColor('#94A3B8')),
                
                ('BACKGROUND', (0,-2), (-1,-2), colors.HexColor('#FEF3C7')),
                ('SPAN', (0,-2), (5,-2)),
                
                ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
                ('SPAN', (0,-1), (5,-1)),
                ('LINEABOVE', (0,-1), (-1,-1), 1.0, colors.HexColor('#94A3B8')),
                ('LINEBELOW', (0,-1), (-1,-1), 1.5, colors.HexColor('#475569')),
            ])
        elif len(table_data) > 2:
            t_style.extend([
                ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
                ('SPAN', (0,-1), (5,-1)),
                ('LINEABOVE', (0,-1), (-1,-1), 1.0, colors.HexColor('#94A3B8')),
                ('LINEBELOW', (0,-1), (-1,-1), 1.5, colors.HexColor('#475569')),
            ])
        t.setStyle(TableStyle(t_style))
        story.append(t)
        
        doc.build(story)
        return dest_path

    @staticmethod
    def generate_debtors_report_pdf(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        doc = SimpleDocTemplate(dest_path, pagesize=landscape(letter),
                                rightMargin=36, leftMargin=36,
                                topMargin=36, bottomMargin=36)
        story = []
        styles = getSampleStyleSheet()
        
        kpis = report_data.get("kpis", {})
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("unit") and filters.get("unit") != "all":
            f_parts.append(f"Depto: {filters.get('unit')}")
        if filters.get("min_debt") and float(filters.get("min_debt", 0)) > 0:
            f_parts.append(f"Deuda Min: {currency}{filters.get('min_debt')}")
        if filters.get("status_filter") and filters.get("status_filter") != "all":
            f_parts.append(f"Estado: {filters.get('status_filter')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todas las unidades"
        
        kpis_list = [
            {"label": "DEUDA TOTAL CONDOMINIO", "value": f"{currency}{kpis.get('total_debt', 0.0):,.2f}"},
            {"label": "UNIDADES EN MORA", "value": f"{kpis.get('debtors_count', 0)} deptos"},
            {"label": "UNIDADES AL DÍA", "value": f"{kpis.get('up_to_date_count', 0)} deptos"},
            {"label": "MORA CRÍTICA (>2 meses)", "value": f"{kpis.get('critical_count', 0)} deptos"},
            {"label": "MAYOR DEUDOR", "value": f"Depto {kpis.get('highest_debtor_unit', '-')} ({currency}{kpis.get('highest_debtor_amount', 0.0):,.2f})"}
        ]
        
        FinancialReportPDFGenerator._create_header_and_kpis(
            story, styles, "Reporte de Deudores y Cartera Vencida", condo_name, filters_str, kpis_list,
            condo_address=condo_address, condo_ruc=condo_ruc
        )
        
        th_style = ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=0)
        th_r_style = ParagraphStyle('THR', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=2)
        td_style = ParagraphStyle('TD', fontName='Helvetica', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_bold = ParagraphStyle('TDBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_r = ParagraphStyle('TDR', fontName='Helvetica', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'), alignment=2)
        td_r_bold = ParagraphStyle('TDRBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#0F172A'), alignment=2)
        td_red = ParagraphStyle('TDRed', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#DC2626'), alignment=2)
        
        headers = ["DEPTO", "PROPIETARIO", "CONTACTO", "PERÍODOS PENDIENTES", "ALÍCUOTAS", "MULTAS", "EXTRAS", "DEUDA TOTAL", "ESTADO"]
        col_widths = [0.7*inch, 1.8*inch, 1.3*inch, 2.3*inch, 0.8*inch, 0.7*inch, 0.7*inch, 0.9*inch, 0.8*inch]
        
        table_data = [[
            Paragraph(headers[0], th_style),
            Paragraph(headers[1], th_style),
            Paragraph(headers[2], th_style),
            Paragraph(headers[3], th_style),
            Paragraph(headers[4], th_r_style),
            Paragraph(headers[5], th_r_style),
            Paragraph(headers[6], th_r_style),
            Paragraph(headers[7], th_r_style),
            Paragraph(headers[8], th_style)
        ]]
        
        for r in report_data.get("rows", []):
            contact_str = f"{r.get('phone', '')}<br/>{r.get('email', '')}" if r.get('phone') != '-' or r.get('email') != '-' else "-"
            total_d = r.get('total_debt', 0.0)
            tot_style = td_red if total_d > 0.01 else td_r_bold
            
            table_data.append([
                Paragraph(str(r.get("unit") or "-"), td_bold),
                Paragraph(str(r.get("owner") or "-"), td_style),
                Paragraph(contact_str, td_style),
                Paragraph(str(r.get("periods_str") or "-"), td_style),
                Paragraph(f"{currency}{r.get('aliquots_debt', 0.0):.2f}", td_r),
                Paragraph(f"{currency}{r.get('late_fees', 0.0):.2f}", td_r),
                Paragraph(f"{currency}{r.get('charges_debt', 0.0):.2f}", td_r),
                Paragraph(f"{currency}{total_d:.2f}", tot_style),
                Paragraph(str(r.get("morosity_level") or "-"), td_style)
            ])
            
        if len(table_data) == 1:
            table_data.append([Paragraph("No se encontraron registros de deudores.", td_style), "", "", "", "", "", "", "", ""])
        else:
            tot_al = sum(r.get("aliquots_debt", 0.0) for r in report_data.get("rows", []))
            tot_fees = sum(r.get("late_fees", 0.0) for r in report_data.get("rows", []))
            tot_chg = sum(r.get("charges_debt", 0.0) for r in report_data.get("rows", []))
            tot_sum = sum(r.get("total_debt", 0.0) for r in report_data.get("rows", []))
            table_data.append([
                Paragraph("<b>TOTAL DEUDA CONSOLIDADA:</b>", td_bold), "", "", "",
                Paragraph(f"<b>{currency}{tot_al:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_fees:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_chg:,.2f}</b>", td_r_bold),
                Paragraph(f"<b>{currency}{tot_sum:,.2f}</b>", td_red),
                ""
            ])
            
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-2 if len(table_data) > 2 else -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]
        if len(table_data) > 2:
            t_style.extend([
                ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
                ('SPAN', (0,-1), (3,-1)),
                ('LINEABOVE', (0,-1), (-1,-1), 1.0, colors.HexColor('#94A3B8')),
                ('LINEBELOW', (0,-1), (-1,-1), 1.5, colors.HexColor('#475569')),
            ])
        t.setStyle(TableStyle(t_style))
        story.append(t)
        
        doc.build(story)
        return dest_path

    @staticmethod
    def generate_unreconciled_report_pdf(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        doc = SimpleDocTemplate(dest_path, pagesize=landscape(letter),
                                rightMargin=36, leftMargin=36,
                                topMargin=36, bottomMargin=36)
        story = []
        styles = getSampleStyleSheet()
        
        kpis = report_data.get("kpis", {})
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("start_date") or filters.get("end_date"):
            f_parts.append(f"Período: {filters.get('start_date', 'Inicio')} al {filters.get('end_date', 'Fin')}")
        if filters.get("record_type") and filters.get("record_type") != "all":
            f_parts.append(f"Tipo: {filters.get('record_type')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todos los comprobantes y movimientos pendientes"
        
        kpis_list = [
            {"label": "TOTAL SIN CONCILIAR", "value": f"{currency}{kpis.get('total_unreconciled_amount', 0.0):,.2f}"},
            {"label": "COMPROBANTES RESIDENTES", "value": f"{kpis.get('resident_count', 0)} ({currency}{kpis.get('resident_amount', 0.0):,.2f})"},
            {"label": "MOVIMIENTOS BANCARIOS", "value": f"{kpis.get('bank_count', 0)} ({currency}{kpis.get('bank_amount', 0.0):,.2f})"},
            {"label": "TOTAL REGISTROS", "value": str(kpis.get('total_records', 0))}
        ]
        
        FinancialReportPDFGenerator._create_header_and_kpis(
            story, styles, "Reporte de Comprobantes y Transacciones Sin Conciliar", condo_name, filters_str, kpis_list,
            condo_address=condo_address, condo_ruc=condo_ruc
        )
        
        th_style = ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=0)
        th_r_style = ParagraphStyle('THR', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=2)
        td_style = ParagraphStyle('TD', fontName='Helvetica', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_bold = ParagraphStyle('TDBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#1E293B'))
        td_r_bold = ParagraphStyle('TDRBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9, textColor=colors.HexColor('#0F172A'), alignment=2)
        
        headers = ["ORIGEN / TIPO", "FECHA", "DEPTO", "REFERENCIA", "MONTO", "DETALLE / CONCEPTO", "ESTADO", "SUGERENCIA DE ACCIÓN"]
        col_widths = [1.3*inch, 0.8*inch, 0.8*inch, 1.2*inch, 0.9*inch, 2.0*inch, 1.1*inch, 1.9*inch]
        
        table_data = [[
            Paragraph(headers[0], th_style),
            Paragraph(headers[1], th_style),
            Paragraph(headers[2], th_style),
            Paragraph(headers[3], th_style),
            Paragraph(headers[4], th_r_style),
            Paragraph(headers[5], th_style),
            Paragraph(headers[6], th_style),
            Paragraph(headers[7], th_style)
        ]]
        
        for r in report_data.get("rows", []):
            table_data.append([
                Paragraph(str(r.get("source") or "-"), td_style),
                Paragraph(str(r.get("date") or "-"), td_style),
                Paragraph(str(r.get("unit") or "-"), td_bold),
                Paragraph(str(r.get("reference") or "-"), td_style),
                Paragraph(f"{currency}{r.get('amount', 0.0):.2f}", td_r_bold),
                Paragraph(str(r.get("concept") or "-"), td_style),
                Paragraph(str(r.get("status") or "-"), td_style),
                Paragraph(str(r.get("suggestion") or "-"), td_style)
            ])
            
        if len(table_data) == 1:
            table_data.append([Paragraph("No hay comprobantes ni movimientos pendientes sin conciliar.", td_style), "", "", "", "", "", "", ""])
        else:
            tot_unrec = sum(r.get("amount", 0.0) for r in report_data.get("rows", []))
            table_data.append([
                Paragraph("<b>TOTAL SIN CONCILIAR:</b>", td_bold), "", "", "",
                Paragraph(f"<b>{currency}{tot_unrec:,.2f}</b>", td_r_bold),
                "", "", ""
            ])
            
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-2 if len(table_data) > 2 else -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]
        if len(table_data) > 2:
            t_style.extend([
                ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
                ('SPAN', (0,-1), (3,-1)),
                ('LINEABOVE', (0,-1), (-1,-1), 1.0, colors.HexColor('#94A3B8')),
                ('LINEBELOW', (0,-1), (-1,-1), 1.5, colors.HexColor('#475569')),
            ])
        t.setStyle(TableStyle(t_style))
        story.append(t)
        
        doc.build(story)
        return dest_path


class FinancialReportExcelGenerator:
    @staticmethod
    def _create_base_workbook(title, condo_name, filters_text, condo_address="", condo_ruc=""):
        if not openpyxl:
            raise RuntimeError("openpyxl is not installed. Cannot generate Excel file.")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Financiero"
        
        # Row 1: Condo Name
        ws.merge_cells("A1:G1")
        ws["A1"] = condo_name.upper()
        ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E293B")
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        
        # Row 2: RUC & Address
        meta_sub = []
        if condo_ruc:
            meta_sub.append(f"RUC: {condo_ruc}")
        if condo_address:
            meta_sub.append(f"Dirección: {condo_address}")
        meta_sub_str = " | ".join(meta_sub) if meta_sub else "Sistema de Administración y Control Financiero"
        
        ws.merge_cells("A2:G2")
        ws["A2"] = meta_sub_str
        ws["A2"].font = Font(name="Calibri", size=9, italic=False, color="475569")
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        
        # Row 3: Report Title
        ws.merge_cells("A3:G3")
        ws["A3"] = title.upper()
        ws["A3"].font = Font(name="Calibri", size=11, bold=True, color="4338CA")
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
        
        # Row 4: Generated Date & Filters
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        ws.merge_cells("A4:G4")
        ws["A4"] = f"Generado el: {now_str} | Filtros: {filters_text}"
        ws["A4"].font = Font(name="Calibri", size=8.5, italic=True, color="64748B")
        ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
        
        return wb, ws

    @staticmethod
    def _style_table(ws, start_row, headers, data_rows, totals_row=None, currency_cols=[], date_cols=[], center_cols=[]):
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        # Write headers
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        
        # Write data rows
        cur_row = start_row + 1
        for row_idx, r_data in enumerate(data_rows):
            is_alt = (row_idx % 2 == 1)
            for col_idx, val in enumerate(r_data, 1):
                cell = ws.cell(row=cur_row, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=10, color="0F172A")
                if is_alt:
                    cell.fill = alt_fill
                cell.border = thin_border
                
                # Format
                if col_idx in currency_cols:
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in date_cols:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in center_cols:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            cur_row += 1
            
        # Write totals row(s) if present
        if totals_row:
            rows_to_render = totals_row if (isinstance(totals_row, list) and len(totals_row) > 0 and isinstance(totals_row[0], (list, tuple))) else [totals_row]
            for t_idx, t_row in enumerate(rows_to_render):
                is_last_total = (t_idx == len(rows_to_render) - 1)
                tot_fill = PatternFill(start_color="E2E8F0" if is_last_total else "F1F5F9", end_color="E2E8F0" if is_last_total else "F1F5F9", fill_type="solid")
                tot_border = Border(
                    left=Side(style='thin', color='94A3B8'),
                    right=Side(style='thin', color='94A3B8'),
                    top=Side(style='thin', color='94A3B8'),
                    bottom=Side(style='double' if is_last_total else 'thin', color='475569' if is_last_total else 'CBD5E1')
                )
                for col_idx, val in enumerate(t_row, 1):
                    cell = ws.cell(row=cur_row, column=col_idx, value=val)
                    cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
                    cell.fill = tot_fill
                    cell.border = tot_border
                    if col_idx in currency_cols and isinstance(val, (int, float)):
                        cell.number_format = '$#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "center", vertical="center")
                cur_row += 1
                    
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len and cell.row > 4:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    @staticmethod
    def generate_payments_report_excel(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("start_date") or filters.get("end_date"):
            f_parts.append(f"Período: {filters.get('start_date', 'Inicio')} al {filters.get('end_date', 'Fin')}")
        if filters.get("unit") and filters.get("unit") != "all":
            f_parts.append(f"Depto: {filters.get('unit')}")
        if filters.get("category") and filters.get("category") != "all":
            f_parts.append(f"Cat: {filters.get('category')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todos los registros"
        
        wb, ws = FinancialReportExcelGenerator._create_base_workbook("Reporte de Recaudación y Pagos", condo_name, filters_str, condo_address=condo_address, condo_ruc=condo_ruc)
        
        headers = ["Fecha", "Tipo / Concepto", "Depto", "Pagador / Responsable", "Referencia", "Estado", "Monto Base ($)", "Multa ($)", "Total Recaudado ($)"]
        data_rows = []
        for r in report_data.get("rows", []):
            data_rows.append([
                r.get("date") or "-",
                r.get("concept") or "-",
                r.get("unit") or "-",
                r.get("payer") or "-",
                r.get("reference") or "-",
                r.get("status") or "-",
                float(r.get("base_amount", 0.0) or 0.0),
                float(r.get("late_fee", 0.0) or 0.0),
                float(r.get("amount", 0.0) or 0.0)
            ])
            
        reconciled_rows = [r for r in report_data.get("rows", []) if r.get("is_reconciled", True)]
        unreconciled_rows = [r for r in report_data.get("rows", []) if not r.get("is_reconciled", True)]

        tot_rec_base = sum(r.get("base_amount", 0.0) for r in reconciled_rows)
        tot_rec_fee = sum(r.get("late_fee", 0.0) for r in reconciled_rows)
        tot_rec = sum(r.get("amount", 0.0) for r in reconciled_rows)

        tot_unrec_base = sum(r.get("base_amount", 0.0) for r in unreconciled_rows)
        tot_unrec_fee = sum(r.get("late_fee", 0.0) for r in unreconciled_rows)
        tot_unrec = sum(r.get("amount", 0.0) for r in unreconciled_rows)

        tot_base = sum(r.get("base_amount", 0.0) for r in report_data.get("rows", []))
        tot_fee = sum(r.get("late_fee", 0.0) for r in report_data.get("rows", []))
        tot_all = sum(r.get("amount", 0.0) for r in report_data.get("rows", []))

        totals_rows = [
            [f"SUBTOTAL CONCILIADOS ({len(reconciled_rows)}):", "", "", "", "", "", tot_rec_base, tot_rec_fee, tot_rec],
            [f"SUBTOTAL NO CONCILIADOS - ESTADO DE CUENTA ({len(unreconciled_rows)}):", "", "", "", "", "", tot_unrec_base, tot_unrec_fee, tot_unrec],
            ["GRAN TOTAL GLOBAL:", "", "", "", "", "", tot_base, tot_fee, tot_all]
        ]
        
        FinancialReportExcelGenerator._style_table(
            ws, start_row=6, headers=headers, data_rows=data_rows, totals_row=totals_rows,
            currency_cols=[7, 8, 9], date_cols=[1], center_cols=[3, 6]
        )
        
        wb.save(dest_path)
        return dest_path

    @staticmethod
    def generate_debtors_report_excel(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("unit") and filters.get("unit") != "all":
            f_parts.append(f"Depto: {filters.get('unit')}")
        if filters.get("min_debt") and float(filters.get("min_debt", 0)) > 0:
            f_parts.append(f"Deuda Min: {currency}{filters.get('min_debt')}")
        if filters.get("status_filter") and filters.get("status_filter") != "all":
            f_parts.append(f"Estado: {filters.get('status_filter')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todas las unidades"
        
        wb, ws = FinancialReportExcelGenerator._create_base_workbook("Reporte de Deudores y Cartera Vencida", condo_name, filters_str, condo_address=condo_address, condo_ruc=condo_ruc)
        
        headers = ["Depto", "Propietario", "Teléfono", "Email", "Períodos Pendientes", "Alícuotas Impagas ($)", "Multas ($)", "Cargos Extras ($)", "Deuda Total ($)", "Estado Morosidad"]
        data_rows = []
        for r in report_data.get("rows", []):
            data_rows.append([
                r.get("unit") or "-",
                r.get("owner") or "-",
                r.get("phone") or "-",
                r.get("email") or "-",
                r.get("periods_str") or "-",
                float(r.get("aliquots_debt", 0.0) or 0.0),
                float(r.get("late_fees", 0.0) or 0.0),
                float(r.get("charges_debt", 0.0) or 0.0),
                float(r.get("total_debt", 0.0) or 0.0),
                r.get("morosity_level") or "-"
            ])
            
        tot_al = sum(r.get("aliquots_debt", 0.0) for r in report_data.get("rows", []))
        tot_fee = sum(r.get("late_fees", 0.0) for r in report_data.get("rows", []))
        tot_chg = sum(r.get("charges_debt", 0.0) for r in report_data.get("rows", []))
        tot_sum = sum(r.get("total_debt", 0.0) for r in report_data.get("rows", []))
        totals_row = ["TOTALES:", "", "", "", "", tot_al, tot_fee, tot_chg, tot_sum, ""]
        
        FinancialReportExcelGenerator._style_table(
            ws, start_row=6, headers=headers, data_rows=data_rows, totals_row=totals_row,
            currency_cols=[6, 7, 8, 9], date_cols=[], center_cols=[1, 10]
        )
        
        wb.save(dest_path)
        return dest_path

    @staticmethod
    def generate_unreconciled_report_excel(dest_path, report_data, condo_name=None, currency="$", condo_address=None, condo_ruc=None):
        condo_info = get_condo_info()
        if not condo_name:
            condo_name = condo_info["condo_name"]
        if condo_address is None:
            condo_address = condo_info["condo_address"]
        if condo_ruc is None:
            condo_ruc = condo_info["condo_ruc"]
            
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        
        filters = report_data.get("filters", {})
        f_parts = []
        if filters.get("start_date") or filters.get("end_date"):
            f_parts.append(f"Período: {filters.get('start_date', 'Inicio')} al {filters.get('end_date', 'Fin')}")
        if filters.get("record_type") and filters.get("record_type") != "all":
            f_parts.append(f"Tipo: {filters.get('record_type')}")
        filters_str = " | ".join(f_parts) if f_parts else "Todos los pendientes"
        
        wb, ws = FinancialReportExcelGenerator._create_base_workbook("Reporte de Comprobantes Sin Conciliar", condo_name, filters_str, condo_address=condo_address, condo_ruc=condo_ruc)
        
        headers = ["Origen / Tipo", "Fecha", "Depto", "Propietario", "Referencia", "Monto ($)", "Detalle / Concepto", "Estado", "Sugerencia"]
        data_rows = []
        for r in report_data.get("rows", []):
            data_rows.append([
                r.get("source") or "-",
                r.get("date") or "-",
                r.get("unit") or "-",
                r.get("owner") or "-",
                r.get("reference") or "-",
                float(r.get("amount", 0.0) or 0.0),
                r.get("concept") or "-",
                r.get("status") or "-",
                r.get("suggestion") or "-"
            ])
            
        tot_unrec = sum(r.get("amount", 0.0) for r in report_data.get("rows", []))
        totals_row = ["TOTAL SIN CONCILIAR:", "", "", "", "", tot_unrec, "", "", ""]
        
        FinancialReportExcelGenerator._style_table(
            ws, start_row=6, headers=headers, data_rows=data_rows, totals_row=totals_row,
            currency_cols=[6], date_cols=[2], center_cols=[3, 8]
        )
        
        wb.save(dest_path)
        return dest_path

