import os
os.environ["TESTING"] = "True"
import json
from sqlmodel import Session, select
from sqlalchemy import text
from sheets_manager import SheetsManager
from receipt_processor import ReceiptProcessor, FinancialReportPDFGenerator, FinancialReportExcelGenerator
from ocr_helper import OCRHelper
from models import Expense, OtherIncome

def test_whatsapp_message_parsing():
    print("Testing WhatsApp message parsing...")
    msg = """Envío comprobante de pago del mes de Abril que hubo esté error.
Nombre: Jenny Portilla Bustamante 
Cédula: 171391685-4 
Departamento: 701
Mes: Abril/ 2026"""
    
    parsed = ReceiptProcessor.parse_whatsapp_message(msg)
    assert parsed["name"] == "Jenny Portilla Bustamante", f"Expected 'Jenny Portilla Bustamante', got '{parsed['name']}'"
    assert parsed["cedula"] == "171391685-4", f"Expected '171391685-4', got '{parsed['cedula']}'"
    assert parsed["unit"] == "701", f"Expected '701', got '{parsed['unit']}'"
    assert parsed["month"] == "Abril", f"Expected 'Abril', got '{parsed['month']}'"
    assert parsed["year"] == 2026, f"Expected 2026, got {parsed['year']}"
    print("[OK] WhatsApp message parsing test PASSED!")

def test_ocr_and_bank_receipt_parsing():
    print("\nTesting OCR text extraction and bank receipt detail parsing...")
    
    # Test OCR simulation for Jenny's receipt
    raw_ocr_text = OCRHelper.extract_text("comprobante_jenny.png")
    assert "Monto: $70.00" in raw_ocr_text
    assert "171391685" in raw_ocr_text
    
    amount, reference, payment_date = ReceiptProcessor.parse_bank_receipt_text(raw_ocr_text)
    assert amount == 70.00, f"Expected amount 70.00, got {amount}"
    assert reference == "171391685", f"Expected reference '171391685', got '{reference}'"
    assert payment_date == "2026-04-18", f"Expected date '2026-04-18', got '{payment_date}'"

    # Test various date formats in vouchers (Spanish words, numbers, abbreviations, mixed)
    sample_text_words = """
    BANCO PICHINCHA
    COMPROBANTE DE TRANSFERENCIA DIRECTA
    FECHA: 12 de Abril del 2026
    MONTO: $70.00
    REFERENCIA: 987654321
    """
    amt, ref, pdate = ReceiptProcessor.parse_bank_receipt_text(sample_text_words)
    assert amt == 70.00
    assert ref == "987654321"
    assert pdate == "2026-04-12"

    sample_text_abbr = """
    BANCO GUAYAQUIL
    FECHA VALOR: 15-Jul-2026 18:30
    MONTO: $60.00
    NRO. REFERENCIA: 445566
    """
    amt2, ref2, pdate2 = ReceiptProcessor.parse_bank_receipt_text(sample_text_abbr)
    assert amt2 == 60.00
    assert ref2 == "445566"
    assert pdate2 == "2026-07-15"

    sample_text_slashes_words = """
    PRODUBANCO
    Fecha de pago: 10/Julio/2026
    Monto: $80.00
    Comprobante: 123456
    """
    amt3, ref3, pdate3 = ReceiptProcessor.parse_bank_receipt_text(sample_text_slashes_words)
    assert amt3 == 80.00
    assert ref3 == "123456"
    assert pdate3 == "2026-07-10"

    # Test voucher variants and ensuring origin account is NOT selected as reference:
    # 1. No. comprobante vs Cuenta de origen
    sample_pichincha = """
    BANCO PICHINCHA
    COMPROBANTE DE TRANSFERENCIA DIRECTA
    Fecha: 15/07/2026 14:30
    Cuenta de origen: 2200123456
    Cuenta destino: 2100987654
    Beneficiario: Condominio El Mirador
    Monto: $70.00
    No. comprobante: 120576232
    Concepto: Alicuota Depto 101
    """
    a_p, r_p, d_p = ReceiptProcessor.parse_bank_receipt_text(sample_pichincha)
    assert a_p == 70.00
    assert r_p == "120576232", f"Expected voucher 120576232, got account {r_p}"
    assert d_p == "2026-07-15"

    # 2. Documento vs Desde la cuenta
    sample_produbanco_doc = """
    PRODUBANCO
    COMPROBANTE ELECTRÓNICO
    Fecha y hora: 12-Jul-2026 18:20
    Desde la cuenta: 1200543210-9
    Hacia la cuenta: 2200987654-1
    Documento: 75602180
    Valor: $70.00
    """
    a_pr, r_pr, d_pr = ReceiptProcessor.parse_bank_receipt_text(sample_produbanco_doc)
    assert a_pr == 70.00
    assert r_pr == "75602180", f"Expected document 75602180, got {r_pr}"
    assert d_pr == "2026-07-12"

    # 3. N. de comprobante vs Cuenta origen
    sample_guayaquil = """
    BANCO GUAYAQUIL
    TRANSFERENCIA DIRECTA
    N. de comprobante: 93484722
    Fecha: 2026/07/10
    Cuenta origen: 0012345678-2
    Monto transferido: $70.00
    """
    a_g, r_g, d_g = ReceiptProcessor.parse_bank_receipt_text(sample_guayaquil)
    assert a_g == 70.00
    assert r_g == "93484722", f"Expected voucher 93484722, got {r_g}"
    assert d_g == "2026-07-10"

    # 4. No. de documento vs Cuenta de débito
    sample_pacifico = """
    BANCO DEL PACÍFICO
    No. de documento: 45209982
    Cuenta de débito: 1045678901
    Monto: $70.00
    Fecha: 05 de Julio 2026
    """
    a_pac, r_pac, d_pac = ReceiptProcessor.parse_bank_receipt_text(sample_pacifico)
    assert a_pac == 70.00
    assert r_pac == "45209982", f"Expected document 45209982, got {r_pac}"
    assert d_pac == "2026-07-05"

    # 5. N° de comprobante vs Cuenta origen
    sample_inter = """
    BANCO INTERNACIONAL
    N° de comprobante: 32563490
    Cuenta origen: 4500123456
    Monto: 70.00
    Fecha: 06/07/2026
    """
    a_i, r_i, d_i = ReceiptProcessor.parse_bank_receipt_text(sample_inter)
    assert a_i == 70.00
    assert r_i == "32563490", f"Expected voucher 32563490, got {r_i}"
    assert d_i == "2026-07-06"

    # 6. Secuencia vs Cuenta de origen
    sample_jep = """
    COOPERATIVA JEP
    Secuencia: 71485490
    Cuenta de origen: 406001234567
    Monto: $70.00
    Fecha: 02-07-2026
    """
    a_j, r_j, d_j = ReceiptProcessor.parse_bank_receipt_text(sample_jep)
    assert a_j == 70.00
    assert r_j == "71485490", f"Expected sequence 71485490, got {r_j}"
    assert d_j == "2026-07-02"

    # 7. No. Operación vs Cta. Origen
    sample_bol = """
    BANCO BOLIVARIANO
    No. Operación: 88776655
    Cta. Origen: 0987654321
    Monto: $70.00
    """
    a_b, r_b, d_b = ReceiptProcessor.parse_bank_receipt_text(sample_bol)
    assert a_b == 70.00
    assert r_b == "88776655", f"Expected operation 88776655, got {r_b}"

    # 8. N° de transacción vs Cuenta
    sample_deuna = """
    DEUNA
    N° de transacción: 71295902
    Cuenta: ****9380
    Monto: $25.00
    """
    a_d, r_d, d_d = ReceiptProcessor.parse_bank_receipt_text(sample_deuna)
    assert a_d == 25.00
    assert r_d == "71295902", f"Expected tx 71295902, got {r_d}"

    # 9. Num. de Documento vs No. de cuenta débito
    sample_bgr = """
    BANCO GENERAL RUMIÑAHUI
    Num. de Documento: 99887766
    No. de cuenta débito: 1100223344
    Monto: $70.00
    """
    a_bgr, r_bgr, d_bgr = ReceiptProcessor.parse_bank_receipt_text(sample_bgr)
    assert a_bgr == 70.00
    assert r_bgr == "99887766", f"Expected doc 99887766, got {r_bgr}"

    # 10. Comprobante No. vs Cuenta de Ahorros
    sample_pich_no = """
    BANCO PICHINCHA
    COMPROBANTE DE TRANSFERENCIA
    Comprobante No. 55443322
    Cuenta de Ahorros: 2200112233
    Valor: $70.00
    """
    a_pn, r_pn, d_pn = ReceiptProcessor.parse_bank_receipt_text(sample_pich_no)
    assert a_pn == 70.00
    assert r_pn == "55443322", f"Expected voucher 55443322, got {r_pn}"

    print("[OK] OCR and bank receipt parsing test PASSED!")

def test_late_fee_calculation():
    print("\nTesting late fee calculations based on voucher deposit date...")
    sm = SheetsManager()
    
    # Day 15 is limit. Payment on April 12 -> Late fee should be 0
    fee_on_time = sm.check_late_fee("2026-04-12", "Abril", 2026)
    assert fee_on_time == 0.0, f"Expected 0.0, got {fee_on_time}"

    # Payment deposited on 12 de Abril del 2026 in words -> Late fee 0
    fee_words_on_time = sm.check_late_fee("12 de Abril del 2026", "Abril", 2026)
    assert fee_words_on_time == 0.0, f"Expected 0.0, got {fee_words_on_time}"
    
    # Payment on April 20 -> Late fee should match config late fee amount
    expected_fee = float(sm.config.get("late_fee_amount", 10.0))
    fee_late = sm.check_late_fee("2026-04-20", "Abril", 2026)
    assert fee_late == expected_fee, f"Expected {expected_fee}, got {fee_late}"

    # Payment deposited on 20 de Abril de 2026 in words -> Late fee applies
    fee_words_late = sm.check_late_fee("20 de Abril de 2026", "Abril", 2026)
    assert fee_words_late == expected_fee, f"Expected {expected_fee}, got {fee_words_late}"

    # Payment on day 15 (limit day) -> On time (0 fee)
    fee_limit_day = sm.check_late_fee("15/04/2026", "Abril", 2026)
    assert fee_limit_day == 0.0, f"Expected 0.0, got {fee_limit_day}"
    
    print("[OK] Late fee calculation test PASSED!")

def test_payment_reconciliation_and_pdf_generation():
    print("\nTesting payment reconciliation and PDF receipt generation...")
    sm = SheetsManager()
    
    # Ensure test fixtures exist
    sm.add_or_update_unit(
        unit_id="701",
        owner="Jenny Portilla",
        phone1="0999999999",
        email1="jenny@gmail.com",
        tenant="",
        phone2="",
        email2="",
        aliquot_base=70.0,
        cedula_owner="171391685"
    )
        
    expected_fee = float(sm.config.get("late_fee_amount", 10.0))
    if not any(b["reference"] == "171391685" for b in sm.data["bank_statement"]):
        sm.data["bank_statement"].append({
            "reference": "171391685",
            "date": "2026-04-18",
            "amount": 70.0 + expected_fee,
            "detail": "DEP. JENNY PORTILLA",
            "reconciled": False
        })
        
    # Clean test scenario
    # Unit 701, Month Abril 2026. Initially 'Pendiente'
    aliquots = sm.data["aliquots"]
    target_aliquot = None
    for a in aliquots:
        if a["unit"] == "701" and a["month"] == "Abril" and a["year"] == 2026:
            target_aliquot = a
            break
            
    if target_aliquot:
        target_aliquot["status"] = "Pendiente"
        target_aliquot["amount"] = 70.0
        target_aliquot["reference"] = ""
        target_aliquot["payment_date"] = ""
        target_aliquot["late_fee"] = 0.0
    else:
        # Create it pending
        target_aliquot = {
            "id": "701-2026-Abril",
            "unit": "701",
            "month": "Abril",
            "year": 2026,
            "amount": 70.0,
            "late_fee": 0.0,
            "payment_date": "",
            "reference": "",
            "status": "Pendiente"
        }
        sm.data["aliquots"].append(target_aliquot)
    
    # Set bank statement match to unreconciled
    for b in sm.data["bank_statement"]:
        if b["reference"] == "171391685":
            b["reconciled"] = False
            break
            
    sm.sync()
    
    # Run reconciliation
    payment_date = "2026-04-18" # Late (deadline is 15th)
    expected_fee = float(sm.config.get("late_fee_amount", 10.0))
    amount = 70.0 + expected_fee
    ref = "171391685"
    
    reconciled = sm.reconcile_payment(unit="701", month="Abril", year=2026, amount=amount, reference=ref, payment_date=payment_date)
    assert reconciled["aliquot"]["status"] == "Pagado"
    assert reconciled["late_fee_paid"] == expected_fee, f"Expected late fee {expected_fee}, got {reconciled['late_fee_paid']}"
    assert reconciled["aliquot"]["reference"] == ref
    
    # Verify bank transaction marked as reconciled
    bank_tx = next(b for b in sm.data["bank_statement"] if b["reference"] == ref)
    assert bank_tx["reconciled"] == True, "Expected bank transaction to be reconciled"
    
    # Test PDF Receipt creation
    pdf_path = os.path.join(os.path.dirname(__file__), "static", "receipts", f"recibo_{reconciled['aliquot']['id']}.pdf")
    
    # Remove existing PDF if any
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
        
    ReceiptProcessor.generate_receipt_pdf(
        dest_path=pdf_path,
        receipt_no=f"REC-AL-{reconciled['aliquot']['id']}",
        resident_name="Jenny Portilla Bustamante",
        unit="701",
        concept="Pago de Alícuota Ordinaria - Mes: Abril / 2026",
        amount=reconciled["aliquot_paid"],
        late_fee=reconciled["late_fee_paid"],
        reference=reconciled["aliquot"]["reference"],
        payment_date=reconciled["aliquot"]["payment_date"],
        abono_deuda=reconciled["abono_deuda"],
        pago_extra=reconciled["pago_extra"],
        saldo_deuda=reconciled["saldo_deuda"]
    )
    
    assert os.path.exists(pdf_path), "Expected PDF receipt file to be created"
    assert os.path.getsize(pdf_path) > 1000, "Expected PDF file to contain content"
    print("[OK] Payment reconciliation and PDF generation test PASSED!")

def test_partial_payment_and_debt_reduction():
    print("\nTesting partial payments (abonos) and debt reduction...")
    sm = SheetsManager()
    
    # 1. Register a mock initial debt (AdditionalCharge) of $100 for unit 701
    charge_id = "test-debt-701"
    sm.data["additional_charges"] = [c for c in sm.data["additional_charges"] if c["id"] != charge_id]
    
    sm.data["additional_charges"].append({
        "id": charge_id,
        "unit": "701",
        "type": "Deuda Inicial",
        "description": "Saldo Deudor Histórico Inicial",
        "amount": 100.0,
        "issue_date": "2026-01-01",
        "status": "Pendiente",
        "payment_date": "",
        "reference": ""
    })
    
    # Reset aliquot to base $70.0
    aliquot = next((a for a in sm.data["aliquots"] if a["unit"] == "701" and a["month"] == "Abril" and a["year"] == 2026), None)
    if not aliquot:
        aliquot = {
            "id": "701-2026-Abril",
            "unit": "701",
            "month": "Abril",
            "year": 2026,
            "amount": 70.0,
            "late_fee": 0.0,
            "paid_amount": 0.0,
            "payment_date": "",
            "reference": "",
            "status": "Pendiente"
        }
        sm.data["aliquots"].append(aliquot)
    else:
        aliquot["amount"] = 70.0
        aliquot["status"] = "Pendiente"
        aliquot["paid_amount"] = 0.0
        aliquot["reference"] = ""
        aliquot["payment_date"] = ""
        aliquot["late_fee"] = 0.0
    
    # 2. Reconcile a payment that is LESS than the aliquot ($40.0 instead of $70.0) -> partial payment (abono) to aliquot
    ref = "partial-ref-1"
    reconciled = sm.reconcile_payment(unit="701", month="Abril", year=2026, amount=40.0, reference=ref, payment_date="2026-04-10") # on time
    
    assert reconciled["aliquot"]["status"] == "Pendiente"
    assert reconciled["aliquot"]["amount"] == 30.0 # remaining aliquot balance
    assert reconciled["aliquot_paid"] == 40.0
    assert reconciled["abono_deuda"] == 0.0
    
    # 3. Now make a payment that pays off the remaining aliquot ($30.0) and pays an extra $50.0 towards the debt -> total $80.0
    ref2 = "partial-ref-2"
    reconciled2 = sm.reconcile_payment(unit="701", month="Abril", year=2026, amount=80.0, reference=ref2, payment_date="2026-04-11") # on time
    
    assert reconciled2["aliquot"]["status"] == "Pagado"
    assert reconciled2["aliquot"]["paid_amount"] == 70.0 # full aliquot paid across abonos
    assert reconciled2["aliquot_paid"] == 30.0
    assert reconciled2["abono_deuda"] == 50.0 # applied to debt
    
    # Check that the initial debt is reduced from $100.0 to $50.0!
    debt = next(c for c in sm.data["additional_charges"] if c["id"] == charge_id)
    assert debt["amount"] == 50.0, f"Expected remaining debt 50.0, got {debt['amount']}"
    assert debt["status"] == "Pendiente"
    
    # 4. Pay the remaining debt separately (using reconcile_additional_charge with amount 50.0)
    ref3 = "partial-ref-3"
    rec_charge = sm.reconcile_additional_charge(charge_id, ref3, "2026-04-12", amount=50.0)
    
    assert rec_charge["charge"]["status"] == "Pagado"
    assert rec_charge["charge"]["amount"] == 0.0
    assert rec_charge["abono_deuda"] == 50.0
    assert rec_charge["saldo_deuda"] == 0.0
    
    # Clean up test charge
    sm.delete_additional_charge(charge_id)
    
    print("[OK] Partial payments and debt reduction tests PASSED!")

def test_undo_payment():
    print("\nTesting undo/delete payment logic...")
    sm = SheetsManager()
    
    # Ensure unit 701 has base $70.0
    unit_701 = next((u for u in sm.data["units"] if u["id"] == "701"), None)
    if not unit_701:
        sm.data["units"].append({
            "id": "701",
            "owner": "Propietario 701",
            "aliquot_base": 70.0
        })
    else:
        unit_701["aliquot_base"] = 70.0
        
    # 1. Reconcile aliquot Abril 2026 for unit 701
    aliquot = next((a for a in sm.data["aliquots"] if a["unit"] == "701" and a["month"] == "Abril" and a["year"] == 2026), None)
    if not aliquot:
        aliquot = {
            "id": "701-2026-Abril",
            "unit": "701",
            "month": "Abril",
            "year": 2026,
            "amount": 70.0,
            "late_fee": 0.0,
            "paid_amount": 0.0,
            "payment_date": "",
            "reference": "",
            "status": "Pendiente"
        }
        sm.data["aliquots"].append(aliquot)
    else:
        aliquot["amount"] = 70.0
        aliquot["status"] = "Pendiente"
        aliquot["paid_amount"] = 0.0
        aliquot["reference"] = ""
        aliquot["payment_date"] = ""
        aliquot["late_fee"] = 0.0
    
    # Ensure bank transaction is unreconciled
    bank_tx = next((b for b in sm.data["bank_statement"] if b["reference"] == "171391685"), None)
    if not bank_tx:
        bank_tx = {
            "reference": "171391685",
            "date": "2026-04-10",
            "amount": 75.0,
            "detail": "TRANSF BANCO PICHINCHA PROPIETARIO 701",
            "reconciled": False
        }
        sm.data["bank_statement"].append(bank_tx)
    else:
        bank_tx["reconciled"] = False
    
    sm.sync()
    
    # Reconcile it
    reconciled = sm.reconcile_payment(unit="701", month="Abril", year=2026, amount=75.0, reference="171391685", payment_date="2026-04-10")
    assert reconciled["aliquot"]["status"] == "Pagado"
    assert bank_tx["reconciled"] == True
    
    # 2. Undo it
    success = sm.undo_payment(reconciled["aliquot"]["id"], "aliquot")
    assert success == True
    
    # Aliquot should be restored to Pendiente, and amount to 70.0
    assert aliquot["status"] == "Pendiente"
    assert aliquot["amount"] == 70.0
    assert aliquot["reference"] == ""
    assert aliquot["payment_date"] == ""
    assert aliquot["late_fee"] == 0.0
    
    # Bank transaction should be unreconciled again
    assert bank_tx["reconciled"] == False
    
    print("[OK] Undo payment tests PASSED!")

def test_advance_payment():
    print("\nTesting advance payments (pagos por adelantado) of several months...")
    sm = SheetsManager()
    
    # Ensure unit 701 has base $70.0
    unit_701 = next((u for u in sm.data["units"] if u["id"] == "701"), None)
    if not unit_701:
        sm.data["units"].append({
            "id": "701",
            "owner": "Propietario 701",
            "aliquot_base": 70.0
        })
    else:
        unit_701["aliquot_base"] = 70.0
        
    # 1. Reset aliquot Abril 2026 for unit 701 to Pendiente, and delete Mayo/Junio/Julio aliquots if any to keep test clean
    sm.data["aliquots"] = [a for a in sm.data["aliquots"] if not (a["unit"] == "701" and a["month"] in ["Mayo", "Junio", "Julio"])]
    
    aliquot_abril = next((a for a in sm.data["aliquots"] if a["unit"] == "701" and a["month"] == "Abril" and a["year"] == 2026), None)
    if not aliquot_abril:
        aliquot_abril = {
            "id": "701-2026-Abril",
            "unit": "701",
            "month": "Abril",
            "year": 2026,
            "amount": 70.0,
            "late_fee": 0.0,
            "paid_amount": 0.0,
            "payment_date": "",
            "reference": "",
            "status": "Pendiente"
        }
        sm.data["aliquots"].append(aliquot_abril)
    else:
        aliquot_abril["amount"] = 70.0
        aliquot_abril["status"] = "Pendiente"
        aliquot_abril["paid_amount"] = 0.0
        aliquot_abril["reference"] = ""
        aliquot_abril["payment_date"] = ""
        aliquot_abril["late_fee"] = 0.0
    
    # Also ensure there is no historical debt (AdditionalCharge) to keep test focused purely on advance payments
    sm.data["additional_charges"] = [c for c in sm.data["additional_charges"] if c["unit"] != "701"]
    
    sm.sync()
    
    # 2. Reconcile a payment of $210.0 for Abril 2026.
    # Aliquot for Abril is $70.0. The remaining $140.0 should pay Mayo 2026 ($70.0) and Junio 2026 ($70.0) in advance!
    ref = "advance-ref-1"
    reconciled = sm.reconcile_payment(unit="701", month="Abril", year=2026, amount=210.0, reference=ref, payment_date="2026-04-10") # on time
    
    # Abril aliquot is paid
    assert reconciled["aliquot"]["status"] == "Pagado"
    assert reconciled["aliquot_paid"] == 70.0
    
    # Mayo and Junio should be created/updated as Pagado!
    assert len(reconciled["advance_aliquots"]) == 2
    assert reconciled["advance_aliquots"][0]["month"] == "Mayo"
    assert reconciled["advance_aliquots"][0]["status"] == "Pagado"
    assert reconciled["advance_aliquots"][1]["month"] == "Junio"
    assert reconciled["advance_aliquots"][1]["status"] == "Pagado"
    
    # Verify records in the database
    aliquot_mayo = next(a for a in sm.data["aliquots"] if a["unit"] == "701" and a["month"] == "Mayo" and a["year"] == 2026)
    assert aliquot_mayo["status"] == "Pagado"
    assert aliquot_mayo["reference"] == ref
    assert aliquot_mayo["payment_date"] == "2026-04-10"
    
    aliquot_junio = next(a for a in sm.data["aliquots"] if a["unit"] == "701" and a["month"] == "Junio" and a["year"] == 2026)
    assert aliquot_junio["status"] == "Pagado"
    assert aliquot_junio["reference"] == ref
    assert aliquot_junio["payment_date"] == "2026-04-10"
    
    # 3. Test PDF generation with advance aliquots
    pdf_path = os.path.join(os.path.dirname(__file__), "static", "receipts", f"recibo_{aliquot_abril['id']}.pdf")
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
        
    ReceiptProcessor.generate_receipt_pdf(
        dest_path=pdf_path,
        receipt_no=f"REC-AL-{aliquot_abril['id']}",
        resident_name="Jenny Portilla Bustamante",
        unit="701",
        concept="Pago de Alícuota Ordinaria - Mes: Abril / 2026",
        amount=reconciled["aliquot_paid"],
        late_fee=reconciled["late_fee_paid"],
        reference=ref,
        payment_date="2026-04-10",
        abono_deuda=reconciled["abono_deuda"],
        pago_extra=reconciled["pago_extra"],
        saldo_deuda=reconciled["saldo_deuda"],
        advance_aliquots=reconciled["advance_aliquots"]
    )
    
    assert os.path.exists(pdf_path)
    assert os.path.getsize(pdf_path) > 1000
    
    # 4. Test Undo Payment - should restore all paid/future aliquots to Pendiente
    success = sm.undo_payment(aliquot_abril["id"], "aliquot")
    assert success == True
    
    assert aliquot_abril["status"] == "Pendiente"
    assert aliquot_mayo["status"] == "Pendiente"
    assert aliquot_mayo["reference"] == ""
    assert aliquot_junio["status"] == "Pendiente"
    assert aliquot_junio["reference"] == ""
    
    print("[OK] Advance payment tests PASSED!")

def test_other_incomes():
    print("Testing other incomes (general revenue not linked to units)...")
    sm = SheetsManager()
    
    # 1. Measure initial summary values (non-destructive test)
    initial_summary = sm.get_summary()
    initial_revenue = initial_summary["total_revenue"]
    initial_balance = initial_summary["net_balance"]
    
    # 3. Add other income
    inc = sm.add_other_income(
        date="2026-07-05",
        concept="Intereses Ganados en Cuenta",
        amount=15.50,
        reference="INT-8899"
    )
    
    assert inc["id"] is not None
    assert inc["amount"] == 15.50
    assert inc["concept"] == "Intereses Ganados en Cuenta"
    assert inc["reference"] == "INT-8899"
    
    # 4. Verify summary calculations updated correctly
    updated_summary = sm.get_summary()
    assert updated_summary["total_revenue"] == initial_revenue + 15.50
    assert updated_summary["net_balance"] == initial_balance + 15.50
    
    # 5. Delete other income
    success = sm.delete_other_income(inc["id"])
    assert success == True
    
    # 6. Verify summary rolled back
    rolled_back_summary = sm.get_summary()
    assert rolled_back_summary["total_revenue"] == initial_revenue
    assert rolled_back_summary["net_balance"] == initial_balance
    
    print("[OK] Other incomes tests PASSED!")

def test_board_members_exoneration():
    print("Testing board members (directiva) registration and aliquot exoneration...")
    sm = SheetsManager()
    
    # 1. Clean test records for Depto 103 and 999 if existing
    sm.delete_unit("999")
    sm.data["aliquots"] = [a for a in sm.data["aliquots"] if not (a["unit"] in ["103", "999"] and a["month"] == "Julio" and int(a.get("year", 0)) == 2026)]
    sm.data["board_members"] = [b for b in sm.data["board_members"] if b.get("unit_id") not in ["103", "999"]]
    with Session(sm.engine) as session:
        session.execute(text("DELETE FROM board_members WHERE unit_id = '103' OR unit_id = '999'"))
        session.execute(text("DELETE FROM aliquots WHERE (unit = '103' OR unit = '999') AND month = 'Julio' AND year = 2026"))
        session.commit()
    sm.load_data()
    
    # Ensure Test Unit 999 exists for test
    sm.add_or_update_unit(
        unit_id="999",
        owner="Carlos Luis Ponce",
        phone1="0999999999",
        email1="carlos@directiva.com",
        tenant="",
        phone2="",
        email2=""
    )
    
    # 2. Add board member for Unit 999
    member = sm.add_board_member(unit_id="999", name="Carlos Luis Ponce", role="Presidente")
    assert member["id"] is not None
    assert member["unit_id"] == "999"
    assert member["name"] == "Carlos Luis Ponce"
    assert member["role"] == "Presidente"
    
    # Case A: exonerate_directiva is FALSE
    sm.update_settings(
        spreadsheet_id=sm.config["spreadsheet_id"],
        google_credentials_json=sm.config["google_credentials_json"],
        use_google_sheets=False,
        late_fee_day=15,
        late_fee_amount=10.0,
        default_aliquot_base=70.0,
        exonerate_directiva=False
    )
    
    # Emit aliquots for Julio 2026
    sm.emit_aliquots_bulk(month="Julio", year=2026)
    
    # Aliquot for 999 should NOT be exonerated
    aliquot_normal = next(a for a in sm.data["aliquots"] if str(a["unit"]) == "999" and str(a["month"]).lower() == "julio" and int(a["year"]) == 2026)
    assert aliquot_normal["amount"] > 0.0
    assert aliquot_normal["status"] == "Pendiente"
    
    # Clean up that aliquot for next test case
    sm.data["aliquots"] = [a for a in sm.data["aliquots"] if not (str(a["unit"]) == "999" and str(a["month"]).lower() == "julio" and int(a.get("year", 0)) == 2026)]
    with Session(sm.engine) as session:
        session.execute(text("DELETE FROM aliquots WHERE unit = '999' AND month = 'Julio' AND year = 2026"))
        session.commit()
    sm.load_data()
    
    # Case B: exonerate_directiva is TRUE
    sm.update_settings(
        spreadsheet_id=sm.config["spreadsheet_id"],
        google_credentials_json=sm.config["google_credentials_json"],
        use_google_sheets=False,
        late_fee_day=15,
        late_fee_amount=10.0,
        default_aliquot_base=70.0,
        exonerate_directiva=True
    )
    
    revenue_before = sm.get_summary()["total_revenue"]
    
    # Emit aliquots for Julio 2026
    sm.emit_aliquots_bulk(month="Julio", year=2026)
    
    # Aliquot for 999 SHOULD be exonerated
    aliquot_exon = next(a for a in sm.data["aliquots"] if str(a["unit"]) == "999" and str(a["month"]).lower() == "julio" and int(a["year"]) == 2026)
    assert aliquot_exon["amount"] == 0.0
    assert aliquot_exon["status"] == "Exonerado"
    assert aliquot_exon["reference"] == "EXONERADO-DIRECTIVA"
    
    # Exonerated directiva must NEVER increase revenue
    summary = sm.get_summary()
    assert summary["total_revenue"] == revenue_before
    
    # Delete board member and test unit
    success = sm.delete_board_member(member["id"])
    assert success == True
    with Session(sm.engine) as session:
        session.execute(text("DELETE FROM aliquots WHERE unit = '999'"))
        session.execute(text("DELETE FROM units WHERE id = '999'"))
        session.commit()
    sm.load_data()
    
    print("[OK] Board members tests PASSED!")

def test_expenses_and_other_incomes_crud():
    print("Testing CRUD operations for expenses and other incomes...")
    sm = SheetsManager()
    
    # 1. Non-destructive CRUD testing on dedicated test items
    initial_exp_count = len(sm.data["expenses"])
    initial_inc_count = len(sm.data["other_incomes"])
    
    # 2. Test Expense CRUD
    exp = sm.add_expense("2026-07-01", "Mantenimiento", "Reparación de ascensor", 120.50)
    assert exp["id"] is not None
    assert exp["category"] == "Mantenimiento"
    assert exp["amount"] == 120.50
    
    # Update Expense
    updated_exp = sm.update_expense(exp["id"], "2026-07-02", "Limpieza", "Limpieza de cisterna", 150.00)
    assert updated_exp is not None
    assert updated_exp["category"] == "Limpieza"
    assert updated_exp["amount"] == 150.00
    assert updated_exp["description"] == "Limpieza de cisterna"
    
    # Verify in DB
    with Session(sm.engine) as session:
        db_exp = session.get(Expense, exp["id"])
        assert db_exp is not None
        assert db_exp.category == "Limpieza"
        assert db_exp.amount == 150.00
        
    # Delete Expense
    success_del_exp = sm.delete_expense(exp["id"])
    assert success_del_exp == True
    assert not any(e["id"] == exp["id"] for e in sm.data["expenses"])
    
    # Verify deleted in DB
    with Session(sm.engine) as session:
        db_exp_del = session.get(Expense, exp["id"])
        assert db_exp_del is None
        
    # 3. Test Other Income CRUD
    inc = sm.add_other_income("2026-07-03", "Alquiler salón", 80.00, "REF-SALON")
    assert inc["id"] is not None
    assert inc["concept"] == "Alquiler salón"
    assert inc["amount"] == 80.00
    
    # Update Other Income
    updated_inc = sm.update_other_income(inc["id"], "2026-07-04", "Alquiler cancha", 100.00, "REF-CANCHA")
    assert updated_inc is not None
    assert updated_inc["concept"] == "Alquiler cancha"
    assert updated_inc["amount"] == 100.00
    assert updated_inc["reference"] == "REF-CANCHA"
    
    # Verify in DB
    with Session(sm.engine) as session:
        db_inc = session.get(OtherIncome, inc["id"])
        assert db_inc is not None
        assert db_inc.concept == "Alquiler cancha"
        assert db_inc.amount == 100.00
        
    # Delete Other Income
    success_del_inc = sm.delete_other_income(inc["id"])
    assert success_del_inc == True
    assert not any(o["id"] == inc["id"] for o in sm.data["other_incomes"])
    
    # Verify deleted in DB
    with Session(sm.engine) as session:
        db_inc_del = session.get(OtherIncome, inc["id"])
        assert db_inc_del is None
        
    print("[OK] Expenses and other incomes CRUD tests PASSED!")

def test_pichincha_statement_parsing():
    print("Testing Banco Pichincha statement parser with OCR text format...")
    from receipt_processor import BankStatementParser
    
    ocr_sample = """
    Página 1 de 2
    BP - CC - 2025 - CTA3703114
    Quito, 6 de julio del 2026
    Señor(a).
    VILLAVICENCIO SOTO DIEGO ROBERTO
    Presente.
    De nuestra consideración:
    En atención a su requerimiento dirigido a Banco Pichincha C.A., nos permitimos adjuntar el detalle de las
    transacciones de la cuenta N° 2216029380, correspondiente del 1 de jul. 2026, al 6 de jul. 2026.
    ID/CI. 0602925026
    Fecha Concepto Número de
    Documento Tipo Cuenta
    Beneficiaria Monto Saldo
    2026-7-6, 6:49
    PM DEP CNB 1717764938003 172079351 Crédito $25,00 $11.610,54
    2026-7-6, 3:08
    PM
    PAGO MEER TRF
    00000000201000811863 152123192 Débito -$112,35 $11.585,54
    2026-7-6, 3:08
    PM IVA MEER TRF 000000002 152123192 Débito -$0,04 $11.697,89
    2026-7-6, 3:08
    PM COM MEER TRF 000000002 152123192 Débito -$0,27 $11.697,93
    2026-7-6,
    10:50 AM DEPOSITO 125303807 Crédito $25,00 $11.698,2
    2026-7-6,
    10:03 AM
    Transf. Directa de Salazar
    Maldonado Liseth Carolina 120576232 Crédito ******9380 $25,00 $11.673,2
    2026-7-6, 9:55
    AM
    Transf. Directa de Salazar
    Maldonado Liseth Carolina 119760533 Crédito ******9380 $25,00 $11.648,2
    2026-7-6, 6:33
    PM
    Transf. Directa de Quintero
    Fuenmayor Janeth Matilde 93484722 Crédito ******9380 $25,00 $11.623,2
    
    Página 2 de 2
    2026-7-6, 3:03
    PM
    Transf. Directa de Cahuasqui Yepez
    Richard Telmo 32563490 Crédito ******9380 $25,00 $11.598,2
    2026-7-3, 9:56
    PM
    Transf. Directa de Grandes Enriquez
    Natalia Elizabeth 75602180 Crédito ******9380 $150,00 $11.573,2
    """
    
    transactions = BankStatementParser.parse_text_statement(ocr_sample)
    
    # We expect exactly 7 credits (5 on Page 1, 2 on Page 2)
    # Debits should be skipped completely.
    assert len(transactions) == 7, f"Expected 7 transactions, got {len(transactions)}"
    
    # Check first transaction details
    tx1 = transactions[0]
    assert tx1["date"] == "2026-07-06", f"Expected date 2026-07-06, got {tx1['date']}"
    assert tx1["reference"] == "172079351", f"Expected reference '172079351', got '{tx1['reference']}'"
    assert tx1["amount"] == 25.0, f"Expected amount 25.0, got {tx1['amount']}"
    assert "DEP CNB" in tx1["detail"], f"Expected DEP CNB in detail, got '{tx1['detail']}'"
    
    # Check Salazar transaction details (tests concept multi-line accumulation)
    tx_salazar = next(t for t in transactions if t["reference"] == "120576232")
    assert tx_salazar["date"] == "2026-07-06"
    assert tx_salazar["amount"] == 25.0
    assert "Salazar Maldonado Liseth" in tx_salazar["detail"], f"Concept not matched properly: '{tx_salazar['detail']}'"
    
    # Check Grandes Enriquez transaction details (tests page 2 with larger amount)
    tx_grandes = next(t for t in transactions if t["reference"] == "75602180")
    assert tx_grandes["date"] == "2026-07-03"
    assert tx_grandes["amount"] == 150.0
    assert "Grandes Enriquez Natalia" in tx_grandes["detail"], f"Concept not matched properly: '{tx_grandes['detail']}'"
    
    print("[OK] Banco Pichincha statement parsing test PASSED!")

def test_misaligned_statement_fallback():
    print("Testing parser fallback for misaligned tables (balance in amount column)...")
    from receipt_processor import BankStatementParser
    import os
    import csv
    
    csv_content = [
        ["Fecha", "Referencia", "Monto", "Concepto"],
        ["2026-07-06", "00172079351", "11610.54", "PM DEP CNB 1717764938003 $25,"],
        ["2026-07-06", "00125303807", "11698.20", "10:50 AM DEPOSITO $25,"],
        ["2026-07-03", "0075602180", "11573.20", "Transf. Directa de Grandes $150,"]
    ]
    
    temp_csv_path = "temp_test_misaligned.csv"
    with open(temp_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(csv_content)
        
    try:
        transactions = BankStatementParser.parse_statement_file(temp_csv_path, ".csv")
        assert len(transactions) == 3, f"Expected 3, got {len(transactions)}"
        
        # Row 1 check: amount should be corrected from 11610.54 to 25.0
        assert transactions[0]["amount"] == 25.0, f"Expected 25.0, got {transactions[0]['amount']}"
        assert transactions[0]["reference"] == "00172079351"
        
        # Row 2 check: amount should be corrected from 11698.2 to 25.0
        assert transactions[1]["amount"] == 25.0, f"Expected 25.0, got {transactions[1]['amount']}"
        
        # Row 3 check: amount should be corrected from 11573.2 to 150.0
        assert transactions[2]["amount"] == 150.0, f"Expected 150.0, got {transactions[2]['amount']}"
        
        print("[OK] Misaligned statement parsing fallback test PASSED!")
    finally:
        if os.path.exists(temp_csv_path):
            os.remove(temp_csv_path)

def test_xlsx_column_mapping_exclusion():
    print("Testing parser column mapping exclusion (Valor Saldo column should not hijack amount)...")
    from receipt_processor import BankStatementParser
    import openpyxl
    import os
    
    wb = openpyxl.Workbook()
    sheet = wb.active
    
    # Header: Valor Saldo comes BEFORE Monto. Under old logic, "Valor Saldo" could match "valor" and hijack the amount column!
    headers = ["Fecha", "Concepto", "Número de Documento", "Valor Saldo", "Monto"]
    sheet.append(headers)
    
    # Write a row: Monto is 70.0, Valor Saldo (balance) is 11610.54
    row = ["2026-07-06", "PM DEP CNB", "172079351", 11610.54, 70.0]
    sheet.append(row)
    
    temp_xlsx_path = "temp_test_mapping.xlsx"
    wb.save(temp_xlsx_path)
    
    try:
        transactions = BankStatementParser.parse_statement_file(temp_xlsx_path, ".xlsx")
        assert len(transactions) == 1, f"Expected 1, got {len(transactions)}"
        
        # Verify that the parsed amount is Monto (70.0) and NOT Valor Saldo (11610.54)
        assert transactions[0]["amount"] == 70.0, f"Expected amount 70.0, got {transactions[0]['amount']}"
        
        print("[OK] Column mapping exclusion test PASSED!")
    finally:
        if os.path.exists(temp_xlsx_path):
            os.remove(temp_xlsx_path)

def test_delete_bank_transaction():
    print("Testing deletion of bank transactions and automatic payment rollback...")
    from sheets_manager import SheetsManager
    from models import BankTransaction, Aliquot
    from sqlmodel import Session, select
    
    sm = SheetsManager()
    
    # 1. Add a dummy unreconciled transaction
    tx_ref = "TX-DELETE-TEST-999"
    sm.data["bank_statement"].append({
        "reference": tx_ref,
        "date": "2026-07-06",
        "amount": 50.0,
        "detail": "Test Transaction",
        "reconciled": False
    })
    sm.sync()
    
    # Verify in DB
    with Session(sm.engine) as session:
        db_tx = session.exec(select(BankTransaction).where(BankTransaction.reference == tx_ref)).first()
        assert db_tx is not None, "Transaction should be in DB"
        
    # 2. Try deleting
    success = sm.delete_bank_transaction(tx_ref)
    assert success == True, "Should successfully delete unreconciled transaction"
    
    # Verify removed from memory and DB
    assert not any(b["reference"] == tx_ref for b in sm.data["bank_statement"]), "Should be removed from memory"
    with Session(sm.engine) as session:
        db_tx_del = session.exec(select(BankTransaction).where(BankTransaction.reference == tx_ref)).first()
        assert db_tx_del is None, "Should be removed from DB"
        
    # 3. Test deleting a reconciled transaction with associated aliquot
    tx_ref_rec = "TX-DELETE-TEST-888"
    aliquot_id = "test-del-al-1"
    sm.data["aliquots"].append({
        "id": aliquot_id,
        "unit": "101",
        "month": "Diciembre",
        "year": 2026,
        "amount": 70.0,
        "paid_amount": 70.0,
        "late_fee": 0.0,
        "payment_date": "2026-12-05",
        "reference": tx_ref_rec,
        "status": "Pagado"
    })
    sm.data["bank_statement"].append({
        "reference": tx_ref_rec,
        "date": "2026-12-05",
        "amount": 70.0,
        "detail": "Test Transaction Reconciled",
        "reconciled": True
    })
    sm.sync()
    
    rev_before = sm.get_summary()["total_revenue"]
    
    # Deleting reconciled bank transaction should succeed and undo the aliquot payment
    success_rec = sm.delete_bank_transaction(tx_ref_rec)
    assert success_rec is True, "Should successfully delete reconciled transaction and undo payment"
    assert not any(b["reference"] == tx_ref_rec for b in sm.data["bank_statement"]), "Reconciled tx should be removed from memory"
    
    al = next((a for a in sm.data["aliquots"] if a["id"] == aliquot_id), None)
    assert al is not None
    assert al["status"] == "Pendiente"
    assert al["paid_amount"] == 0.0
    
    rev_after = sm.get_summary()["total_revenue"]
    assert rev_after == rev_before - 70.0, f"Total revenue should decrease by 70.0, got {rev_before} -> {rev_after}"
    
    # Clean up test aliquot
    sm.data["aliquots"] = [a for a in sm.data["aliquots"] if a["id"] != aliquot_id]
    with Session(sm.engine) as session:
        db_al = session.exec(select(Aliquot).where(Aliquot.id == aliquot_id)).first()
        if db_al:
            session.delete(db_al)
            session.commit()
    sm.sync()
    
    print("[OK] Bank transaction deletion test PASSED!")

def test_fused_pdf_text_parsing():
    print("Testing PDF/OCR statement parsing with fused amount/reference columns...")
    from receipt_processor import BankStatementParser
    
    sample_text = """
Fecha Concepto Número de Documento Tipo Cuenta Beneficiaria Monto Saldo
2026-7-6, 6:49
PM DEP CNB 1717764938003
 $25,00172079351 Crédito $11.610,54
2026-7-6,
10:50 AM DEPOSITO
 $25,00125303807 Crédito $11.698,2
2026-7-3, 9:56
PM
Transf. Directa de Grandes Enriquez
Natalia Elizabeth
 $150,0075602180 Crédito ******9380 $11.573,2
"""
    transactions = BankStatementParser.parse_text_statement(sample_text)
    assert len(transactions) == 3, f"Expected 3 transactions, got {len(transactions)}"
    
    # 1. First transaction
    assert transactions[0]["date"] == "2026-07-06"
    assert transactions[0]["amount"] == 25.0, f"Expected 25.0, got {transactions[0]['amount']}"
    assert transactions[0]["reference"] == "172079351", f"Expected 172079351, got {transactions[0]['reference']}"
    assert "DEP CNB 1717764938003" in transactions[0]["detail"]
    
    # 2. Second transaction
    assert transactions[1]["date"] == "2026-07-06"
    assert transactions[1]["amount"] == 25.0
    assert transactions[1]["reference"] == "125303807"
    assert "DEPOSITO" in transactions[1]["detail"]
    
    # 3. Third transaction
    assert transactions[2]["date"] == "2026-07-03"
    assert transactions[2]["amount"] == 150.0
    assert transactions[2]["reference"] == "75602180"
    assert "Grandes Enriquez" in transactions[2]["detail"]
    
    print("[OK] Fused PDF/OCR text parsing test PASSED!")

def test_receipt_recipient_selection():
    print("Testing receipt recipient selection and defaults...")
    from sheets_manager import SheetsManager
    
    # 1. Initialize SheetsManager (non-destructive)
    sm = SheetsManager()
    
    # 2. Add a unit with both owner and tenant
    sm.add_or_update_unit(
        unit_id="801",
        owner="Juan Perez",
        phone1="099999999",
        email1="juan@perez.com",
        tenant="Maria Lopez",
        phone2="088888888",
        email2="maria@lopez.com"
    )
    
    # 3. Add a unit with owner only (no tenant)
    sm.add_or_update_unit(
        unit_id="802",
        owner="Carlos Gomez",
        phone1="077777777",
        email1="carlos@gomez.com",
        tenant="",
        phone2="",
        email2=""
    )
    
    # Check default setting is "inquilino"
    assert sm.config.get("receipt_recipient_type", "inquilino") == "inquilino"
    
    # Test helper get_receipt_recipient_name
    # Case A: Unit 801 (has tenant) -> should return tenant (Maria Lopez)
    name_801_default = sm.get_receipt_recipient_name("801")
    assert name_801_default == "Maria Lopez", f"Expected Maria Lopez, got {name_801_default}"
    
    # Case B: Unit 802 (no tenant) -> should fall back to owner (Carlos Gomez)
    name_802_default = sm.get_receipt_recipient_name("802")
    assert name_802_default == "Carlos Gomez", f"Expected Carlos Gomez, got {name_802_default}"
    
    # Case C: Override recipient type to "propietario" globally via settings
    sm.update_settings(
        spreadsheet_id=sm.config.get("spreadsheet_id", ""),
        google_credentials_json=sm.config.get("google_credentials_json", ""),
        use_google_sheets=sm.config.get("use_google_sheets", False),
        late_fee_day=sm.config.get("late_fee_day", 15),
        late_fee_amount=sm.config.get("late_fee_amount", 10.0),
        initial_bank_balance=sm.config.get("initial_bank_balance", 0.0),
        default_aliquot_base=sm.config.get("default_aliquot_base", 70.0),
        receipt_recipient_type="propietario"
    )
    
    # Verify global config updated
    assert sm.config["receipt_recipient_type"] == "propietario"
    
    # With global config "propietario", Unit 801 should return owner (Juan Perez)
    name_801_prop = sm.get_receipt_recipient_name("801")
    assert name_801_prop == "Juan Perez", f"Expected Juan Perez, got {name_801_prop}"
    
    # Case D: Individual override passed in (e.g. override to "inquilino" specifically)
    name_801_overridden = sm.get_receipt_recipient_name("801", recipient_override="inquilino")
    assert name_801_overridden == "Maria Lopez", f"Expected Maria Lopez, got {name_801_overridden}"
    
    # Case E: Unit has individual setting set to "propietario" (overriding global config "inquilino")
    sm.config["receipt_recipient_type"] = "inquilino"
    sm.save_config()
    sm.add_or_update_unit(
        unit_id="801",
        owner="Juan Perez",
        phone1="099999999",
        email1="juan@perez.com",
        tenant="Maria Lopez",
        phone2="088888888",
        email2="maria@lopez.com",
        receipt_recipient_type="propietario"
    )
    name_801_unit_prop = sm.get_receipt_recipient_name("801")
    assert name_801_unit_prop == "Juan Perez", f"Expected Juan Perez (unit override), got {name_801_unit_prop}"
    
    # Case F: Unit has individual setting set to "inquilino" (overriding global config "propietario")
    sm.config["receipt_recipient_type"] = "propietario"
    sm.save_config()
    sm.add_or_update_unit(
        unit_id="801",
        owner="Juan Perez",
        phone1="099999999",
        email1="juan@perez.com",
        tenant="Maria Lopez",
        phone2="088888888",
        email2="maria@lopez.com",
        receipt_recipient_type="inquilino"
    )
    name_801_unit_inq = sm.get_receipt_recipient_name("801")
    assert name_801_unit_inq == "Maria Lopez", f"Expected Maria Lopez (unit override), got {name_801_unit_inq}"
    
    # Clean up settings
    sm.config["receipt_recipient_type"] = "inquilino"
    sm.save_config()
    print("[OK] Receipt recipient selection and defaults test PASSED!")

def test_incremental_aliquots_emission():
    print("Testing incremental aliquot emission (only missing aliquots are generated)...")
    from sheets_manager import SheetsManager
    
    sm = SheetsManager()
    test_month = "Noviembre"
    test_year = 2098
    
    # Clean only leftover test records for this future test period
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE year = {test_year} AND month = '{test_month}'"))
        session.execute(text("DELETE FROM units WHERE id IN ('TEST-901', 'TEST-902', 'TEST-903')"))
        session.commit()
    sm.load_data()
    
    # 2. Add two test units
    sm.add_or_update_unit("TEST-901", "Owner One", "099999991", "one@test.com", "", "", "")
    sm.add_or_update_unit("TEST-902", "Owner Two", "099999992", "two@test.com", "", "", "")
    
    total_units_before = len(sm.data["units"])
    
    # 3. Emit aliquots for test period (should emit exactly total_units_before aliquots)
    emitted = sm.emit_aliquots_bulk(test_month, test_year)
    assert emitted == total_units_before, f"Expected {total_units_before} emitted aliquots, got {emitted}"
    
    # 4. Try to emit again for same period (should emit 0 aliquots, as all exist)
    emitted_again = sm.emit_aliquots_bulk(test_month, test_year)
    assert emitted_again == 0, f"Expected 0 emitted aliquots on duplicate run, got {emitted_again}"
    
    # 5. Add a third test unit
    sm.add_or_update_unit("TEST-903", "Owner Three", "099999993", "three@test.com", "", "", "")
    
    # 6. Emit aliquots again for test period (should emit ONLY the missing one, count = 1)
    emitted_incremental = sm.emit_aliquots_bulk(test_month, test_year)
    assert emitted_incremental == 1, f"Expected 1 emitted aliquot for the new unit, got {emitted_incremental}"
    
    # Clean up test artifacts
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE year = {test_year} AND month = '{test_month}'"))
        session.execute(text("DELETE FROM units WHERE id IN ('TEST-901', 'TEST-902', 'TEST-903')"))
        session.commit()
    sm.load_data()
    
    print("[OK] Incremental aliquot emission test PASSED!")

def test_historical_tenant_preservation():
    print("Testing preservation of historical tenant info when tenant is changed...")
    from sheets_manager import SheetsManager
    from notifications import NotificationManager
    
    sm = SheetsManager()
    test_month = "Septiembre"
    test_year = 2098
    
    # Clean up only test unit 1001 if left over
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '1001' AND year = {test_year} AND month = '{test_month}'"))
        session.execute(text("DELETE FROM units WHERE id = '1001'"))
        session.commit()
    sm.load_data()
    
    # 2. Add unit 1001 with original tenant
    sm.add_or_update_unit(
        unit_id="1001",
        owner="Juan Propietario",
        phone1="099999991",
        email1="juan@prop.com",
        tenant="Maria InquilinaOriginal",
        phone2="088888881",
        email2="maria@inq.com",
        receipt_recipient_type="inquilino"
    )
    
    # 3. Emit aliquot for test period
    sm.emit_aliquots_bulk(test_month, test_year)
    
    aliquot = next(a for a in sm.data["aliquots"] if a["unit"] == "1001" and a["month"] == test_month and a["year"] == test_year)
    assert aliquot["tenant_name"] == "Maria InquilinaOriginal", f"Expected Maria InquilinaOriginal, got {aliquot['tenant_name']}"
    
    # 4. Reconcile/pay aliquot
    sm.reconcile_payment(
        unit="1001",
        month=test_month,
        year=test_year,
        amount=70.0,
        reference="REF-SEPT-1001",
        payment_date="2026-09-05"
    )
    
    aliquot = next(a for a in sm.data["aliquots"] if a["unit"] == "1001" and a["month"] == test_month and a["year"] == test_year)
    assert aliquot["paid_by"] == "Maria InquilinaOriginal"
    
    # 5. Change tenant of unit 1001 to a new one
    sm.add_or_update_unit(
        unit_id="1001",
        owner="Juan Propietario",
        phone1="099999991",
        email1="juan@prop.com",
        tenant="Pedro InquilinoNuevo",
        phone2="088888882",
        email2="pedro@inq.com",
        receipt_recipient_type="inquilino"
    )
    
    # 6. Verify aliquot still preserves the historical tenant details!
    aliquot_after = next(a for a in sm.data["aliquots"] if a["unit"] == "1001" and a["month"] == test_month and a["year"] == test_year)
    assert aliquot_after["tenant_name"] == "Maria InquilinaOriginal", f"Expected historical Maria InquilinaOriginal, got {aliquot_after['tenant_name']}"
    assert aliquot_after["paid_by"] == "Maria InquilinaOriginal"
    assert aliquot_after["email_tenant"] == "maria@inq.com"
    assert aliquot_after["phone_tenant"] == "088888881"
    
    # 7. Check that notifications for this receipt still target the historical tenant Maria
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp_path = tmp.name
        
    sent_logs = NotificationManager.send_receipt_notifications(
        sm=sm,
        unit_id="1001",
        receipt_pdf_path=tmp_path,
        receipt_no=f"REC-AL-1001-{test_year}-{test_month}",
        resident_name=aliquot_after["paid_by"],
        concept=f"Alícuota {test_month}",
        amount=70.0,
        payment_date="2026-09-05",
        reference="REF-SEPT-1001"
    )
    
    # Check that one of the sent log targets is indeed Maria InquilinaOriginal (and not Pedro)
    assert any("Maria InquilinaOriginal" in str(log) for log in sent_logs), f"Expected Maria in notification logs, got: {sent_logs}"
    assert not any("Pedro InquilinoNuevo" in str(log) for log in sent_logs), f"Expected Pedro NOT to be notified, got: {sent_logs}"
    
    # Clean up temp file and test records
    try:
        os.remove(tmp_path)
    except:
        pass
        
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '1001' AND year = {test_year} AND month = '{test_month}'"))
        session.execute(text("DELETE FROM units WHERE id = '1001'"))
        session.commit()
    sm.load_data()
    
    print("[OK] Historical tenant preservation test PASSED!")

def test_discrepancy_name_filtering():
    print("Testing discrepancy name filtering rules...")
    from app import is_valid_name_for_mismatch_check
    
    assert not is_valid_name_for_mismatch_check("No")
    assert not is_valid_name_for_mismatch_check("no")
    assert not is_valid_name_for_mismatch_check("de")
    assert not is_valid_name_for_mismatch_check("la")
    assert not is_valid_name_for_mismatch_check("N/A")
    assert not is_valid_name_for_mismatch_check("-")
    assert not is_valid_name_for_mismatch_check("")
    
    assert is_valid_name_for_mismatch_check("Juan")
    assert is_valid_name_for_mismatch_check("Maria Lopez")
    
    print("[OK] Discrepancy name filtering rules test PASSED!")

def test_additional_charges_crud_and_app_mode():
    print("Testing additional charges CRUD and app_mode switching...")
    from sheets_manager import SheetsManager
    manager = SheetsManager()
    
    assert "app_mode" in manager.config
    
    # 1. Create a charge
    charge = manager.add_additional_charge(unit="101", type="Deuda Inicial", description="Saldo Inicial", amount=150.0)
    assert charge["amount"] == 150.0
    assert charge["description"] == "Saldo Inicial"
    
    # 2. Update the charge
    updated = manager.update_additional_charge(charge["id"], amount=120.0, description="Saldo Inicial Corregido")
    assert updated["amount"] == 120.0
    assert updated["description"] == "Saldo Inicial Corregido"
    
    # 3. Delete the charge
    success = manager.delete_additional_charge(charge["id"])
    assert success is True
    assert not any(c["id"] == charge["id"] for c in manager.data["additional_charges"])
    
    print("[OK] Additional charges CRUD and app_mode switching test PASSED!")

def test_abonos_and_revenue():
    print("Testing abonos and additional charges revenue accounting...")
    from sheets_manager import SheetsManager
    manager = SheetsManager()
    
    # 1. Non-destructive: measure initial revenue
    initial_revenue = manager.get_summary()["total_revenue"]
    
    # Create test charge of $100 for unit 101
    charge = manager.add_additional_charge(unit="101", type="Cuota Extraordinaria", description="Ascensor Test Abono", amount=100.0)
    assert charge["amount"] == 100.0
    
    # Revenue is still initial_revenue
    assert manager.get_summary()["total_revenue"] == initial_revenue
    
    # 2. Make a partial payment (abono) of $40 on the charge
    rec_res = manager.reconcile_additional_charge(charge_id=charge["id"], reference="TX-AC-1", payment_date="2026-08-16", amount=40.0)
    assert rec_res is not None
    assert rec_res["abono_deuda"] == 40.0
    assert rec_res["saldo_deuda"] == 60.0
    
    # Check that revenue increased by exactly $40.0!
    summary = manager.get_summary()
    assert summary["total_revenue"] == initial_revenue + 40.0, f"Expected {initial_revenue + 40.0} revenue after abono, got {summary['total_revenue']}"
    
    # 3. Make a full payment of the remaining $60 on the charge
    rec_res2 = manager.reconcile_additional_charge(charge_id=charge["id"], reference="TX-AC-2", payment_date="2026-08-16", amount=60.0)
    assert rec_res2 is not None
    assert rec_res2["abono_deuda"] == 60.0
    assert rec_res2["saldo_deuda"] == 0.0
    
    # Check that revenue increased by exactly $100.0!
    summary = manager.get_summary()
    assert summary["total_revenue"] == initial_revenue + 100.0, f"Expected {initial_revenue + 100.0} revenue after full payment, got {summary['total_revenue']}"
    
    # 4. Test Aliquot Partial Payment (Abono a Alícuota)
    # Create a test aliquot for Unit 101
    test_al_id = "101-2029-Diciembre"
    manager.data["aliquots"] = [a for a in manager.data["aliquots"] if a["id"] != test_al_id]
    test_al = {
        "id": test_al_id,
        "unit": "101",
        "month": "Diciembre",
        "year": 2029,
        "amount": 70.0,
        "late_fee": 0.0,
        "paid_amount": 0.0,
        "payment_date": "",
        "reference": "",
        "status": "Pendiente",
        "comprobante_url": ""
    }
    manager.data["aliquots"].append(test_al)
    manager.sync()
    
    # Revenue is initial + 100
    assert manager.get_summary()["total_revenue"] == initial_revenue + 100.0
    
    # Make a partial payment (abono) of $25 on this aliquot
    al_rec = manager.reconcile_payment(
        unit="101",
        month="Diciembre",
        year=2029,
        amount=25.0,
        reference="TX-AL-ABONO-25",
        payment_date="2029-12-05",
        late_fee=0.0
    )
    assert al_rec["aliquot_paid"] == 25.0
    assert al_rec["aliquot"]["amount"] == 45.0
    assert al_rec["aliquot"]["paid_amount"] == 25.0
    assert al_rec["aliquot"]["status"] == "Pendiente"
    
    # Total revenue must now include the $25 abono -> initial + 100 + 25 = initial + 125
    summary = manager.get_summary()
    assert summary["total_revenue"] == initial_revenue + 125.0, f"Expected {initial_revenue + 125.0} revenue after aliquot abono, got {summary['total_revenue']}"
    
    # 5. Undo the aliquot payment
    undone_al = manager.undo_payment(test_al_id, "aliquot")
    assert undone_al is True
    summary = manager.get_summary()
    assert summary["total_revenue"] == initial_revenue + 100.0, f"Expected {initial_revenue + 100.0} after aliquot undo, got {summary['total_revenue']}"
    
    # 6. Undo the charge payment
    undone = manager.undo_payment(charge["id"], "charge")
    assert undone is True
    
    # Since payments are undone, revenue should go back to initial_revenue!
    summary = manager.get_summary()
    assert summary["total_revenue"] == initial_revenue, f"Expected {initial_revenue} revenue after undo, got {summary['total_revenue']}"
    
    # Clean up test records
    manager.delete_additional_charge(charge["id"])
    manager.delete_aliquot(test_al_id)
    
    print("[OK] Abonos and additional charges revenue accounting test PASSED!")

def test_db_url_configuration():
    print("Testing DB connection URL construction and configuration...")
    from sheets_manager import SheetsManager
    manager = SheetsManager()
    orig_config = dict(manager.config)
    
    # 1. Test sqlite URL construction
    manager.config["db_type"] = "sqlite"
    url = manager.build_db_url(force_config_db_type=True)
    assert "sqlite:///" in url
    
    # 2. Configure postgres parameters
    manager.config["db_type"] = "postgres"
    manager.config["db_host"] = "localhost"
    manager.config["db_port"] = 5432
    manager.config["db_user"] = "myuser"
    manager.config["db_password"] = "mypassword"
    manager.config["db_name"] = "mydb"
    manager.config["db_custom_url"] = ""
    
    # Verify connection string format
    postgres_url = manager.build_db_url(force_config_db_type=True)
    assert "postgresql+psycopg://myuser:mypassword@127.0.0.1:5432/mydb" in postgres_url or "postgresql://" in postgres_url
    
    # 3. Configure custom connection URL
    manager.config["db_custom_url"] = "postgresql://custom_user:custom_pass@custom_host:9999/custom_db"
    custom_url = manager.build_db_url(force_config_db_type=True)
    assert "custom_user:custom_pass@custom_host:9999/custom_db" in custom_url
    
    # Restore original config
    manager.config.update(orig_config)
    
    print("[OK] DB connection URL construction test PASSED!")

def test_edit_aliquot_amount():
    print("Testing aliquot amount editing for unpaid aliquots...")
    from sheets_manager import SheetsManager
    manager = SheetsManager()
    
    test_unit = "TEST-EDIT-901"
    
    # Clean previous test records for TEST-EDIT-901
    with Session(manager.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM units WHERE id = '{test_unit}'"))
        session.commit()
    manager.load_data()
    
    # 1. Create a unit and emit an aliquot
    manager.add_or_update_unit(
        unit_id=test_unit,
        owner="Pedro Perez",
        phone1="099999999",
        email1="pedro@edit.com",
        tenant="",
        phone2="",
        email2="",
        aliquot_base=70.0,
        receipt_recipient_type="general"
    )
    
    manager.emit_aliquots_bulk("Enero", 2027)
    aliquot_id = f"{test_unit}-2027-Enero"
    
    # Find aliquot
    aliquot = next((a for a in manager.data["aliquots"] if a["id"] == aliquot_id), None)
    assert aliquot is not None
    assert aliquot["amount"] == 70.0
    assert aliquot["status"] == "Pendiente"
    
    # 2. Update amount to 85.0 (unpaid)
    updated = manager.update_aliquot_amount(aliquot_id, 85.0)
    assert updated is not None
    assert updated["amount"] == 85.0
    
    # Verify in DB/memory reload
    aliquot = next((a for a in manager.data["aliquots"] if a["id"] == aliquot_id), None)
    assert aliquot["amount"] == 85.0
    
    # 3. Pay the aliquot
    manager.update_aliquot_status(aliquot_id, "Pagado", "2027-01-10", "REF-901")
    assert aliquot["status"] == "Pagado"
    
    # 4. Attempt to edit paid aliquot should raise ValueError
    try:
        manager.update_aliquot_amount(aliquot_id, 100.0)
        assert False, "Should have raised ValueError when editing paid aliquot"
    except ValueError as e:
        assert "ya ha sido pagada" in str(e)
        
    # 5. Attempt to set negative amount on unpaid aliquot should raise ValueError
    manager.emit_aliquots_bulk("Febrero", 2027)
    aliquot_feb = f"{test_unit}-2027-Febrero"
    try:
        manager.update_aliquot_amount(aliquot_feb, -50.0)
        assert False, "Should have raised ValueError for negative amount"
    except ValueError as e:
        assert "no puede ser negativo" in str(e)
        
    # Clean up test records
    with Session(manager.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM units WHERE id = '{test_unit}'"))
        session.commit()
    manager.load_data()
        
    print("[OK] Aliquot amount editing test PASSED!")

def test_role_access_control():
    print("Testing Role-Based Access Control (RBAC) middleware...")
    import asyncio
    from app import role_access_middleware
    from starlette.requests import Request
    from starlette.responses import Response

    async def run_checks():
        async def fake_call_next(req):
            return Response("OK", status_code=200)

        # 1. Owner mutating expenses -> 403
        scope = {
            "type": "http",
            "method": "POST",
            "path": "/api/expenses",
            "headers": [(b"x-user-role", b"owner")],
        }
        req = Request(scope)
        resp = await role_access_middleware(req, fake_call_next)
        assert resp.status_code == 403, f"Expected 403 for owner, got {resp.status_code}"

        # 2. Tenant mutating settings -> 403
        scope = {
            "type": "http",
            "method": "POST",
            "path": "/api/settings",
            "headers": [(b"x-user-role", b"tenant")],
        }
        req = Request(scope)
        resp = await role_access_middleware(req, fake_call_next)
        assert resp.status_code == 403, f"Expected 403 for tenant, got {resp.status_code}"

        # 3. Owner reading summary -> 200
        scope = {
            "type": "http",
            "method": "GET",
            "path": "/api/summary",
            "headers": [(b"x-user-role", b"owner")],
        }
        req = Request(scope)
        resp = await role_access_middleware(req, fake_call_next)
        assert resp.status_code == 200, f"Expected 200 for owner GET, got {resp.status_code}"

        # 4. Admin mutating expenses -> 200
        scope = {
            "type": "http",
            "method": "POST",
            "path": "/api/expenses",
            "headers": [(b"x-user-role", b"admin")],
        }
        req = Request(scope)
        resp = await role_access_middleware(req, fake_call_next)
        assert resp.status_code == 200, f"Expected 200 for admin POST, got {resp.status_code}"

    asyncio.run(run_checks())
    print("[OK] Role-Based Access Control (RBAC) test PASSED!")

def test_database_config_encryption_and_initial_balances():
    print("Testing Database Config Storage, Key Encryption, and Separate Initial Balances Table...")
    from models import SystemConfig, InitialBalance
    from crypto_helper import decrypt_secret
    
    manager = SheetsManager()
    orig_cfg = dict(manager.config)
    
    real_db_pwd = manager.config.get("db_password", "1351")
    
    # 1. Update system settings with sensitive keys
    manager.update_settings(
        spreadsheet_id="test_sheet_123",
        google_credentials_json='{"type": "service_account", "private_key": "my_secret_key"}',
        use_google_sheets=False,
        late_fee_day=15,
        late_fee_amount=12.50,
        initial_bank_balance=350.0,
        default_aliquot_base=75.0,
        smtp_host="smtp.gmail.com",
        smtp_port=587,
        smtp_user="condominio.admin@gmail.com",
        smtp_password="MySecretSmtpPassword999!",
        smtp_from="condominio.admin@gmail.com",
        twilio_sid="ACtest123456",
        twilio_token="TwilioTokenSecret789",
        twilio_whatsapp_from="whatsapp:+14155238886",
        whatsapp_provider="twilio",
        meta_wa_token="EAAG_Meta_Secret_Token_456",
        meta_wa_phone_number_id="10987654321",
        meta_wa_verify_token="MetaVerifySecretToken",
        condo_name="Condominio Vista Hermosa",
        db_password=real_db_pwd
    )
    
    # 2. Check direct database row in SystemConfig table
    with Session(manager.engine) as session:
        db_cfg = session.exec(select(SystemConfig).where(SystemConfig.id == 1)).first()
        assert db_cfg is not None, "SystemConfig record not found in database!"
        assert db_cfg.condo_name == "Condominio Vista Hermosa"
        assert db_cfg.late_fee_amount == 12.50
        
        # Verify passwords are ENCRYPTED in the database (Fernet token starting with gAAAAA)
        assert db_cfg.smtp_password.startswith("gAAAAA"), "SMTP password must be encrypted in database"
        assert db_cfg.smtp_password != "MySecretSmtpPassword999!", "Plaintext SMTP password leaked in database"
        assert decrypt_secret(db_cfg.smtp_password) == "MySecretSmtpPassword999!", "Decrypted SMTP password mismatch"
        
        assert db_cfg.twilio_token.startswith("gAAAAA"), "Twilio token must be encrypted in database"
        assert decrypt_secret(db_cfg.twilio_token) == "TwilioTokenSecret789"
        
        assert db_cfg.meta_wa_token.startswith("gAAAAA"), "Meta token must be encrypted in database"
        assert decrypt_secret(db_cfg.meta_wa_token) == "EAAG_Meta_Secret_Token_456"
        
        assert db_cfg.db_password.startswith("gAAAAA"), "DB password must be encrypted in database"
        assert decrypt_secret(db_cfg.db_password) == real_db_pwd

    # 3. Check InitialBalance table separately
    with Session(manager.engine) as session:
        db_bal = session.exec(select(InitialBalance).where(InitialBalance.id == 1)).first()
        assert db_bal is not None, "InitialBalance record not found in database!"
        assert db_bal.initial_bank_balance == 350.0

    # 4. Update initial balances using dedicated method
    manager.update_initial_balances(
        initial_bank_balance=1200.0,
        initial_reserve_fund=500.0,
        cut_off_date="2026-01-01",
        description="Saldo inicial consolidado",
        notes="Aprobado en asamblea general"
    )
    
    with Session(manager.engine) as session:
        db_bal2 = session.exec(select(InitialBalance).where(InitialBalance.id == 1)).first()
        assert db_bal2.initial_bank_balance == 1200.0
        assert db_bal2.initial_reserve_fund == 500.0
        assert db_bal2.cut_off_date == "2026-01-01"
        
    summary = manager.get_summary()
    assert summary["initial_bank_balance"] == 1200.0

    # 5. Restore original configuration and initial balances
    manager.config.update(orig_cfg)
    manager.save_config()
    manager.update_initial_balances(
        initial_bank_balance=float(orig_cfg.get("initial_bank_balance", 0.0)),
        initial_reserve_fund=float(orig_cfg.get("initial_reserve_fund", 0.0))
    )
    
    print("[OK] Database Config Storage, Key Encryption, and Separate Initial Balances test PASSED!")

def test_delete_aliquot_superadmin_only():
    print("Testing Aliquot Deletion restricted to Superadmin only...")
    from fastapi import Request, HTTPException, Response
    from app import sm, delete_aliquot, role_access_middleware
    import asyncio
    
    # 1. Create a temporary unit and aliquot
    sm.add_or_update_unit(
        unit_id="888",
        owner="Propietario Temporal",
        phone1="0991112223",
        email1="temp888@test.com",
        tenant="",
        phone2="",
        email2=""
    )
    test_aliquot = {
        "id": "888-2026-Diciembre",
        "unit": "888",
        "month": "Diciembre",
        "year": 2026,
        "amount": 75.0,
        "late_fee": 0.0,
        "payment_date": "",
        "reference": "",
        "status": "Pendiente",
        "comprobante_url": ""
    }
    sm.data["aliquots"].append(test_aliquot)
    sm.sync()
    
    assert any(a["id"] == "888-2026-Diciembre" for a in sm.data["aliquots"])
    
    # 2. Try deleting as tenant/owner via middleware -> 403 Forbidden
    async def run_middleware_checks():
        async def fake_call_next(request):
            return Response("OK", status_code=200)
            
        scope = {
            "type": "http",
            "method": "DELETE",
            "path": "/api/aliquots/888-2026-Diciembre",
            "headers": [(b"x-user-role", b"tenant")],
        }
        resp = await role_access_middleware(Request(scope), fake_call_next)
        assert resp.status_code == 403, f"Expected 403 for tenant, got {resp.status_code}"
        
        scope["headers"] = [(b"x-user-role", b"owner")]
        resp = await role_access_middleware(Request(scope), fake_call_next)
        assert resp.status_code == 403, f"Expected 403 for owner, got {resp.status_code}"

    asyncio.run(run_middleware_checks())
    
    # 3. Try deleting as regular admin via endpoint -> 403 Forbidden
    req_admin = Request({
        "type": "http",
        "method": "DELETE",
        "path": "/api/aliquots/888-2026-Diciembre",
        "headers": [(b"x-user-role", b"admin")]
    })
    try:
        delete_aliquot("888-2026-Diciembre", req_admin)
        assert False, "Admin should not be able to delete aliquot!"
    except HTTPException as e:
        assert e.status_code == 403
        assert "Superadministrador" in e.detail

    # 4. Try deleting without role header -> 403 Forbidden
    req_anon = Request({
        "type": "http",
        "method": "DELETE",
        "path": "/api/aliquots/888-2026-Diciembre",
        "headers": []
    })
    try:
        delete_aliquot("888-2026-Diciembre", req_anon)
        assert False, "Anonymous should not be able to delete aliquot!"
    except HTTPException as e:
        assert e.status_code == 403

    # 5. Delete as superadmin -> 200 OK
    req_super = Request({
        "type": "http",
        "method": "DELETE",
        "path": "/api/aliquots/888-2026-Diciembre",
        "headers": [(b"x-user-role", b"superadmin")]
    })
    resp = delete_aliquot("888-2026-Diciembre", req_super)
    assert resp["status"] == "success"
    
    # Verify aliquot is deleted from memory and database
    assert not any(a["id"] == "888-2026-Diciembre" for a in sm.data["aliquots"])
    with Session(sm.engine) as session:
        from models import Aliquot
        db_a = session.exec(select(Aliquot).where(Aliquot.id == "888-2026-Diciembre")).first()
        assert db_a is None, "Aliquot was not deleted from database!"
        
    # 6. Try deleting non-existent aliquot -> 404
    try:
        delete_aliquot("888-2026-Diciembre", req_super)
        assert False, "Non-existent aliquot should return 404!"
    except HTTPException as e:
        assert e.status_code == 404

    # Clean up test unit
    sm.delete_unit("888")
    
    print("[OK] Aliquot deletion restricted to Superadmin test PASSED!")

def test_individual_aliquot_emission():
    print("Testing Individual Aliquot Emission per Unit...")
    from sheets_manager import SheetsManager
    from app import sm, emit_aliquots, AliquotEmitRequest
    from fastapi import HTTPException
    
    # 1. Ensure test unit 101 exists
    sm.add_or_update_unit(
        unit_id="101",
        owner="Juan Perez",
        phone1="0991234567",
        email1="juan@test.com",
        tenant="",
        phone2="",
        email2="",
        aliquot_base=70.0
    )
    
    # Clean up any pre-existing test aliquot
    test_id = "101-2028-Noviembre"
    sm.delete_aliquot(test_id)
    
    # 2. Emit individually via SheetsManager
    al = sm.emit_aliquot_individual(unit="101", month="Noviembre", year=2028)
    assert al is not None
    assert al["id"] == test_id
    assert al["unit"] == "101"
    assert al["month"] == "Noviembre"
    assert al["year"] == 2028
    assert al["amount"] == 70.0
    assert al["status"] == "Pendiente"
    assert any(a["id"] == test_id for a in sm.data["aliquots"])
    
    # 3. Emitting same individual aliquot again should raise ValueError
    try:
        sm.emit_aliquot_individual(unit="101", month="Noviembre", year=2028)
        assert False, "Should raise ValueError for already existing aliquot"
    except ValueError as e:
        assert "ya existe" in str(e)
        
    # 4. Emitting for non-existent unit should raise ValueError
    try:
        sm.emit_aliquot_individual(unit="99999", month="Noviembre", year=2028)
        assert False, "Should raise ValueError for non-existent unit"
    except ValueError as e:
        assert "no existe" in str(e)
        
    # 5. Emit individually with custom amount
    test_id_custom = "101-2028-Diciembre"
    sm.delete_aliquot(test_id_custom)
    al_custom = sm.emit_aliquot_individual(unit="101", month="Diciembre", year=2028, amount=85.50)
    assert al_custom["amount"] == 85.50
    assert al_custom["id"] == test_id_custom
    
    # 6. Test via API endpoint emit_aliquots
    sm.add_or_update_unit(
        unit_id="102",
        owner="Maria Lopez",
        phone1="0992345678",
        email1="maria@test.com",
        tenant="",
        phone2="",
        email2="",
        aliquot_base=75.0
    )
    test_id_api = "102-2028-Noviembre"
    sm.delete_aliquot(test_id_api)
    req = AliquotEmitRequest(unit_id="102", month="Noviembre", year=2028, amount=75.0)
    resp = emit_aliquots(req)
    assert resp["status"] == "success"
    assert resp["emitted"] == 1
    assert any(a["id"] == test_id_api for a in sm.data["aliquots"])
    
    # 7. Clean up test aliquots
    sm.delete_aliquot(test_id)
    sm.delete_aliquot(test_id_custom)
    sm.delete_aliquot(test_id_api)
    
    print("[OK] Individual Aliquot Emission test PASSED!")

def test_reference_edit_and_auto_reconcile():
    print("Testing Bank Reference Editing and Automatic Reconciliation...")
    from app import (
        sm, update_aliquot, update_additional_charge, update_bank_transaction_endpoint,
        AliquotUpdate, ChargeUpdate, BankTransactionUpdate
    )
    from sqlmodel import Session, text
    
    test_unit = "TEST-REF-SYS1"
    
    # 0. Clean test records
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM additional_charges WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM units WHERE id = '{test_unit}'"))
        session.execute(text("DELETE FROM bank_statement WHERE reference LIKE 'TEST-SYS-REF-%'"))
        session.commit()
    sm.load_data()

    # Create test unit
    sm.add_or_update_unit(
        unit_id=test_unit,
        owner="Propietario Test AutoRec",
        phone1="0988776655",
        email1="autorec@test.com",
        tenant="",
        phone2="",
        email2="",
        aliquot_base=80.0,
        receipt_recipient_type="general"
    )

    # 1. Add bank statement transaction
    test_ref_1 = "TEST-SYS-REF-1001"
    sm.add_bank_transaction("2026-09-18", test_ref_1, 80.0, "TRANSF BANCO TEST SYS")
    
    # Emit aliquot for test unit
    sm.emit_aliquot_individual(unit=test_unit, month="Octubre", year=2028, amount=80.0)
    aliquot_id = f"{test_unit}-2028-Octubre"
    
    aliquot = next((a for a in sm.data["aliquots"] if a["id"] == aliquot_id), None)
    assert aliquot is not None
    assert aliquot["status"] == "Pendiente"

    # 2. Test Partial / Substring Reference does NOT auto-reconcile (strict exact matching)
    partial_ref = "TEST-SYS-REF-10"
    res_partial = update_aliquot(aliquot_id, AliquotUpdate(amount=80.0, reference=partial_ref, auto_reconcile=True))
    assert res_partial["reconciled"] is False
    aliquot = next((a for a in sm.data["aliquots"] if a["id"] == aliquot_id), None)
    assert aliquot["status"] == "Validación Manual"
    assert aliquot["reference"] == partial_ref

    # 3. Test Exact Reference Match DOES auto-reconcile
    res_exact = update_aliquot(aliquot_id, AliquotUpdate(amount=80.0, reference=test_ref_1, auto_reconcile=True))
    assert res_exact["reconciled"] is True
    assert res_exact["aliquot"]["status"] == "Pagado"
    
    bank_tx = next((b for b in sm.data["bank_statement"] if b["reference"] == test_ref_1), None)
    assert bank_tx["reconciled"] is True

    # 4. Test Additional Charge Reference Auto-Reconcile
    test_ref_charge = "TEST-SYS-REF-2002"
    sm.add_bank_transaction("2026-09-18", test_ref_charge, 35.0, "TRANSF MULTA TEST SYS")
    
    charge = sm.add_additional_charge(
        unit=test_unit,
        type="Multa",
        description="Multa prueba auto-reconciliacion",
        amount=35.0
    )
    charge_id = charge["id"]
    
    res_charge = update_additional_charge(
        charge_id,
        ChargeUpdate(amount=35.0, description="Multa prueba auto-reconciliacion", reference=test_ref_charge, auto_reconcile=True)
    )
    assert res_charge["reconciled"] is True
    assert res_charge["charge"]["status"] == "Pagado"
    
    bank_tx_c = next((b for b in sm.data["bank_statement"] if b["reference"] == test_ref_charge), None)
    assert bank_tx_c["reconciled"] is True

    # 5. Test Bank Transaction Edit (Reverse Auto-Reconcile)
    target_ref_3 = "TEST-SYS-REF-3003"
    sm.emit_aliquot_individual(unit=test_unit, month="Noviembre", year=2028, amount=80.0)
    aliquot_id_3 = f"{test_unit}-2028-Noviembre"
    update_aliquot(aliquot_id_3, AliquotUpdate(amount=80.0, reference=target_ref_3, auto_reconcile=True))
    
    aliquot_3 = next((a for a in sm.data["aliquots"] if a["id"] == aliquot_id_3), None)
    assert aliquot_3["status"] == "Validación Manual"
    
    typo_ref = "TEST-SYS-REF-3003-TYPO"
    sm.add_bank_transaction("2026-09-18", typo_ref, 80.0, "DEPOSITO NOV TEST")
    
    res_bank_update = update_bank_transaction_endpoint(
        old_reference=typo_ref,
        req=BankTransactionUpdate(reference=target_ref_3, date="2026-09-18", amount=80.0, detail="DEPOSITO NOV TEST CORREGIDO")
    )
    assert res_bank_update["reconciled"] is True
    
    aliquot_3 = next((a for a in sm.data["aliquots"] if a["id"] == aliquot_id_3), None)
    assert aliquot_3["status"] == "Pagado"

    # 6. Clean up test records
    with Session(sm.engine) as session:
        session.execute(text(f"DELETE FROM aliquots WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM additional_charges WHERE unit = '{test_unit}'"))
        session.execute(text(f"DELETE FROM units WHERE id = '{test_unit}'"))
        session.execute(text("DELETE FROM bank_statement WHERE reference LIKE 'TEST-SYS-REF-%'"))
        session.commit()
    sm.load_data()

    print("[OK] Bank Reference Editing and Auto-Reconciliation test PASSED!")

def test_financial_reports():
    print("\nTesting Financial Reports Module (Payments, Debtors, Unreconciled)...")
    sm = SheetsManager()
    output_dir = os.path.join(os.path.dirname(__file__), "static", "receipts")
    os.makedirs(output_dir, exist_ok=True)
    
    # 1. Payments Report
    p_data = sm.get_payments_report_data()
    assert "records" in p_data
    assert "kpis" in p_data
    assert "total_collected" in p_data["kpis"]
    pdf_p = os.path.join(output_dir, "sys_test_rep_pagos.pdf")
    FinancialReportPDFGenerator.generate_payments_report_pdf(pdf_p, p_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(pdf_p)
    
    xlsx_p = os.path.join(output_dir, "sys_test_rep_pagos.xlsx")
    FinancialReportExcelGenerator.generate_payments_report_excel(xlsx_p, p_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(xlsx_p)

    # 2. Debtors Report
    d_data = sm.get_debtors_report_data(status_filter="all")
    assert "records" in d_data
    assert "kpis" in d_data
    assert "total_debt" in d_data["kpis"]
    pdf_d = os.path.join(output_dir, "sys_test_rep_deudores.pdf")
    FinancialReportPDFGenerator.generate_debtors_report_pdf(pdf_d, d_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(pdf_d)

    xlsx_d = os.path.join(output_dir, "sys_test_rep_deudores.xlsx")
    FinancialReportExcelGenerator.generate_debtors_report_excel(xlsx_d, d_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(xlsx_d)

    # 3. Unreconciled Report
    u_data = sm.get_unreconciled_report_data()
    assert "records" in u_data
    assert "kpis" in u_data
    assert "total_unreconciled_amount" in u_data["kpis"]
    pdf_u = os.path.join(output_dir, "sys_test_rep_sin_conciliar.pdf")
    FinancialReportPDFGenerator.generate_unreconciled_report_pdf(pdf_u, u_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(pdf_u)

    xlsx_u = os.path.join(output_dir, "sys_test_rep_sin_conciliar.xlsx")
    FinancialReportExcelGenerator.generate_unreconciled_report_excel(xlsx_u, u_data, condo_name="Condominio Test", currency="$")
    assert os.path.exists(xlsx_u)

    print("[OK] Financial Reports (Preview, PDF, Excel) test PASSED!")

def test_no_debt_certificate_generation():
    print("\nTesting No Debt Certificate PDF Generation and Validation...")
    from receipt_processor import generate_no_debt_certificate_pdf
    
    test_pdf_path = os.path.join(os.path.dirname(__file__), "static", "receipts", "test_cert_no_debt_101.pdf")
    if os.path.exists(test_pdf_path):
        os.remove(test_pdf_path)
        
    res_path = generate_no_debt_certificate_pdf(
        dest_path=test_pdf_path,
        unit_id="101",
        owner_name="Juan Perez",
        president_name="Carlos Mendoza",
        treasurer_name="Sofia Herrera",
        condo_name="Condominio El Mirador"
    )
    
    assert os.path.exists(test_pdf_path), "Certificate PDF was not generated on disk."
    assert os.path.getsize(test_pdf_path) > 1000, "Certificate PDF file is empty or corrupted."
    
    # Test endpoint logic directly
    from app import download_no_debt_certificate
    from fastapi import HTTPException
    
    # Find a unit with no debt or with debt
    sm = SheetsManager()
    debtors = sm.get_debtors()
    debtor_units = [d["unit"] for d in debtors if d["total_debt"] > 0]
    
    # Check debtor unit fails with 400
    if debtor_units:
        d_unit = debtor_units[0]
        try:
            download_no_debt_certificate(d_unit)
            assert False, f"Expected HTTPException 400 for debtor unit {d_unit}"
        except HTTPException as he:
            assert he.status_code == 400
            assert "posee deudas pendientes" in he.detail
            
    # Clean up test pdf
    if os.path.exists(test_pdf_path):
        os.remove(test_pdf_path)
        
    print("[OK] No Debt Certificate PDF generation & validation test PASSED!")

def test_condo_address_and_ruc_config_and_reports():
    print("\nTesting Condo Address and RUC Database Configuration & Document Rendering...")
    import pypdf
    import openpyxl
    
    sm = SheetsManager()
    
    # 1. Update settings with condo_address and condo_ruc
    orig_condo_name = sm.config.get("condo_name", "Condominio El Mirador")
    orig_condo_ruc = sm.config.get("condo_ruc", "")
    orig_condo_address = sm.config.get("condo_address", "")

    test_ruc = "1792345678001"
    test_address = "Av. República de El Salvador N36-84 y Naciones Unidas, Quito"
    test_condo_name = "Condominio Torres del Valle"
    
    res = sm.update_settings(
        condo_name=test_condo_name,
        condo_ruc=test_ruc,
        condo_address=test_address
    )
    assert res["status"] == "success"
    assert sm.config.get("condo_ruc") == test_ruc
    assert sm.config.get("condo_address") == test_address
    assert sm.config.get("condo_name") == test_condo_name
    
    # Verify persistence in Database
    from models import SystemConfig
    with Session(sm.engine) as session:
        db_cfg = session.exec(select(SystemConfig)).first()
        assert db_cfg is not None
        assert db_cfg.condo_ruc == test_ruc
        assert db_cfg.condo_address == test_address
        assert db_cfg.condo_name == test_condo_name
        
    output_dir = os.path.join(os.path.dirname(__file__), "static", "receipts")
    os.makedirs(output_dir, exist_ok=True)
    
    # 2. Test No Debt Certificate generation with RUC & Address
    cert_pdf_path = os.path.join(output_dir, "test_cert_with_ruc_addr.pdf")
    if os.path.exists(cert_pdf_path):
        os.remove(cert_pdf_path)
        
    from receipt_processor import generate_no_debt_certificate_pdf
    generate_no_debt_certificate_pdf(
        dest_path=cert_pdf_path,
        unit_id="201",
        owner_name="Rodrigo Morales",
        president_name="Carlos Mendoza",
        treasurer_name="Sofia Herrera",
        condo_name=test_condo_name,
        condo_address=test_address,
        condo_ruc=test_ruc
    )
    assert os.path.exists(cert_pdf_path)
    
    # Read PDF text to verify RUC and Address are present
    reader = pypdf.PdfReader(cert_pdf_path)
    cert_text = ""
    for page in reader.pages:
        cert_text += page.extract_text() or ""
        
    assert test_ruc in cert_text, f"Expected RUC '{test_ruc}' in certificate PDF text"
    assert "Torres del Valle" in cert_text or "Naciones Unidas" in cert_text
    
    # 3. Test Financial Reports (PDF)
    pdf_rep_path = os.path.join(output_dir, "test_payments_rep_ruc.pdf")
    p_data = sm.get_payments_report_data()
    FinancialReportPDFGenerator.generate_payments_report_pdf(
        dest_path=pdf_rep_path,
        report_data=p_data,
        condo_name=test_condo_name,
        currency="$",
        condo_address=test_address,
        condo_ruc=test_ruc
    )
    assert os.path.exists(pdf_rep_path)
    rep_reader = pypdf.PdfReader(pdf_rep_path)
    rep_text = ""
    for page in rep_reader.pages:
        rep_text += page.extract_text() or ""
    assert test_ruc in rep_text, f"Expected RUC '{test_ruc}' in payments report PDF"
    assert test_condo_name.lower() in rep_text.lower(), "Expected condo name in payments report PDF"
    assert "naciones unidas" in rep_text.lower(), "Expected address in payments report PDF"
    
    # 4. Test Financial Reports (Excel)
    excel_rep_path = os.path.join(output_dir, "test_payments_rep_ruc.xlsx")
    FinancialReportExcelGenerator.generate_payments_report_excel(
        dest_path=excel_rep_path,
        report_data=p_data,
        condo_name=test_condo_name,
        currency="$",
        condo_address=test_address,
        condo_ruc=test_ruc
    )
    assert os.path.exists(excel_rep_path)
    wb = openpyxl.load_workbook(excel_rep_path)
    ws = wb.active
    cell_a1 = str(ws['A1'].value or "")
    cell_a2 = str(ws['A2'].value or "")
    assert test_condo_name.lower() in cell_a1.lower()
    assert test_ruc in cell_a2
    assert "naciones unidas" in cell_a2.lower() or "república" in cell_a2.lower()
    wb.close()
    
    # 5. Test Payment Receipt (PDF) with RUC & Address
    receipt_pdf_path = os.path.join(output_dir, "test_receipt_ruc.pdf")
    ReceiptProcessor.generate_receipt_pdf(
        dest_path=receipt_pdf_path,
        receipt_no="REC-TEST-RUC-01",
        resident_name="Rodrigo Morales",
        unit="201",
        concept="Pago de Alícuota Mensual",
        amount=70.0,
        late_fee=0.0,
        reference="99887766",
        payment_date="2026-04-20",
        condo_name=test_condo_name,
        condo_address=test_address,
        condo_ruc=test_ruc
    )
    assert os.path.exists(receipt_pdf_path)
    rcpt_reader = pypdf.PdfReader(receipt_pdf_path)
    rcpt_text = ""
    for page in rcpt_reader.pages:
        rcpt_text += page.extract_text() or ""
    assert test_ruc in rcpt_text, f"Expected RUC '{test_ruc}' in payment receipt PDF"
    assert test_condo_name.lower() in rcpt_text.lower() or "naciones unidas" in rcpt_text.lower()
    
    # Clean up test artifacts
    for p in [cert_pdf_path, pdf_rep_path, excel_rep_path, receipt_pdf_path]:
        if os.path.exists(p):
            try:
                os.remove(p)
            except:
                pass

    # Restore original condo settings
    sm.update_settings(
        condo_name=orig_condo_name,
        condo_ruc=orig_condo_ruc,
        condo_address=orig_condo_address
    )
                
    print("[OK] Condo Address and RUC Database Configuration & Document Rendering test PASSED!")

if __name__ == "__main__":
    print("=== STARTING SYSTEM INTEGRATION TESTS ===")
    test_whatsapp_message_parsing()
    test_ocr_and_bank_receipt_parsing()
    test_late_fee_calculation()
    test_payment_reconciliation_and_pdf_generation()
    test_partial_payment_and_debt_reduction()
    test_undo_payment()
    test_advance_payment()
    test_other_incomes()
    test_board_members_exoneration()
    test_expenses_and_other_incomes_crud()
    test_pichincha_statement_parsing()
    test_misaligned_statement_fallback()
    test_xlsx_column_mapping_exclusion()
    test_delete_bank_transaction()
    test_fused_pdf_text_parsing()
    test_receipt_recipient_selection()
    test_incremental_aliquots_emission()
    test_historical_tenant_preservation()
    test_discrepancy_name_filtering()
    test_additional_charges_crud_and_app_mode()
    test_abonos_and_revenue()
    test_db_url_configuration()
    test_edit_aliquot_amount()
    test_role_access_control()
    test_database_config_encryption_and_initial_balances()
    test_delete_aliquot_superadmin_only()
    test_individual_aliquot_emission()
    test_reference_edit_and_auto_reconcile()
    test_financial_reports()
    test_no_debt_certificate_generation()
    test_condo_address_and_ruc_config_and_reports()
    print("\n=========================================")
    print("SUCCESS: ALL 31 TESTS PASSED SUCCESSFULLY! THE SYSTEM IS 100% OPERATIONAL.")
    print("=========================================")
